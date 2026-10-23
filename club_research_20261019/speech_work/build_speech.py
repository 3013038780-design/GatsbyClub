from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).parent;OUT=ROOT.parent/'output'
sections=[]
for s in json.load(open(ROOT/'speech_content.json')):
 if s['title'].startswith('活动概念逐项'):
  for k in range(0,len(s['body']),8):
   sub=dict(s);sub['body']=s['body'][k:k+8]
   first=sub['body'][0][:2];last=sub['body'][-1][:2]
   sub['title']='活动概念逐项备讲 '+first+'至'+last
   sections.append(sub)
 else:sections.append(s)
doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.75);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
sec.header_distance=Cm(.6);sec.footer_distance=Cm(.6)
for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Songti SC')
 pp=st._element.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
normal=doc.styles['Normal'];normal.font.size=Pt(11.5);normal.paragraph_format.line_spacing=1.28;normal.paragraph_format.space_after=Pt(8)
doc.styles['Title'].font.size=Pt(25);doc.styles['Title'].paragraph_format.space_after=Pt(13)
doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 1'].paragraph_format.space_after=Pt(13)
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(10);doc.styles['Heading 2'].paragraph_format.space_after=Pt(7)
header=sec.header.paragraphs[0];header.text='北京香港私享俱乐部  网站讲解稿';header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('合伙人讨论  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
def para(t,size=None,bold=False,after=8):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(after);r=p.add_run(t);r.bold=bold
 if size:r.font.size=Pt(size)
 return p

def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[12.5,4.3]
 for i,(x,w)in enumerate(zip(headers,widths)):t.columns[i].width=Cm(w);t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=x
 for j,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if j==0:pr.append(OxmlElement('w:tblHeader'))
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'233E50' if j==0 else ('F1F4F5' if j%2==0 else 'FFFFFF'));cp.append(sh)
   bd=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bd.append(x)
   cp.append(bd);m=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'50' if side in ['top','bottom'] else '100');x.set(qn('w:type'),'dxa');m.append(x)
   cp.append(m)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.15
    if i==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
     r.font.size=Pt(9.5)
     if j==0:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 para('',size=2,after=0)

def link(p,label,url):
 rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');rp.append(sz);r.append(rp);tx=OxmlElement('w:t');tx.text=label;r.append(tx);h.append(r);p._p.append(h)

# Opening page includes practical use instructions and clickable Site.
p=doc.add_paragraph('网站配套讲解与合伙人讨论');p.paragraph_format.space_after=Pt(8)
for r in p.runs:r.font.size=Pt(10)
doc.add_heading('北京香港私享俱乐部网站讲解稿',0)
para('讲解人  项目发起人     听众  合伙人与核心执行团队',size=10,after=10)
para('使用方式：正文按第一人称撰写，可直接照读；每节的网页位置与操作提示用于同步演示，备讲用于解释思考和回应追问。金额沿用网站进取定价版，原始调研日期为2026年10月19日。',size=10,after=10)
z=para('',after=15);link(z,'打开配套网站  私密访问需登录','https://beijing-hongkong-private-circle.easy-tern-3117.chatgpt.site')
doc.add_heading(sections[0]['title'],1)
para('网页位置  '+sections[0]['where']+'  |  '+sections[0]['op'],size=9.5,after=12)
for x in sections[0]['body']:para(x)
para(sections[0]['note'],size=9.2,after=0)

for idx,s in enumerate(sections[1:],1):
 tag=doc.add_paragraph(s['where']);tag.paragraph_format.page_break_before=True;tag.paragraph_format.keep_with_next=True;tag.paragraph_format.space_after=Pt(6)
 for r in tag.runs:r.font.size=Pt(9.5);r.bold=True
 doc.add_heading(s['title'],1)
 para('操作提示  '+s['op'],size=9.5,after=13)
 for x in s['body']:
  if s['title'].startswith('活动概念逐项'):
   first,rest=x.split('。',1);z=para(first+'。',bold=True);z.add_run(rest).bold=False
  elif s['title'].startswith('合伙人可能') and '\n' in x:
   first,rest=x.split('\n',1);z=para(first,bold=True);z.add_run('\n'+rest).bold=False
  else:para(x,size=10.5 if s['table'] else None,after=6 if s['table'] else 8)
 if s['table']:table(*s['table'])
 if s['note']:
  n=para(s['note'],size=9.2,after=0);n.paragraph_format.line_spacing=1.2

doc.core_properties.title='北京香港私享俱乐部网站讲解稿'
doc.core_properties.subject='网站逐节口播 操作提示 商业思考 财务演示 合伙人问答'
doc.core_properties.author=''
doc.save(OUT/'北京香港私享俱乐部网站讲解稿.docx')
print('Created',OUT/'北京香港私享俱乐部网站讲解稿.docx')
