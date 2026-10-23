from pathlib import Path
import json, math, csv
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output'
OUT.mkdir(exist_ok=True)

# All financial parameters below are proposed operating assumptions, not market quotations.
PRODUCTS=[
 dict(id='dinner',name='精选主题晚宴',unit='人',cap=16,price=1680,var=620,fixed=5200),
 dict(id='game',name='原创任务及角色夜',unit='人',cap=20,price=1380,var=420,fixed=6200),
 dict(id='city',name='城市线索协作日',unit='人',cap=24,price=1280,var=380,fixed=8500),
 dict(id='couple',name='伴侣共创工作坊',unit='对',cap=6,price=5280,var=1980,fixed=8000),
 dict(id='retreat',name='三天两晚主题旅行',unit='人',cap=12,price=11800,var=7100,fixed=22000),
 dict(id='party',name='季度主题派对',unit='人',cap=60,price=1680,var=680,fixed=36000),
 dict(id='hk',name='香港联动私享晚宴',unit='人',cap=16,price=2980,var=1380,fixed=16000),
 dict(id='private',name='私人组局策划费',unit='单',cap=1,price=19800,var=9900,fixed=0),
]
SCHEDULE={
 'dinner':[1,1,1,1,1,1,2,2,2,2,2,2],
 'game':[0,1,0,1,0,1,0,1,1,1,1,1],
 'city':[0,0,1,0,1,0,1,0,0,1,0,0],
 'couple':[0,1,1,0,1,1,0,1,1,0,1,1],
 'retreat':[0,0,0,0,0,1,0,1,0,1,0,1],
 'party':[0,0,0,0,0,1,0,0,0,0,0,1],
 'hk':[0,0,0,0,0,0,1,0,0,0,1,0],
 'private':[0,0,0,0,0,0,1,1,1,1,1,1],
}
SINGLE=[0,0,0,4,5,6,7,7,7,8,8,8]
COUPLE=[0,0,0,1,1,2,2,2,3,3,3,3]
SCENARIOS={
 '保守':dict(single=32,couple=8,counts=[12,4,2,4,2,1,0,2],paid=[10,12,14,4,8,36,10,1],fixed=42000,start=90000,capital=300000),
 '基准':dict(single=60,couple=20,counts=[18,8,4,8,4,2,2,6],paid=[13,16,19,5,10,48,13,1],fixed=60000,start=160000,capital=600000),
 '增长':dict(single=110,couple=40,counts=[24,12,6,12,6,3,4,10],paid=[15,18,22,6,12,54,15,1],fixed=95000,start=250000,capital=1200000),
}
def dist(total, weights):
    raw=[total*w/sum(weights) for w in weights]
    ans=[int(x) for x in raw]
    for i in sorted(range(12),key=lambda i:raw[i]-ans[i],reverse=True)[:total-sum(ans)]:ans[i]+=1
    return ans
def model(name,pricefactor=1.,seatfactor=1.,fixedextra=0):
    s=SCENARIOS[name]
    ns=dist(s['single'],SINGLE); nc=dist(s['couple'],COUPLE)
    schedules={p['id']:dist(n,SCHEDULE[p['id']]) for p,n in zip(PRODUCTS,s['counts'])}
    monthly=[]; cumfees=0.; revfees=0.; people=0; cash=s['capital']; previous_prepay=0
    event_refund=.03
    for m in range(12):
        fee=ns[m]*9800+nc[m]*16800
        cumfees+=fee; people+=ns[m]+2*nc[m]
        fee_rev=cumfees/12.; revfees+=fee_rev
        gross=cost=nextcost=0.
        for p,paid in zip(PRODUCTS,s['paid']):
            n=schedules[p['id']][m]
            seats=paid if p['id']=='private' else paid*seatfactor
            gross+=n*seats*p['price']*pricefactor
            cost+=n*(seats*p['var']+p['fixed'])
            if m<11:nextcost+=schedules[p['id']][m+1]*(seats*p['var']+p['fixed'])
        refunds=gross*event_refund
        membercost=people*600/12+(ns[m]+2*nc[m])*150
        start=s['start']*([.5,.25,.25]+[0]*9)[m]
        fixed=s['fixed']+fixedextra
        profit=gross-refunds+fee_rev-cost-membercost-fixed-start
        prepay=.1*nextcost
        delta=cashflow=gross-refunds+fee-cost-membercost-fixed-start-(prepay-previous_prepay)
        previous_prepay=prepay;cash+=cashflow
        deferred=cumfees-revfees
        reserve=.05*deferred
        monthly.append(dict(month=m+1,new_single=ns[m],new_couple=nc[m],people=people,member_cash=fee,member_revenue=fee_rev,event_gross=gross,refund=refunds,event_cost=cost,member_cost=membercost,fixed=fixed,start=start,profit=profit,cashflow=cashflow,cash=cash,deferred=deferred,reserve=reserve,available=cash-reserve,prepay=prepay))
    result=dict(name=name,monthly=monthly,member_cash=sum(r['member_cash'] for r in monthly),member_revenue=revfees,event_gross=sum(r['event_gross'] for r in monthly),event_net=sum(r['event_gross']-r['refund'] for r in monthly),event_cost=sum(r['event_cost'] for r in monthly),member_cost=sum(r['member_cost'] for r in monthly),profit=sum(r['profit'] for r in monthly),cashflow=sum(r['cashflow'] for r in monthly),end_cash=monthly[-1]['cash'],min_cash=min(s['capital'],*[r['available'] for r in monthly]),deferred=monthly[-1]['deferred'],reserve=monthly[-1]['reserve'])
    result['revenue']=result['event_net']+revfees
    return result
M={n:model(n) for n in SCENARIOS}
B=M['基准']
assert sum(SINGLE)==60 and sum(COUPLE)==20
assert all(sum(SCHEDULE[p['id']])==n for p,n in zip(PRODUCTS,SCENARIOS['基准']['counts']))
for name,r in M.items():
    assert abs(r['cashflow']-(r['profit']+r['deferred']))<1e-6
    assert abs(r['end_cash']-SCENARIOS[name]['capital']-r['cashflow'])<1e-6
(ROOT/'financial_model.json').write_text(json.dumps(dict(products=PRODUCTS,scenarios=SCENARIOS,results=M),ensure_ascii=False,indent=2))
with (ROOT/'financial_monthly.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(B['monthly'][0]));w.writeheader();w.writerows(B['monthly'])

SOURCES=[
('S01','BoL用户提供资料','本次对话附件及用户陈述','', '五张截图；不盈利为用户提供的信息，未获得财务报表。'),
('S02','Modernist 官网活动','官方活动记录','https://www.modernist.club/','混合历史与新活动，未标年份的项目不作当月日历使用。'),
('S03','Modernist Gallery','官方历史图库','https://www.modernist.club/gallery','Gatsby、Studio 54及爵士晚宴等历史活动。'),
('S04','The Supper Club Membership','官方公开会籍页面','https://members.thesupperclubinc.com/membership','美元报价；分档权益及750美元入会费；未获得经营利润。'),
('S05','The Supper Club Terms','官方服务条款','https://thesupperclubinc.com/terms-of-services','Passport旅行另收费；以对应方案最终条款为准。'),
('S06','Parlor Membership Overview','官方会员手册','https://parlor.social/Parlor%20Membership.pdf','无固定地点、兴趣分组、分档；手册未列清晰有效期。'),
('S07','Mandala Club Calendar','官方当前活动及空间页','https://www.mandala.club/','含2026年10月活动；已过去与尚未举行项目分开理解。'),
('S08','Mandala Member Benefits','官方会员说明','https://www.mandala.club/enquire','个人、伴侣及30岁以下方案；未公开标准价格。'),
('S09','67 Pall Mall Singapore Events','官方活动日历','https://www.67pallmall.com/the-singapore-club/club-events/','盲品、WTF威士忌调配等；实际活动价格需会员登录。'),
('S10','67 Pall Mall Membership','官方会费说明','https://www.67pallmall.com/the-singapore-club/membership/','标准年费S$4800及入会费S$2000；Joint资格需咨询。'),
('S11','The Battery Societies','官方兴趣社群目录','https://legacy.thebatterysf.com/societies','Singles、Supper Club、Terroir等。'),
('S12','The Battery Events','官方活动日历','https://thebatterysf.com/events','含2026年10月爵士、台球、主题派对等。'),
('S13','The Battery Charter','官方会籍章程','https://thebatterysf.com/charter','会籍、来宾、资料与社区规则；不推断盈利。'),
('S14','Carlyle and Co','官方网站','https://www.carlyleandco.com/','成员定位及申请；未使用媒体旧报价作为现价。'),
('S15','Chef Nero Family Feast','官方历史活动页','https://login.carlyleandco.com/Chinese_New_Year/Chef_Nero_s_Family_Feast','2025蛇年活动；HK$1398，酒配另加HK$400。'),
('S16','Estelle Community','官方权益页','https://estellecommunity.com/','城市及乡间空间联动；活动数量属于品牌自述。'),
('S17','The Arts Club Culture','官方艺术文化页','https://theartsclub.co.uk/arts-culture/','音乐、早餐谈话、工作室参访；详细日历会员登录。'),
('S18','The Offline Club France Retreat','官方历史产品页','https://www.theoffline-club.com/retreat-locations/5-day-retreat-in-france','2026年夏季历史产品€795—1095；页面天数标题与夜数表述不完全一致。'),
('S19','Secret Cinema About','官方体验说明','https://www.secretcinema.com/about','参与故事世界；规模数字是品牌自述，未用于市场测算。'),
('S20','Crystal Maze Live Experience','官方线下体验页','https://the-crystal-maze.com/london/the-experience/','将综艺机制转为售票线下产品的案例。'),
('S21','Taskmaster','官方节目与任务目录','https://taskmaster.tv/shows','以节目任务介绍为玩法依据，未进行整季逐集评测。'),
('S22','Game Changer','官方Dropout节目页','https://watch.dropout.tv/game-changer-new','每轮规则变化、参与者探索规则的机制。'),
('S23','The Fortune Hotel','ITV官方节目发布','https://www.itv.com/presscentre/presscentre/media-releases/fortune-hotel-opens-its-doors-itv1-and-itvx','2023公布的首季格式；用于机制研究。'),
('S24','The Devils Plan','Netflix官方介绍','https://about.netflix.com/en/news/the-devils-plan-death-room-trailer','2025年第二季策略与联盟介绍。'),
('S25','Jet Lag Tag Eur It','Nebula官方节目页','https://nebula.tv/jetlag/season/3','欧洲追逐主题；不假设节目全部规则可直接用于商业活动。'),
('S26','The Amazing Race','CBS官方节目介绍','https://www.cbs.com/shows/amazing_race/about/','双人队、旅行与任务机制；不照搬淘汰和高强度。'),
('S27','Culinary Class Wars','Netflix官方节目说明','https://media.netflix.com/en/only-on-netflix/81728365','料理竞赛；非普通会员烹饪能力评估工具。'),
('S28','Love Is Blind Pods','Netflix官方制作介绍','https://www.netflix.com/tudum/articles/love-is-blind-pods?inapp=true','声音先行交流；原节目求婚机制不纳入产品。'),
('S29','Come Dine With Me','官方节目频道说明','https://www.youtube.com/watch?v=VokDXK6MSf0','2026-09-02官方频道视频说明，轮值做东及评分。'),
('S30','The Traitors Missions','USA Network官方节目文章','https://www.usanetwork.com/usa-insider/traitors-best-missions-season-1','团队任务及身份博弈，只借鉴虚构游戏机制。'),
('S31','Mandala Instagram','官网链接并实读公开主页','https://www.instagram.com/mandalaclub.sg/','已查看简介、精选栏目、网格与日期；具体帖子全文触发登录。'),
('S32','Carlyle Instagram','官网链接并实读公开主页','https://www.instagram.com/carlyleandco/','已读公开简介、What’s on及Drink and dine栏目；深帖未完整读取。'),
('S33','Battery Instagram','官网链接并实读公开主页','https://www.instagram.com/thebatterysf/','已读主页及Membership、Battery Powered栏目；不推断转化率。'),
('S34','Offline Club Instagram','官网链接并实读公开主页','https://www.instagram.com/theoffline_club/','已读Retreats、Your City、Events等公开栏目；不以粉丝数估营收。'),
('S35','Arts Club Instagram','官网链接并实读公开主页','https://www.instagram.com/the_arts_club/','已读公开简介、网格和日期；详细活动以官网补证。'),
]
(ROOT/'source_register.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2))

doc=Document()
sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
styles=doc.styles
for el in list(styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
for el in list(doc.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[st].font.name='Arial';styles[st].font.color.rgb=RGBColor(0,0,0)
    styles[st]._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'),'Songti SC')
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.line_spacing=1.35
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Title'].font.size=Pt(27)
styles['Heading 1'].font.size=Pt(20)
styles['Heading 1'].paragraph_format.space_after=Pt(13)
styles['Heading 2'].font.size=Pt(12)
styles['Heading 2'].paragraph_format.space_before=Pt(9)
styles['Heading 2'].paragraph_format.space_after=Pt(5)
styles['Subtitle'].font.size=Pt(12)
header=sec.header.paragraphs[0]
header.text='北京香港私享社交俱乐部  |  研究与商业方案'
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026年10月19日  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
page_counter=0
def page(title,tag=None):
    global page_counter
    needs_break=page_counter>0
    page_counter+=1
    if tag:
        p=doc.add_paragraph(tag);p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.page_break_before=needs_break
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.font.size=Pt(9);r.bold=True
    heading=doc.add_heading(title,0 if page_counter==1 else 1)
    if not tag:heading.paragraph_format.page_break_before=needs_break
def p(t,bold=False,small=False):
    z=doc.add_paragraph();r=z.add_run(t);r.bold=bold
    if small:
        r.font.size=Pt(8.5);z.paragraph_format.line_spacing=1.2
    return z
def h(t):doc.add_heading(t,2)
def table(headers,rows,widths=None,small=9,padding=80):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for j,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        if j==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        for i,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[i])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'233E50' if j==0 else ('F2F5F6' if j%2==0 else 'FFFFFF'));pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for e0 in ['top','bottom','left','right']:
                e=OxmlElement('w:'+e0);e.set(qn('w:w'),str(padding if e0 in ['top','bottom'] else 80));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.17
                for run in para.runs:
                    run.font.size=Pt(small)
                    if j==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def link(para,label,url):
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245475');pr.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size);run.append(pr)
    tx=OxmlElement('w:t');tx.text=label;run.append(tx);el.append(run);para._p.append(el)
def money(n):return f'{n/10000:,.1f}'

page('北京香港私享社交俱乐部','海外案例研究与商业计划')
p('给创始人与合伙人的经营建议',True)
p('以精选成员、原创共同体验和长期关系服务为核心。北京形成稳定活动与会员密度，香港承担精选跨城体验与合作节点。单身、恋爱中和已婚会员均有清楚的参加理由。')
h('建议先做什么')
p('先成立有审核与主理人服务的轻资产俱乐部，使用已有及合作餐厅、酒吧与活动空间；以三类可复购活动验证付费，再开放年度会籍。把婚恋作为明确自愿的产品线，把伴侣与共同兴趣产品同时做好。')
h('研究带来的三个判断')
p('一是活动安排本身能成为会员产品：The Supper Club与Parlor值得借鉴。二是同一批成员可以在多种小社群里留下来：The Battery提供了结构参照。三是有趣可以被设计：规则、协作、轮值做东与共同作品，比单纯升级场地更值得优先投入。[S04—S06，S11，S21—S30]')
h('先看现金流的真实来源')
p(f'按本方案的基准假设，首年经营收入约{money(B["revenue"])}万元、税前亏损约{money(-B["profit"])}万元。会员预收款使现金表现更好，但年末仍有约{money(B["deferred"])}万元未履约会费。当前假设支持小规模验证，尚不足以把项目称作稳定盈利的现金流业务。',True)
p('所有本项目价格、人数与成本为测试假设；海外品牌能公开收费，不代表其实际盈利已经被证实。BoL的调性作为参照；其不盈利的信息由本项目发起人提供，未经财务核验。[S01]')
p('2026年10月19日  |  北京为主场  香港联动',small=True)

page('如何使用这份方案','阅读导航')
table(['部分','解决的问题'],[
('案例与证据 3至10页','BoL、9个海外会员俱乐部、5个官方Instagram主页观察；分清能借鉴的机制和不能推出的结论。'),
('玩法与产品 11至20页','10档节目、3类线下体验，以及32个活动创意和8个重点产品。'),
('经营与测算 21至27页','会员权益、北京香港运营、单场经济、首年三情景、月度现金与次年季度测算。'),
('试点与合伙 28至30页','90天验证、投入门槛、成员关系管理、人员分工和合伙人决策。'),
('来源索引 31至34页','35条来源记录；官网、条款、官方节目页与Instagram逐项可回溯。')],[3.2,13.8],10)
h('证据的使用方式')
p('正文的[S编号]对应末尾来源索引。官方页面证明其公开产品或表述；用户提供资料用于理解BoL。活动的商业价值、北京香港适配度及本项目方案，是基于这些材料作出的判断。')
p('Instagram已直接检查Mandala、Carlyle、Battery、Offline Club及Arts Club的公开主页、栏目与可见网格。部分帖子全文要求登录；对其正文和互动转化不作未经核实的描述。节目研究依据官方规则、任务介绍和发布资料，不把本方案包装成全季观片评分。')
h('本轮最关键的资料缺口')
p('没有海外俱乐部可比的经审计单店利润，也没有BoL的成本明细；不据此宣称哪一家“已经验证盈利”。本项目的实际可触达成员、价格接受度、场地内部结算和主理人供给，需要在90天内通过实际付费与工时记录补齐。')
p('市场测算从你能邀请的人开始：愿意来一次的人、愿意再来的人、愿意购买会籍的人。泛婚恋人口和高净值总人数不直接转化为可获得收入。')

page('BoL提供了什么参照','本地起点')
p('五张截图显示了很清楚的产品层次：母品牌使用海归与菁英的商业和生活成长定位；活动有行业酒会、高尔夫社交及AI交流；BoL Selected单独承接审核制单身私宴。小程序同时提供VIP、发现和活动入口。[S01]')
table(['可确认的设计','对本项目的启发'],[
('母品牌覆盖较宽','单身、伴侣及已婚成员可以共同参加明确的兴趣活动；婚恋产品保留独立入口。'),
('Selected强调审核与限定','筛选应对应活动需求和实际可交付的成员组合，避免只有文案稀缺。'),
('AI体验区与6—8人驻桌','做有内容的小桌交流，给成员完整表达与听见他人的机会。'),
('联合主办','可结合不同主理人的成员、内容和空间；具体分工与结算需要另行设计。')],[4.3,12.7],10)
h('用户补充的经营事实')
p('本项目发起人提供的信息是BoL并不盈利；目前没有可判断亏损来源的成本资料。票价、人员配置、招募成本和场地成本，都需要分别核验。')
h('我们需要新增的经营能力')
p('同一场活动能否在计入招募、策划、现场服务和售后后仍有贡献利润；会员能否因为持续的服务购买会籍；活动能否由第二位主理人稳定复现。这三个问题，比复制视觉风格更决定经营结果。')
p('BoL资料中的“10w+”“90%以上海归”记为品牌自述，不直接作为付费规模、已核实成员质量或本项目预测的输入。游艇画面也不等于拥有游艇资产。[S01]')

page('九个俱乐部的比较','研究样本')
table(['对象','主要机制','适合借鉴','投入判断'],[
('Modernist / 美国','会员晚餐、主题夜、行业交流','成员主理与专业／玩乐混合','有会所；用合作场地改写'),
('The Supper Club / 美国','精选餐桌、分档会籍、旅行','活动次数与权益定价','优先研究活动交付'),
('Parlor / 美国','无固定地点、兴趣分组、App','轻资产城市社交日历','接近首期结构'),
('Mandala / 新加坡','日常仪式、文化活动、伴侣会籍','高低频结合、共同兴趣','会所成本不可直接照搬'),
('67 Pall Mall / 新加坡等','葡萄酒专业内容、活动、旅行','主题专业性、动手体验','先做酒类专题系列'),
('The Battery / 美国','会籍内的兴趣与身份小组','单身线与长期社群并行','成员主理需有人管理'),
('Carlyle and Co / 香港','文化、音乐、餐饮与私密服务','主人的接待感与持续体验','作为服务参照与合作方向'),
('Estelle Community / 英国','城市社交与乡间度假联动','城市相识后周末加深','复制路径；保持轻资产'),
('The Arts Club / 英国','艺术、音乐、参访与家庭内容','全年内容编排与家庭延伸','借专业伙伴建立内容')],[4.1,4.3,5.2,3.4],9)
p('来源[S02—S17]。上述机制可以在公开材料中找到；会员满意度、单店盈利与本项目可复制性不能由品牌存在时间或装修规模推出。')
h('优先次序')
p('经营结构优先研究The Supper Club和Parlor；活动组合研究Modernist与Mandala；会员分组研究The Battery；服务与审美研究Carlyle、Estelle和Arts Club。67 Pall Mall帮助建立“专业内容＋共同体验”的细分产品。')

page('Modernist让成员成为内容','深度案例 01')
h('事实')
p('官网活动记录包括会员私人晚餐、Gatsby主题派对、爵士晚宴、校友社交与AI主题交流；图库保留了部分历史活动。记录混有不同年份，不将整页当作2026年当月日历。未核实其统一现行会费及利润。[S02—S03]')
h('值得学的地方')
p('成员既来参加活动，也能带来自己的兴趣、朋友和内容。专业交流之后有餐桌与派对，安静的人和外向的人都可能找到适合自己的时段。由此可形成多种连接入口，减少每场都从陌生人开始的招募负担。')
h('北京版本')
p('一次24人的“未来客厅”：前45分钟由一位成员展示真实项目或作品，随后四张6人餐桌围绕四个生活问题交流，最后进入可自由选择的音乐与饮品时段。技术内容只占开头一段，主理人控制节奏，避免变成连续路演。')
p('每个月固定一个小主题，邀请一位成员与一位外部嘉宾共同主理。候选内容包括AI创作、建筑、电影、美食、体育与城市文化；成员贡献列入权益和激励，但不强迫每人提供资源。')
h('盈利怎么检验')
p('先按场售票，明确内容、餐饮和主理人服务各自的成本。专业嘉宾的加入是否提高购票与复购，要与普通晚宴比较；赞助只在实际签约后入账。主理人的邀请与协调工时必须记录。')
h('不能误学的地方')
p('名人、校友标签与热闹照片容易制造表面吸引力，却不能保证每位成员得到一次完整的交流。我们应限制总人数和单桌人数、安排明确主持，并向双方确认后再做活动后的引荐。')

page('把活动日历做成会员产品','深度案例 02')
h('The Supper Club的收费结构')
p('公开会籍页列出：首次申请的Quattro为US$2500／年含本州4次活动；Explorer为US$5000／年含全国8次；另有不限次数方案及US$750入会费。旅行需结合条款看额外收费，不能从“有旅行权益”理解成费用全包。[S04—S05]')
h('Parlor的组织结构')
p('官方手册介绍无固定地点、精选活动、兴趣Circles及移动端入口，并区分Standard与Premium。手册不构成本轮可核实的现行收费表。[S06]')
h('经营启发')
p('会员付费理由可以是省下每次寻找局、判断人和安排时间的精力。你们需要持续交付一份可信、适合且能预订的社交日历，并把“哪些活动已经含在会费内”说清楚。')
table(['方案','用户得到什么','经营约束'],[
('按次售票','低承诺尝试一场','需要持续获客，收入波动大'),
('有限次数系列卡','一段连续体验及可预期价格','需要同批成员的出席与替补机制'),
('年度会籍＋活动另购','持续筛选、组局与优先预约','不能只卖报名资格，需有基础交付'),
('不限次数','高使用频率会员更容易理解','重度用户挤占容量，暂不建议首期')],[4,6,7],10)
h('首期选择')
p('用单次售票验证前三个月；第四个月达到门槛后开放有限权益会籍。年度活动均在付款前列出容量和预约方式；优先预约保留48小时窗口，不承诺所有活动都有位置。')

page('Mandala把社交变成固定仪式','深度案例 03')
h('事实')
p('官网2026年10月日历列有First Pour新旧会员相聚、Proudly Pitch a Friend朋友推荐单身嘉宾、黑胶聆听、调酒、跑步及家庭午餐；会员说明另设个人、伴侣与30岁以下类型，未公开完整现行会费。[S07—S08]')
h('值得学的地方')
p('大活动用来形成记忆，小活动让人有熟悉的到访理由。“每月知道哪一天可以来”能减轻成员安排生活的负担；不同关系状态的人可以在日历里找到自己的位置。')
h('北京版本的四个固定栏目')
table(['栏目','具体组织','收费方向'],[
('第一杯','新会员与老会员同桌，主持人介绍共同兴趣','会员基础小桌权益，限次数'),
('朋友介绍朋友','自愿报名的朋友用3分钟介绍单身嘉宾，双方确认后联系','单身专场门票'),
('周三客厅','黑胶、桌面协作、人物故事轮换；一晚只做一个主题','按场票或系列卡'),
('周末共创','伴侣、朋友分别成组制作菜单、短片或旅行路线','工作坊或包场费')],[3.2,9,4.8],10)
h('对会员生命周期的价值')
p('脱单后可以转入伴侣会籍，继续参加餐饮、音乐和运动活动；原来的朋友仍在。关系状态变化只改变适合参与的产品，不需要重新证明自己属于这个社群。')
p('首期只固定一至两个高频栏目，做出稳定体验后扩展。维护十个空群和办十个浅活动，都会消耗主理人；每个栏目都有负责人、人数上限与取消阈值。')

page('细分兴趣与城市乡间联动','深度案例 04')
h('67 Pall Mall')
p('新加坡官网标准年费为S$4800 nett，通常另收S$2000入会费，部分类别例外。活动日历有香槟盲品及威士忌调配体验；该体验让参与者带走个人调配的小瓶作品。会籍与活动票需分别核对。[S09—S10]')
p('启发：一项真正可参与的兴趣，可以让成员长期回来。北京可先做“风味实验室”四期：嗅觉、盲品、搭配、成员共同创作。每期都能单独参加，也能串成连续体验；提供等价的无酒精选项。')
h('Estelle Community')
p('官网把城市会所与乡间Estelle Manor连接，并介绍餐饮、康体及会员活动。会员使用多个空间的权益，是其公开的产品结构；不据此推断新项目需要自建同类设施。[S16]')
p('启发：先在城里认识，再选择一个周末一起出行。旅行的售卖对象应优先来自已经参加过活动的人，降低陌生人对长时间共处的顾虑。北京可以先联动近郊，再从香港组织短途体验。')
h('本项目的差异化测试')
p('同时测试“稀缺场地”与“共同作品”谁更影响购买：用成本相近的两场活动，一场侧重场所，一场侧重参与；比较付款转化、完整交谈次数、再次报名及转介绍。样本只支持方向判断，不能拿小样本做成功率宣传。')
h('资本投入选择')
p('首年不以会所资产为核心。场地合作按单场结算，建立两套替补供应商；购买的是按需使用的服务和时间。把会员费投入在可稳定复现的邀请、策划与接待能力上。')

page('长期社群怎样容纳单身与伴侣','深度案例 05')
h('The Battery')
p('官方社群目录中，Singles与Supper Club、Terroir及其他兴趣小组并列；活动日历包括爵士、台球、游戏夜与主题派对。章程也规定了会籍、来宾及社区使用方式。[S11—S13]')
p('启发：单身不是会员唯一的身份。成员可以因为吃饭、音乐和运动持续联系；单身产品有自己的筛选和报名，不必让每次聚会承担撮合任务。')
h('Carlyle and Co')
p('官网强调成员的个性、兴趣和故事。2025年的官方厨师共桌晚宴公开价为HK$1398，配酒另加HK$400；这只是历史活动价格，不能当作现行会费或本项目定价。[S14—S15]')
p('启发：好的接待感来自主理人知道谁来、谁适合坐在一起、谁希望安静，以及什么时候应该介绍两个人认识。把这些服务写进现场流程，才能让“私享”成为可感知的体验。')
h('The Arts Club')
p('官网列出音乐、早餐谈话、艺术工作室参访等文化内容，并有家庭相关入口；完整会员日历需要登录。[S17]')
p('启发：伴侣与家庭内容可以增加留存，但首年以成人共同体验为主。亲子托管、儿童课程和家庭旅行需要另一套能力，放入后续产品池，不在首期会籍中承诺。')
h('共同原则')
p('年龄、家庭背景和职业是活动组合的输入；同时要看参加动机、好奇心、交流方式和实际出席。用成员之间愿意再次见面的关系质量，评价俱乐部是否正在形成。')

page('Instagram的表达可以怎样学','官方账号观察')
table(['已实读的主页','公开可见内容','对本项目的启发'],[
('@mandalaclub.sg','What’s On、Race Week HQ、Mandala Minis、Membership；网格有活动与人物内容','活动做成有名字的连续栏目，人物故事与活动信息并行'),
('@carlyleandco','What’s on、Drink and dine、Curious等栏目；简介列Harbour与Town Club','建立统一的主人语气，用好奇和体验吸引成员'),
('@thebatterysf','Membership、Battery Powered、About栏目','把参加、归属与社群价值分别说清楚'),
('@theoffline_club','Retreats、Your City、Events、Vibes等栏目','从内容页清楚进入城市活动与旅行，不只展示气氛'),
('@the_arts_club','公开简介、活动与餐饮网格及发布日；深帖需登录','用餐饮、文化和人物组合出持续的生活场景')],[4,6,7],9.5)
p('观察范围为公开主页、栏目与可见网格；具体帖子全文、视频完整内容、私信及会员内容未完整取得。以上不包含社交媒体转化率或运营效果推断。[S31—S35]')
h('本项目的内容结构')
p('建议每月12条内容：4条活动预告、3条去标识的真实互动片段、2条经同意的成员／主理人故事、2条活动方法或作品、1条会籍说明。分别引向申请、具体活动和旅行咨询；先记录来源与付款，不追求粉丝数字。')
h('调性落到细节')
p('视觉使用一致的邀请函、餐桌、光线和现场动作；文案写出“这晚具体做什么、谁适合来、怎么参加”。成员的家庭财富和职业名头不作为主要宣传素材。可拍摄的人与时段提前约定，给安静参与者同样完整的体验。')
p('主页开设活动日历、如何加入、单身专场、伴侣共创、周末旅行五个清楚入口。年度会籍的服务内容与限制直接可读，避免只能私信询问。')

page('十个节目里值得转化的玩法','综艺机制 01')
table(['节目与证据','好玩的机制','拟转化的原创产品'],[
('Taskmaster [S21]','同一目标可以有意想不到的解法，短任务让不同能力被看见','周末任意门：一组完成构图、声音、道具任务，再共同展示'),
('Game Changer [S22]','参与中发现每轮变化的规则，错误也能形成共同笑点','今晚规则会变：渐进公开规则的客厅小游戏，给每人跳过权'),
('The Fortune Hotel [S23]','双人合作、有限线索、交换选择，在漂亮酒店里形成连续悬念','秘密行李箱：箱内是故事线索与体验券，晚宴后揭晓，无现金下注'),
('The Devils Plan [S24]','信息与资源不对称使成员需要沟通、协商和合作','智识客厅：三轮原创逻辑与协作任务，结束后复盘不同解法'),
('Jet Lag [S25]','城市交通与地理本身成为游戏棋盘','城市密码：在预先踩点的区域收集线索，以公共交通和步行串联')],[4.5,6,6.5],9.5)
h('如何筛选适合高端会员的玩法')
p('优先选择规则可在五分钟解释、能形成多次小组交流、体力门槛低、可以中途休息、结束后有共同记忆的机制。保留悬念与惊喜，但事前说明活动强度、时长和大致类型。')
h('保持参与者的主体感')
p('公开表现不舒服的人可以承担摄影构思、道具、记录或观察任务；设置小组成果和趣味奖，避免公开给人气、财富、魅力打分。虚构博弈只发生在游戏内，不能用于评价真实诚信。')
p('上述节目属于机制来源。对外活动使用独立主题、原创任务和视觉，不宣称为官方综艺体验或授权项目。')

page('从旅行到关系的五种机制','综艺机制 02')
table(['节目与证据','可用的互动结构','拟转化的原创产品'],[
('The Amazing Race [S26]','两人协作、路线选择与分段任务','双人城市协作日：每段轮换领航，不设置赶不上就出局'),
('Culinary Class Wars [S27]','食材限制、团队分工、作品呈现','一晚限定餐厅：用一个主题共同设计两道菜和待客方式'),
('Love Is Blind [S28]','先听声音，再与外在形象相遇','声音先行夜：三轮主题对话，随后自愿见面和自由交谈'),
('Come Dine With Me [S29]','轮值做东使同一批人体验不同待客方式','四席轮值晚宴：四期由不同成员主理餐桌，专业厨房执行'),
('The Traitors [S30]','合作任务、隐藏身份与晚宴仪式串联','庄园谜夜：虚构角色与信息交换，所有人全程参与')],[4.5,6,6.5],9.5)
h('三个适合先试的组合')
p('轻量组合：Taskmaster＋Game Changer，适合第一次见面，成本主要在主持与原创设计。餐桌组合：料理竞赛＋轮值做东，适合你已有的餐厅和酒吧。周末组合：城市任务＋秘密行李箱＋晚宴，适合较高客单与连续相处。')
h('从节目到活动的关键改写')
p('电视观众喜欢看淘汰与冲突；付费参与者更在意体验公平、参与时间和体面。产品需要让每个人都有事做、有自由选择，也有独处和休息空间。剧情只规定处境，不规定真实情感结果。')
h('伴侣专场的尺度')
p('旅行偏好、安排分工和未来想象可成为共同任务。深层关系沟通由双方私下完成；需要专业支持的环节由合适的关系工作者主持，不把公开揭露矛盾作为娱乐内容。')

page('三个真实的线下转化样本','体验产品参照')
h('Secret Cinema 参与故事世界')
p('官方介绍的核心是让观众进入故事环境并参与其中。[S19] 对本项目最有价值的是活动前的邀请、到场的角色入口、空间中的任务与最后的共同结局。首期可以只使用一家餐厅的三个区域，控制布景投入。')
h('The Crystal Maze Live Experience 规则可以售票')
p('该官方线下体验以团队任务与主题区域组织参与。[S20] 启发是把一套可复现的游戏系统做成产品，而非每场从零编一次综艺。你们可建立一套人物卡、任务模块和主持手册，再按主题更换外层故事。')
h('The Offline Club 留白本身也是产品')
p('官方法国旅行页展示小团体、共同餐桌和可自由选择的乡间时间，2026年夏季页面价格为€795—1095；属于历史产品，天数与夜数表述需向经营者确认。[S18] 其启发是密集任务之外，需要允许不表演、不发言也能舒服共处。')
h('建议的活动设计比例')
p('首轮周末原型可采用40%轻任务、30%自由交流、20%餐饮与共同作品、10%回顾。比例是设计起点，不是研究验证结论；用活动反馈调整。每一段写明为什么安排它，删掉只为了显得丰富的环节。')
h('一条可复购的路径')
p('第一次以好玩的主题来；第二次愿意和相识的人一起完成新体验；第三次主动带一个朋友来；随后才考虑会籍或旅行。报名记录应帮助主理人识别这条路径，不能仅记录总人次。')

ACTS=[
('声音先行','3轮声音交流后自愿见面','单身','半日'),('朋友提案','朋友用3分钟介绍单身嘉宾','单身','晚间'),('双人地图','轮换领航完成城市线索','单身','一日'),('秘密晚宴','虚构身份＋三幕线索＋共桌','单身／兴趣','晚间'),
('四席晚宴','同批成员四期轮值主理餐桌','兴趣／单身','系列'),('风味盲盒','用嗅觉线索寻找搭配伙伴','兴趣','半日'),('最后一杯','组队设计一款酒或无酒精饮品','兴趣','晚间'),('一晚餐厅','共同设计菜单及待客体验','兴趣／伴侣','半日'),
('任意门任务','开放式创意任务与小组作品','兴趣','半日'),('规则会变','渐进发现规则的客厅游戏','兴趣','晚间'),('庄园谜夜','角色线索与共同剧情结局','兴趣／伴侣','周末'),('秘密行李箱','交换线索箱解锁下一段体验','兴趣／伴侣','周末'),
('黑胶客厅','每人带一首歌和一个故事','兴趣','晚间'),('城市声音','收集声音后制作一段配乐','兴趣','一日'),('72秒短片','组队拍一个主题短片','兴趣／伴侣','半日'),('艺术侦探','依据作品线索完成原创故事','兴趣','半日'),
('看见彼此','各自策划对方喜欢的半天','伴侣','一日'),('婚前周末','预算、家庭期待与共同选择','伴侣','周末'),('未来地图','共同制作未来一年体验清单','伴侣','半日'),('周年之旅','用纪念物与路线重构共同回忆','伴侣','定制'),
('帆船轮值','教练带领换位协作与晚餐','兴趣／伴侣','一日'),('球场搭档','分级高尔夫或板式网球任务','兴趣','半日'),('山间来信','轻徒步、写信与小桌晚餐','兴趣／伴侣','周末'),('无屏周末','阅读、散步、长桌与留白','兴趣／伴侣','周末'),
('主人的一日','成员带小组体验其熟悉领域','兴趣','一日'),('幕后参访','真实流程任务＋经营者餐桌','兴趣','一日'),('收藏物语','用一件物件讲一个人生故事','兴趣','晚间'),('互助圆桌','每人一个问题、一段经验分享','兴趣','系列'),
('港岛密码','街区故事任务与隐藏餐桌','兴趣／伴侣','一日'),('北京来客','北京成员迎接香港小组共创','兴趣','周末'),('季末盛宴','前几期共同作品与朋友派对','全体分区','晚间'),('朋友家宴','两对伴侣或朋友共同待客','伴侣／朋友','系列'),
]
for i in range(2):
    page('三十二个活动概念'+(' 上' if i==0 else ' 下'),'原创活动池')
    table(['活动','核心玩法','参加者','时长'],ACTS[i*16:(i+1)*16],[3,8,3.3,2.7],9.3)
    p('本页为待验证创意。单身活动仅面向确认有相识意愿者；伴侣活动以双方共同报名为基础；兴趣活动明确开放范围。表中的斜线表示可以分场改编，并非默认混场。',small=True)

CARDS=[
('精选餐桌','16人／3小时／建议测试1680元每人','会员或审核访客；单身版本单独招募。','提前收集一个愿意谈的话题；18:30迎接与两人轻聊，19:00四人桌第一段用餐，19:45按共同兴趣调换一次座位，20:30主理人发起共同选择，21:00自由交流。','主理人1、接待1、餐厅服务。需要安静度足够的分桌空间。固定成本5200元含主题准备、现场引导与增量场地，餐饮等620元每人。','至少10个实付席位才确认首场；13人时单场贡献8580元，未扣3%退款假设与固定团队费。','后续向每人发送下一期适合的活动；双方确认后引荐。衡量至少两次完整交谈、30日再次付款和推荐。'),
('今晚规则会变','20人／3小时／建议测试1380元每人','单身或兴趣分场；不要求表演经验。','4人小组完成构图、听觉和道具三轮原创任务；每轮只新增一个规则，成员可选择解题、制作、记录或展示；最后共同揭晓解法。','主持1、执行2；道具箱与可移动桌椅。固定6200元，餐饮道具等420元每人；主持人提前彩排。','至少12个实付席位；16人时贡献9160元。每人至少参与两轮，不能因早期输赢离场。','同一套模块可换皮成电影、城市、节日主题。观察是否因任务而愿意再次见面，而非只评价主持人好笑。'),
('城市密码','24人／6小时／建议测试1280元每人','按步行和交流偏好招募；伴侣与单身分场。','区域限在已踩点的3公里范围，4人队完成声音、风味、观察三站任务；中途休息并换领航；终点为长桌晚餐或早晚餐。','领队1、站点协作2、备援1；固定8500元包含踩点、人员和任务制作，餐饮交通等380元每人。','至少16个实付席位；19人时贡献8600元。雨天切换室内三站，事前告知备用方案。','不同街区可形成三期；记录同队相识与后续复购。活动评价不使用跑速或体力排名。'),
('伴侣共创工作坊','6对／4小时／建议测试5280元每对','恋爱中、订婚或已婚双方共同报名；以共同体验为目的。','独立写下各自期待，随后共同设计一份纪念日晚餐或旅行；在预算、自由时间与家庭参与三个情境中协作；完成后私下交流并制作一张共同计划卡。','关系工作者1、主理人1、餐饮人员；固定8000元，餐饮材料等1980元每对。不承担关系治疗承诺。','至少4对实付；5对时贡献8500元。双方都可跳过问题；不向其他成员展示私人答案。','下次产品可换成角色互换、周年之旅或长期共同计划。分别征求两人的体验反馈，避免只听付款者。'),
('秘密行李箱周末','12人／三天两晚／建议测试11800元每人','参加过至少一次活动者优先；单身或伴侣分别成团。','入住领取角色与线索箱；第一晚轻任务，次日分段协作与自由体验，晚宴揭晓故事；第三天由成员自由选择散步、创作或安静早餐。','旅行执行伙伴、主理人及现场协作；固定22000元，酒店餐饮交通等7100元每人。报价不含到出发城市的大交通。','至少8人实付再确认，最多12人；10人时贡献25000元。单身默认独立入住，不安排陌生人强制同房。','旅行后安排一次轻聚；记录推荐、返场和每位成员的实际参与时间。需要可退出任务与室内替代段。'),
('香港私享联动','16人／晚宴＋半日可选／建议测试2980元每人','香港本地成员为基础，北京访客自愿参加。','一位当地成员共同主理，前半段探索一个城市主题，晚间分桌交流；次日参访单独选择、单独报价。','香港主理人＋活动制作伙伴；固定16000元包括协调、北京执行人员差旅分摊及场地，变动1380元每人。均为人民币预算。','至少12人实付；13人时贡献4800元。若只靠北京成员飞行补位，则推迟，先建立香港本地招募。','香港当地活动与北京互访分别记录。会员自己的机酒自理，付款前提供明确费用范围。'),
('季末主题盛宴','60人／4小时／建议测试1680元每人','已参与成员及受邀朋友；单身意愿另行自选。','展示本季短片、音乐或菜单作品；前段小桌接待，中段共同节目，后段自由音乐与交流；设置安静区域。','制作经理、接待与执行伙伴；固定36000元，餐饮等680元每人。人多不等于有效连接。','至少40人实付；48人时贡献12000元。首期不靠派对赚钱，前面的小活动形成稳定返场后再办。','衡量朋友申请与付费转化，扣除免费嘉宾及品牌招待席。赞助不进入基准收入。'),
('私人组局服务','每单8—16人／建议策划服务费19800元起','成员发起的生日、朋友聚会或纪念日；明确目的和名单范围。','需求面谈，提供两套可比流程，确定邀请、餐桌、活动与服务，再执行与回顾。基础报价包含一次改稿和一次现场主理。','供应商由客户直接结算；本项目只将策划费入收入。单均交付成本假设9900元，超出服务范围另报价。','每单贡献9900元；预留20—30小时服务容量。不把供应商总消费当成平台收入。','通过成熟的模块组合提高效率；受邀者自愿申请俱乐部，不默认进入会员库。'),
]
for i in range(4):
    page('重点产品'+str(i*2+1)+'与'+str(i*2+2),'八个可测试原型')
    for card in CARDS[i*2:(i+1)*2]:
        name,spec,audience,flow,resource,threshold,follow=card
        h(name);p(spec,True)
        for label,txt in [('适合谁',audience),('怎样玩',flow),('交付与成本',resource),('确认门槛',threshold),('后续连接',follow)]:
            z=p(label+'  '+txt);z.paragraph_format.space_after=Pt(5)
    p('本页价格、成本和贡献为本项目设计假设；贡献尚未扣团队工资、创业投入、税费及退款。须以报价、实付与工时复核。',small=True)

page('三天两晚怎样让剧情服务关系','旅行原型的节奏')
table(['时间','共同体验','成员仍保留的选择'],[
('周五15:00—18:00','到达、独立入住、领取线索箱；主理人轻介绍','可先休息；不录制真实隐私访谈'),
('周五18:30—21:00','四人桌晚餐＋第一幕故事；一次轻任务','不参与表演者可做观察、道具或记录'),
('周六09:30—12:00','两段小组任务，单身场可自愿换搭档','体力相近分组，有坐下完成的等价任务'),
('周六12:00—16:00','午餐和至少两小时自由时间；可选散步、料理或阅读','允许独处，不按出席任务数量评价成员'),
('周六16:30—21:00','选择一项共同作品；晚宴解锁最终线索','游戏内身份当晚结束，保留友好回顾'),
('周日09:00—12:00','自由早餐、写给未来的信、离开','单身意愿私下双向确认；伴侣反馈由双方私下完成')],[4.2,7.2,5.6],9.8)
h('单身版与伴侣版分别修改什么')
p('单身版增加自愿轮换、小组相识与活动后双向引荐；伴侣版增加共同选择、互相主理半天和未来计划。共享景色、餐饮和原创故事制作能力，招募和关系服务分别组织。')
h('故事如何复用')
p('核心故事保留三幕、六个角色功能、三套任务模块；可以更换城市、季节、食物与艺术元素。不要每次重做一整套制作。每场回顾记录哪一段需要主持救场、哪一段让成员自然继续交流，下一次据此修改。')

page('会员为哪些持续服务付费','会籍设计')
table(['类型','建议测试价','明确交付'],[
('审核访客','无年费，活动单购','适合具体活动即可申请，不拥有浏览全部成员的权限。'),
('个人会籍','9800元／12个月','面谈建档、每季一次需求更新、每年2次会员小桌、48小时优先预约、定向组局与双向引荐请求服务。'),
('伴侣会籍','16800元／两人／12个月','双方独立资料与共同会籍，每人2次会员小桌、共创活动优先预约、每季一次共同需求更新。'),
('私人服务','按具体服务另报价','专门引荐、私人聚会及旅行策划单独定义交付与人数，不捆绑无限承诺。')],[3.1,4.1,9.8],10)
p('全年2次小桌权益按会籍年计算，每次12席，计划在开放会籍后每月举办2场。基准年末100名实际成员对应完整会籍年200席权益，全年滚动容量需至少216席。首年晚入会者的剩余权益延续至各自会籍到期。')
h('9800元为什么可能成立')
p('价格检验的是持续省心的社交安排、被了解的需求与可信的引荐服务。每季需求更新必须真的发生；无法稳定安排基础小桌和回应请求，就暂缓收年度会费。与三个月付费活动相比，年费要带来独立增量价值。')
h('关系状态变化')
p('单身转伴侣时，可以按剩余期限补足价差升级；双方均保留独立资料。伴侣关系变化后，以剩余权益转为个人服务方案，事前约定结算，避免出现另一人被删除或被迫分享信息。')
p('会籍不自动获得相亲资格；单身产品每次确认参加意愿。引荐需要双方同意，回应与安排属于交付，成功恋爱或商业合作不属于结果承诺。')

page('北京主场怎样连接香港','获客与城市密度')
h('北京的起点')
p('首期选择一个交通方便、能保持安静交流的合作餐厅或酒吧作为固定见面点，并准备两个替补场地。创始人的校友、朋友和既有餐饮资源是候选招募渠道，使用前确认实际可触达人数与邀请意愿，不假设机构背书。')
h('香港的起点')
p('先找一位能持续做成员接待的当地合作主理人，建立30—40名候选人的直接联系，再测试一场12—16人活动。北京成员的出行意愿作为额外需求，不能承担填满香港活动的基本盘。首年只测两场香港活动，先不承担第二套全职团队。')
table(['渠道','最初动作','实际记录'],[
('创始人直接邀请','每周10—15个一对一邀请，不群发泛名单','受邀、回复、付费、到场、再次付费'),
('成员推荐','参加过一次后，自愿推荐1位朋友','每个推荐人的实际付款与成员体验'),
('合作主理人','与艺术、餐饮、运动或行业内容伙伴共创一场','单场成本、来源、二次报名、分成'),
('社交内容','以活动细节、共同作品及自愿人物故事引向申请','内容来源到报名及付款的转化')],[3.3,7.2,6.5],10)
h('真实供给比名册更重要')
p('婚恋线按年龄段、长期所在地、关系意愿和彼此择偶范围看可行组合。名单里有人不等于会来，也不等于双方互相愿意认识。某组缺乏合适组合时，先候补和补充招募，允许该组活动延后。')
p('非单身线独立记录共同兴趣和双方意愿，不依赖单身产品的转化来证明需求。两条线同等研究、分别核算，资本分配随实际复购和交付能力调整。')

page('单场活动的收入与成本','单位经济')
rows=[]
for pr,paid,n in zip(PRODUCTS,SCENARIOS['基准']['paid'],SCENARIOS['基准']['counts']):
    cm=paid*(pr['price']-pr['var'])-pr['fixed']
    be=math.ceil(pr['fixed']/(pr['price']-pr['var'])) if pr['fixed'] else 1
    rows.append((pr['name'],f'{pr["price"]:,}/{pr["unit"]}',f'{paid}/{pr["cap"]}',f'{pr["var"]:,}',f'{pr["fixed"]:,}',f'{cm:,}',str(be)))
table(['产品','单价元','付费／容量','单位变动','场次固定','单场贡献','贡献打平'],rows,[3.6,2.4,2.2,2.2,2.2,2.2,2.2],8.8)
p('单场贡献＝实付单位数×（单价－单位变动成本）－场次固定成本。打平指场次贡献为零，未覆盖全职团队、创业投入、税费和退款；不能作为公司盈利门槛。私人服务只记录服务费，供应商由客户直接付款。',small=True)
h('成本边界')
p('变动成本包括餐饮、按人材料、住宿及按人交通；场次固定包括外部主持或专业人士、原创制作分摊、现场执行和增量场地。长期成员运营及常规制作工资进入月固定费，外部执行与长期岗位不重复计算。关联餐厅按可持续内部结算计入，不能按零场地、零服务算利润。')
h('旅行与香港的口径')
p('基准旅行按整包体验收入和全部相关供应商成本列示，11800元包含两晚住宿、餐食和活动，不含到出发城市的大交通；若改为代理模式，则只计策划／代理费，相关供应商流水不再列收入。香港样本2980元为人民币预算，正式销售另做港币报价，不直接套汇率。')
p('全年活动收入统一扣3%的退款／改期损失假设；付款席位之外的免费嘉宾另列成本。赞助、投资收益和未签约的品牌收入在基准情景中为零。')

page('首年三种经营结果','财务情景')
table(['指标 万元','保守','基准','增长'],[
('年末个人／伴侣账户','32／8','60／20','110／40'),
('年末实际成人成员','48','100','190'),
('活动及私人服务场次','27','52','77'),
('会费现金收入',*[money(M[n]['member_cash']) for n in M]),
('会费履约收入',*[money(M[n]['member_revenue']) for n in M]),
('活动净收入',*[money(M[n]['event_net']) for n in M]),
('经营收入合计',*[money(M[n]['revenue']) for n in M]),
('活动直接成本',*[money(M[n]['event_cost']) for n in M]),
('会员履约与核验成本',*[money(M[n]['member_cost']) for n in M]),
('全年团队固定费','50.4','72.0','114.0'),
('一次性启动费用','9.0','16.0','25.0'),
('税前经营结果',*[money(M[n]['profit']) for n in M]),
('假设投入资金','30.0','60.0','120.0'),
('年末现金余额',*[money(M[n]['end_cash']) for n in M]),
('年末未履约会费',*[money(M[n]['deferred']) for n in M]),
('年内最低可用现金',*[money(M[n]['min_cash']) for n in M]),
],[6.8,3.4,3.4,3.4],9.3)
p('均为设计假设，非报价或回报承诺。经营结果在启动费用当期列支后、税前计算；未包含融资和股东分红。表中可用现金已减去未履约会费5%的退款流动性准备；未来履约成本仍需按滚动预算预留。首年假设无年度续费，次年才检验留存。',small=True)

page('基准情景的每月现金','首年滚动测算')
table(['月份','新增个／对','会费收款','活动净收','经营结果','期末现金','未履约会费'],[
(f'M{r["month"]}',f'{r["new_single"]}／{r["new_couple"]}',money(r['member_cash']),money(r['event_gross']-r['refund']),money(r['profit']),money(r['cash']),money(r['deferred'])) for r in B['monthly']
],[2,2.5,2.5,2.5,2.5,2.5,2.5],9.2,padding=20)
p('金额单位万元；M1为实际启动月，并非指定自然月份。会籍从M4开放，新入会按完整年费收款，管理测算从当月起按12个月平均确认。活动当月完成并结算；预付下月直接活动成本的10%，基准模型M12余额未含下一年项目定金。',small=True)
h('为什么现金与利润不一样')
p(f'基准年会费收款{money(B["member_cash"])}万元，确认会费收入{money(B["member_revenue"])}万元，差额{money(B["deferred"])}万元对应后续月份的服务。年末现金{money(B["end_cash"])}万元包含股东投入和尚未履约的年费，不能直接作为可分配利润。')
h('固定费与创业投入')
p('基准每月6万元：两位全职岗位及雇主成本3.6万元、创始人劳务1.2万元、内容5000元、系统3000元、行政4000元。启动费16万元分三个月列支，覆盖原型、基础系统、品牌内容与招募；未含长期租赁押金或装修。')
p('流动性要求：每月滚动安排90天履约预算。未履约会费的5%仅用于退款压力测试；另需预留实际履约成本。未履约年费不用于股东分配。',small=True)

page('什么变化最容易使利润消失','压力测试与招募漏斗')
sens=[('基准',model('基准')),('实付席位平均下降20%',model('基准',seatfactor=.8)),('票价平均下降10%',model('基准',pricefactor=.9)),('固定运营每月增加1万元',model('基准',fixedextra=10000)),('活动及私人服务提价7%',model('基准',pricefactor=1.07)),('固定运营每月减少9000元',model('基准',fixedextra=-9000))]
table(['情形','税前经营结果 万元','变化原因'],[(a,money(r['profit']),b) for (a,r),b in zip(sens,['参照情景','场次固定成本不随空席下降','交付成本仍需发生','全年增加12万元固定支出','数量不变时接近打平；须验证需求','服务能力不下降时接近打平；须核实'])],[6.4,4.3,6.3],9.3)
h('需求端需要多少真实人')
p('基准一年680个付费成人参加人次，伴侣产品按两人计，另有6单私人服务。它们不是680个独立新客。假设每位独立付费参与者年均参加2.4次，需要约284位真实付费参与者；其中60个个人会籍与20个伴侣会籍共100人，不能重复计成新增客源。')
p('若300个可触达候选人中，60%表达兴趣、其中50%首次付款，则只有90位首购者。以每人2.4次计，无法支持基准活动量。需要持续转介绍和伙伴带来的新增客源，或主动减少场次。上述转化率用于压力测试，尚未获得验证。')
h('会籍也需要单独验证')
p('从约284位独立付费参与者转成100名实际成员，人数转化约35%。对一个新品牌，这是较有要求的假设。应分别看个人与伴侣转化，不能用免费嘉宾或重复报名充当新会员。提价或降成本只是单变量测算，不能同时假设不会影响需求和交付。')
h('投入要随验证节奏变化')
p('用户资源多，并不自动等于稳定的同时间出席与付费。保守情景减少活动和人员；增长情景则增加主理人及制作能力。首年最应避免的是先建大团队，再依靠更多活动补固定成本。')

page('次年要靠留存验证俱乐部','第二年季度模型')
# Illustrative second-year continuation: explicit cohort renewals, quarterly admissions, no fee increases.
active_s=60;active_c=20;cash2=B['end_cash'];deferred2=B['deferred'];qrows=[];year2profit=0
for q in range(1,5):
    # Cohorts expire M16..M24; Q2-Q4 renew original monthly cohorts at 70%, expected counts allowed.
    starti=(q-2)*3+3
    exp_s=sum(SINGLE[starti:starti+3]) if q>=2 else 0
    exp_c=sum(COUPLE[starti:starti+3]) if q>=2 else 0
    renew_s=.7*exp_s;renew_c=.7*exp_c
    active_s-=.3*exp_s;active_c-=.3*exp_c
    news=8;newc=3
    avg_s=active_s+news/2;avg_c=active_c+newc/2
    active_s+=news;active_c+=newc
    receipt=(news+renew_s)*9800+(newc+renew_c)*16800
    recog=(avg_s*9800+avg_c*16800)/4
    # Per quarter counts 6,3,1,3,1,1,1,2; same paid-seat assumptions as Year 1.
    qcounts=[6,3,1,3,1,1,1,2]
    gross=sum(n*paid*pr['price'] for n,paid,pr in zip(qcounts,SCENARIOS['基准']['paid'],PRODUCTS))
    direct=sum(n*(paid*pr['var']+pr['fixed']) for n,paid,pr in zip(qcounts,SCENARIOS['基准']['paid'],PRODUCTS))
    memcost=(avg_s+2*avg_c)*600/4+(news+newc*2)*150
    fixed=210000
    profit=gross*.97+recog-direct-memcost-fixed
    cfo=gross*.97+receipt-direct-memcost-fixed
    cash2+=cfo;deferred2+=receipt-recog;year2profit+=profit
    qrows.append((f'Q{q}',f'{active_s+active_c:.1f}',money(receipt),money(recog),money(gross*.97),money(profit),money(cash2)))
table(['季度','期末账户','会费收款','会费收入','活动净收','经营结果','期末现金'],qrows,[2,2.4,2.5,2.5,2.5,2.6,2.5],9)
p('单位万元；账户数允许期望值小数，不代表实际可以出售半个会籍。每季新增8个个人、3个伴侣账户；原首年到期批次续费率70%，未到期账户继续有效。新增按季内均匀入会近似确认收入，未使用票价涨幅。',small=True)
p('每季活动假设为6场晚宴、3场任务夜、1场城市、3场伴侣工作坊、1次旅行、1场派对、1场香港活动、2单私人组局；实付人数沿用首年基准。每月固定费从6万元升至7万元，体现服务容量扩张。上表为季度近似，实际按每个会籍起止日和合同履约更新。')
h('第二年的关键不是继续预售')
p(f'该组假设下，次年税前经营结果约{money(year2profit)}万元。应首先确认续费、老成员再次参加、新成员融入，以及活动质量能否由主理团队维持。若需要持续靠创始人免费投入和新增年费弥补交付，俱乐部尚未形成健康经营。')
h('增长的停止条件')
p('出现小桌权益连续两个月难预约、成员请求平均超过三个工作日未回应、主理人长期超负荷，先暂停新增会籍并补充交付。新增城市需当地持续供给、至少两场付费验证与一位可独立主持者，不以一次爆满扩张。')

page('团队如何稳定地交付','人员与系统')
table(['角色','主要责任','基准工时与结算'],[
('创始人','首批邀请、成员质量、内容方向及关键伙伴','每月约40小时已计劳务预算；额外投入另记录'),
('成员运营','申请面谈、需求更新、报名组合、售后引荐','全职；常规请求三个工作日内回应'),
('制作运营','场地与供应商、原创任务、现场流程、财务台账','全职；采购及执行容量每月预排'),
('活动主理人','主持、交流引导、现场观察与回顾','按场结算，费用进入场次固定成本'),
('香港合作主理人','本地邀请、场地、接待与伙伴协调','先按项目预算结算，不先设香港全职团队')],[3.5,7,6.5],9.5)
p('基准52场付费或私人项目，加会籍基础小桌约18场；平均每月约6场，高峰月份更密。普通场次约需15—25个团队工时，旅行另有伙伴执行；两位全职不能同时跨城或兼任全部专业角色。连续两个月累计超过可用工时80%，增加兼职或减场。')
h('首版小程序与后台')
p('先提供邀请申请、面谈与审核状态、活动报名／付款／候补、会籍权益、参加资格、活动反馈及双向引荐。后台记录关系阶段、城市、兴趣、出席、权益消耗和服务请求。早期用现成表单和会员工具跑通，再按真实需求定制。')
p('资料区分本人自述、已查看证明及待更新信息；成员自己选择分享范围。业务人员按服务需要访问，婚恋与兴趣参加资格独立。创始人私人择偶使用与其他成员相同的报名和双向意愿流程。')
h('合作与资源机会')
p('企业参访、经营经验、供需引荐可成为会员自愿选择的内容。投资合作另行判断项目，不作为入会回报或基准收入；成员信息与商业合作保持明确的使用目的。')

page('九十天验证与投入顺序','先取得实际付费证据')
table(['阶段','动作与人数','需要得到的证据'],[
('第1—14天','访谈24人：8名单身、8名伴侣双方或已婚参与者、8名伙伴／主理人；盘点候选名单','过去实际参加什么、付过多少钱、下一次愿为什么付款；双方意见分别记录'),
('第15—35天','一场精选晚宴＋一场任务夜；邀请转实付，不以免费满场验收','两个活动各达到确认门槛；记录招募、制作和服务工时'),
('第36—60天','一场伴侣共创＋重复一场表现更好的原型','至少4对实付，分别问双方体验；能由另一位主理人复现'),
('第61—90天','做一次连续体验的第二期；香港伙伴小型预访谈','30日再次付费率与推荐付费；香港具备12位明确付费意向后再开场')],[3,7.3,6.7],9.5)
h('继续投入的门槛')
p('至少三场收费活动达到各自确认线，合计至少40位独立付费成年人；在有后续场次可选的参与者中，30日再次付费率目标25%以上；至少70%反馈有两次以上愿意继续的完整交谈；退款／严重体验问题逐项解决。数字是经营门槛，不是行业基准。')
p('会籍开启另需：至少20份针对明确价格及权益的入会意向申请、未来六个月可公布的活动安排，以及基础小桌容量。正式开放后首批至少10个账户实付，才按基准扩充。若伴侣线或单身线未达标，分别修改产品，不以另一条线的成绩替代。')
h('90天投入上限')
p('按同一基准模型，建议先审批30万元净现金占用上限。M1—M3包含18万元固定运营费和16万元启动费，扣除活动净贡献、计入供应商预付款后，占用约28.7万元。达标后再释放年度60万元预算中的剩余额度。若选择兼职主理、现成工具及更少场次，须另做精简预算，不能声称仍交付同样的基准方案。')

page('合伙人需要拍板的事项','经营选择')
table(['事项','建议默认','需要怎样验证'],[
('母品牌','精选成员的社交俱乐部，单身与伴侣分入口','客户能否用一句话解释为何参加'),
('首期资源','合作场地，两位核心执行岗位按节奏到位','至少两套可持续场地报价及独立主持能力'),
('首个产品','晚宴、任务夜、伴侣共创同时测试','实付、交付成本、双方反馈分别核算'),
('会费','前三个月以按场为主，达门槛再开年费','独立服务价值与预约容量'),
('香港','合作主理人，首年最多两场测试','本地付费供给及差旅后贡献'),
('资金','先30万元试点，后续按里程碑释放','每月现金、未来履约与招募成本'),
('股权与激励','资本、成员来源、执行劳动分开讨论','先记录3个月实际贡献，再谈长期比例')],[3.3,7,6.7],9.5)
h('让贡献可结算')
p('资本按实缴记录；持续执行按岗位工资或项目费；成员来源奖励在真实收款、活动完成和退款期后结算，不出售成员名册。关联餐厅、酒吧及其他伙伴先确认报价、质量标准与替补条款。品牌合作与成员招待席逐场入账。')
h('把关系价值保留在可持续经营里')
p('创始人可以得到更丰富的朋友关系、行业经验和合作机会；成员也应得到同样明确的价值。社群成立的标志，是两个人即使不依赖你在场，也愿意因为共同体验继续联系，并愿意为下一次值得的安排付费。')
p('最适合现在的动作，是用高质量的小规模产品获得真实付款和复购证据，再决定会籍规模、人员和香港频次。视觉调性决定第一次是否有吸引力，交付与关系质量决定能否长期经营。',True)

for i in range(0,len(SOURCES),9):
    page('资料来源与核验记录 '+str(i//9+1),'来源索引')
    for sid,title,kind,url,note in SOURCES[i:i+9]:
        z=p(sid+'  '+title,True);z.paragraph_format.space_after=Pt(2)
        z=p(kind+'。'+note,small=True);z.paragraph_format.space_after=Pt(2)
        if url:
            z=doc.add_paragraph();link(z,'打开原始来源',url);z.paragraph_format.space_after=Pt(7)
    p('访问与整理日期：2026年10月19日。公开价格保留原币种，不作为北京香港供应商报价；网页更新后以经营者现行说明为准。',small=True)

doc.core_properties.title='北京香港私享社交俱乐部研究与商业方案'
doc.core_properties.subject='海外俱乐部 综艺机制 产品设计与经营测算'
doc.core_properties.author=''
doc.core_properties.keywords='北京 香港 私享社交 俱乐部 商业计划'
path=OUT/'北京香港私享社交俱乐部研究与商业方案.docx'
doc.save(path)
print(json.dumps({'file':str(path),'planned_pages':page_counter,'scenarios':{n:{k:r[k] for k in ['revenue','profit','cashflow','end_cash','deferred','min_cash']} for n,r in M.items()}},ensure_ascii=False,indent=2))
