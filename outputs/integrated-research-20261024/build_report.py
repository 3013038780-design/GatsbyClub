from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
B=Path(__file__).parent;D=json.loads((B/'统一研究数据.json').read_text());O=B/'交付文件';Q=B/'qa';S={s['id']:s for s in D['sources']};used=[]
font='Source Han Serif CN'
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name,size,bold in [('Normal',11.5,False),('Title',25,True),('Subtitle',12,False),('Heading 1',19,True),('Heading 2',16,True),('Heading 3',12.5,True)]:
 st=d.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor(0,0,0)
 rf=st.element.get_or_add_rPr().get_or_add_rFonts()
 for k in list(rf.attrib):del rf.attrib[k]
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),font)
 st.paragraph_format.line_spacing=Pt(19);st.paragraph_format.space_after=Pt(7)
 if name=='Title':st.paragraph_format.line_spacing=Pt(34);st.paragraph_format.space_after=Pt(12)
 if name.startswith('Heading'):st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12)
 for n in st.element.xpath('.//w:pBdr'):n.getparent().remove(n)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('高端线下社群市场调研与对标报告  ');r.font.size=Pt(8.5)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def p(t,lead=False):
 q=d.add_paragraph()
 if lead and '：'in t:
  a,b=t.split('：',1);q.add_run(a+'：').bold=True;q.add_run(b)
 else:q.add_run(t)
 return q
def h(t,l=2):return d.add_paragraph(t,'Heading '+str(l))
def page(t,l=2):
 q=h(t,l);q.paragraph_format.page_break_before=True;return q
def link(q,label,url,anchor=None,size=9):
 z=OxmlElement('w:hyperlink')
 if url:z.set(qn('r:id'),q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 elif anchor:z.set(qn('w:anchor'),anchor)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'255878');pr.append(col);sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz);r.append(pr);tx=OxmlElement('w:t');tx.text=label;r.append(tx);z.append(r);q._p.append(z)
def cite(ids):
 q=d.add_paragraph();q.paragraph_format.line_spacing=Pt(13);q.paragraph_format.space_after=Pt(5)
 for i,k in enumerate(ids):
  s=S[k]
  if k not in used:used.append(k)
  if i:q.add_run('  ')
  if s['url']:link(q,k+' '+s['name'],s['url'])
  else:q.add_run(k+' '+s['name']+'  微信小程序内查看').font.size=Pt(9)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tc.append(mar)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF1');tc.append(sh)
   for q in c.paragraphs:
    q.paragraph_format.line_spacing=Pt(15.2);q.paragraph_format.space_after=Pt(2)
    for r in q.runs:r.font.size=Pt(10);r.bold=(i==0)
 bor=OxmlElement('w:tblBorders')
 for side in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'CDD4DA');bor.append(e)
 t._tbl.tblPr.append(bor)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def picture(name):
 q=d.add_paragraph();q.paragraph_format.line_spacing=1;q.add_run().add_picture(str(Q/name),width=Cm(16.5))
F=FontProperties(fname='/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
plt.rcParams.update({'axes.spines.top':False,'axes.spines.right':False,'axes.spines.left':False,'font.size':11})
fig,ax=plt.subplots(figsize=(9.5,2.8));labels=['中国内地','中国含港澳台'];vals=[10.6,12.6]
ax.barh(labels,vals,color=['#25485D','#A5B8C2'],height=.45)
ax.set_yticks(range(2),labels,fontproperties=F);ax.set_xlim(0,14.5);ax.set_xlabel('亿元总资产家庭  万户',fontproperties=F)
for i,v in enumerate(vals):ax.text(v+.2,i,str(v),va='center')
ax.set_title('财富人群规模  截至2025年1月1日',fontproperties=F);fig.tight_layout();fig.savefig(Q/'财富人群.png',dpi=170);plt.close(fig)
fig,ax=plt.subplots(figsize=(9.5,3));labels=['25岁及以下','26至30岁','31至35岁','36至40岁','41至45岁','45岁以上'];vals=[8,19,24,19,18,12]
ax.bar(range(6),vals,color='#31556A');ax.set_xticks(range(6),labels,fontproperties=F,fontsize=10);ax.set_ylim(0,29);ax.set_ylabel('样本占比 %',fontproperties=F)
for i,v in enumerate(vals):ax.text(i,v+.6,str(v)+'%',ha='center')
ax.set_title('2026胡润品质生活调查样本年龄  n=470',fontproperties=F);fig.tight_layout();fig.savefig(Q/'样本年龄.png',dpi=170);plt.close(fig)

fig,ax=plt.subplots(figsize=(9.5,3));ys=[2021,2022,2023,2024,2025];vs=[75,83,91,98,103.5]
ax.plot(ys,vs,color='#31556A',linewidth=2.4,marker='o');ax.set_xticks(ys);ax.set_ylim(65,115);ax.set_ylabel('市场规模  亿元人民币',fontproperties=F)
for x,y in zip(ys,vs):ax.text(x,y+2,str(y),ha='center')
ax.set_title('互联网婚恋市场规模  网经社电数宝口径',fontproperties=F);fig.tight_layout();fig.savefig(Q/'婚恋市场五年.png',dpi=170);plt.close(fig)

fig,ax=plt.subplots(figsize=(9.5,2.9));labels=['独自报名','自述单身并愿交友','居住大城市'];vals=[80,60,75]
ax.barh(labels,vals,color=['#31556A','#5E7B8E','#9AAEB9'],height=.5)
ax.set_yticks(range(3),labels,fontproperties=F);ax.invert_yaxis();ax.set_xlim(0,100);ax.set_xlabel('品牌自述占比 %；各指标人群可重叠',fontproperties=F)
for i,v in enumerate(vals):ax.text(v+1,i,str(v)+'%',va='center')
ax.set_title('FTLO旅行者结构  官网未列调查年份与样本量',fontproperties=F);fig.tight_layout();fig.savefig(Q/'旅行者结构.png',dpi=170);plt.close(fig)

d.add_paragraph('高端线下社群\n市场调研与对标报告','Title')
d.add_paragraph('商业社交 婚恋交友 社交旅行','Subtitle')
d.add_paragraph('2026年10月24日  合伙人讨论材料','Subtitle')
h('本次结论')
p('建议围绕同一批高资产客户建立持续线下社群，商业交流、婚恋专场和旅行分别设置参加目的与准入。共享会员服务和活动日历，先在一个可触达客户的城市验证。',True)
p('首次付费的理由要具体：商务客买适合讨论真实问题的同桌；单身客买可信的同行条件与认识机会；旅行客买愿意相处的人、舒适的行程和有人负责组织。资产条件用来审核，报名理由还要落实到每场活动。')
p('再次参加依靠持续匹配和共同经历。BoL的活动日历、Founder Feast的同桌联系、社群从城市活动延伸度假，都值得学习。建议先按场与按团测试，积累真实再次参加记录后再决定年费权益。')
table(['方向','主要学习对象','我们优先学什么'],[['商业','BoL、Founder Feast','会籍与活动分层；人工筛选、小桌配对和餐费分账'],['婚恋','BoL Selected、钻石婚恋','半天私宴、男女准入与互选；顾问和资源分工'],['旅行','奇语、OWR；补充Offline Club','小团内容、行前准备；社群主持人与目的地合作交付']],[2.2,6.5,8.1])
p('研究判断：值得开展小规模付费验证。首阶段重点是获得一批适配客户并把一次相处做好；现有资料尚不足以支持重资产会所、高价顾问团队或同时铺开多城。具体收费以产品规格、可落实成本和客户反馈共同确定。',True)

page('阅读目录')
toc=[]
for label in ['第一部分 市场调研','1 市场规模与需求基础','2 客户生活与消费偏好','3 商业社交供给与收费','4 婚恋市场与客群比较','5 旅行市场与产品比较','第二部分 对标','6 选择依据与名单','7 BoL商务会籍','8 Founder Feast','9 BoL Selected','10 钻石婚恋','11 奇语旅行','12 Other Way Round','13 综合判断','附录 既有成果与资料入口']:
 q=p('');toc.append((q,label));q.add_run(label)
p('Excel先看“市场调研”和“对标产品”两张表。机构名录供扩搜回查，来源与互动记录供核验，不需要从原始数据逐条读起。')

page('第一部分 市场调研',1);h('1 市场规模与需求基础')
p('现有付费供给至少有三种：商业社交按会籍、课程或活动收费；婚恋按活动、顾问或定制项目收费；旅行按团期和包含服务收费。它们在同一客户的生活中相遇，但市场统计口径并不一致。')
picture('财富人群.png')
table(['指标','最新可用数值','统计范围与时间'],[
 ['亿元总资产家庭','内地10.6万户','胡润模型估计；截至2025-01-01，报告2026-06-30发布'],
 ['结婚登记','676.5万对，增长10.8%','2025全年年度终值；2026年9月公报'],
 ['国内居民出游','65.22亿人次','2025全年，全国；按人次'],
 ['国内旅游花费','6.30万亿元','2025全年，全国国内居民出游花费']],[4.1,5,7.7]);cite(['A02','A01','A04'])
p('对我们的意义：A8若按人民币资产亿元理解，应进一步确定是个人还是家庭资产。内地10.6万户是家庭总资产估计，只用于衡量财富人群基础；不能直接乘以单身比例和会费得出营业规模。',True)

page('1 市场份额应怎样判断')
p('目前没有取得覆盖中国高端线下社群、口径一致的总收入与企业前五收入资料。因此，本报告不编制CR5，也不把粉丝量当作市场份额。这个缺口只影响份额排序，仍可按具体产品比较客户、价格、经营方式和可见持续运营。')
table(['我们要了解什么','本版能够回答','需要新增的证据'],[
 ['总体有多大','官方婚姻与旅游统计、财富家庭数量及消费调查','高资产人群实际购买线下社群服务的年度支出'],
 ['主要玩家是谁','商业、婚恋、旅行各类供给地图；总Excel收录'+str(len(D['market']))+'条机构与品牌线索','同年度、同地区和同服务定义的收入排名'],
 ['谁做得更好','有日期的产品、活动、会员与续费披露分别比较','统一口径到场、付费会员、留存及贡献数据'],
 ['我们能进入哪里','明确准入的小型晚宴、餐桌社群与旅行组织机制','本地付费客户反馈和真实供应商成本']],[3.5,7,6.3])
h('城市不能混在一起判断',3)
p('胡润报告列北京亿元家庭18200户、上海16500户；广东16700户是全省，不能当深圳市。商业组织的地域也不同：EO华东以上海和长三角为中心，BoL本轮核查的是深圳产品。未来选择首发城市，应同时看可触达客户和能交付的活动，不能只按全国财富排名。');cite(['A02','B02','BOL'])
p('市场名单与对标名单分开：市场名单保留行业广度，对标名单才回答我们现在向谁学。只要某品牌仍有独特收费或组织机制，未进入重点名单也不会从资料库消失。')

page('2 客户生活与消费偏好')
picture('样本年龄.png')
p('2026胡润品质生活调查访问470位千万资产人士，其中70位亿元资产以上；78%来自北京、上海、广州和深圳。样本平均年龄36岁，41%已婚已育。客户范围包含有家庭的人，支持商业和生活方式活动独立于婚恋专场存在。')
table(['客户生活线索','调查结果','对活动组织的启发'],[
 ['时间','平均每周工作45小时，每年假期21天','缩短低效环节；旅行要值得占用有限假期'],
 ['消费方向','计划服务体验支出增加12%','把价值落实到同行者和交付，而非只展示豪华场地'],
 ['运动','徒步16%、跑步15%；高尔夫和网球各7%','不预设所有高资产客人都爱高尔夫'],
 ['休闲','温泉／SPA／足疗、运动、独处是主要放松方式','安排可选择的参与深度与休息空间']],[3.3,6,7.5]);cite(['A03'])
p('以上是调查样本，不作为某机构付费客户画像。每家客户描述在对标部分分别列年龄、工作、资产或收入、性格及兴趣；没有披露的字段保留为空。这样既能写具体，也不把职业想象成收入。')

page('3 商业社交有哪些供给')
table(['组织及付费对象','主要买到什么','本次经营证据'],[
 ['长江商学院\n企业决策者','教育项目、同学共同学习和校友网络','当前2026级收费和申请条件可读；累计学员不是当前付费人数'],
 ['EO华东\n企业所有者和创始人','固定同伴交流、培训及全球组织资源','当前官网110名；全球2024／25续会率91.6%，与华东口径分列'],
 ['木兰汇与年会\n女性经营者和商业参与者','论坛、路演、同行接触和考察','2026第18届；近千成员等为组织自述'],
 ['北京香港马会会所\n个人或公司登记代表及家属','会所使用与可预约服务，消费另结','会籍和月费明确，未披露当前有效会员与利润'],
 ['Founders Network\n分经营阶段科技创业者','导师、线上学习、分会聚会及高档闭门活动','自述600+创始人；有当期活动日历，无利润披露']],[5,6.1,5.7]);cite(['B01','B02','B04','B05','B07','PO13'])
p('这些供给共同说明：客户付费依据可以是共同学习、可信同行、长期小组或日常使用。社群的活动数量只是交付的一部分；是否能把适合的人组织在一起，决定了同类活动的差异。')

page('3 商业社交具体怎样收费')
table(['产品','收费单位','包含与另付'],[
 ['长江2026级EMBA','99.8万元／完整项目\n申请费2000元','按教育项目招生录取；不能当年费社群报价'],
 ['EO全球基础费用','2630美元／年\n首次另3500美元','地方分会会费及活动费用按当地另列'],
 ['木兰2026年会普通','8800元／该届活动','简餐、无桌卡；5800早鸟适用期已结束'],
 ['木兰2026年会VIP','17800元／该届活动','3次自助、前排桌位姓名牌、赠1普通票、视频和合影'],
 ['木兰2026年会SVIP','37800元／该届活动','嘉宾桌餐、前排沙发、赠2普通票、正装照、视频与合影'],
 ['Founders Network','75／134／249美元月展示价，按年付','年额900／1608／2988美元；更高档增加大会、导师及闭门活动'],
 ['BoL年费','488／2888元／年','基础酒会与更多活动资格分层；私宴另付'],
 ['Founder Feast','39.99加元／月等','晚餐席位和匹配；餐酒另付']],[4.2,5.1,7.5]);cite(['B01','EOFEES','B06','B07','BOL','B12'])
p('定价差异来自卖的东西不同：同样叫“会员”，可以是教育、会所准入、同伴服务，也可以只是餐桌组织。我们需要先明确一笔费用覆盖哪些工作，再比较相近产品，不能把上表做成一个行业平均价。')

page('4 婚恋市场要分清哪些产品')
table(['机构与产品','客户及当前收费证据','本轮判断'],[
 ['钻石婚恋','整体定位含个人或家庭千万以上及条件优秀者；爱情定制30万、60万、100万起、家办1000万起','高资产服务深度参照；单次活动价另查'],
 ['BoL Selected','本次男士家庭5000万或年入100万；女士三条件选二；1688元展示价','最接近我们的小型审核私宴'],
 ['HIMMR','985／211／QS100高知人群；当前线下活动与1V1统一报价未取得','保留高知活动与匹配供给参照'],
 ['我主良缘','白领、公务员及企业主等；官网派聚旅行社交品牌入口仍在','没有带价现售旅行产品，暂不定为旅行主对标'],
 ['珍爱网','线上及线下顾问服务；当前统一线下报价未取得','市场覆盖对象；不以App订阅代替门店价'],
 ['金凤凰','深港本地及再婚等细分；9月20日活动99元，长期红娘报价未公开','低价线下产品补充，非高资产主对标']],[3.6,8,5.2]);cite(['C01','BOL','M06','M10','M08','M11','M12'])
p('行业中至少有两笔不同的钱：参加一场活动的钱，以及顾问持续帮忙筛选、推荐、复盘的钱。我们的初创产品更接近前者；要进入高价定制，需要另有顾问能力、候选资源和服务合同。')

page('4 婚恋市场规模和竞争格局')
picture('婚恋市场五年.png')
p('根据网经社电子商务研究中心发布的《2025年度中国数字生活市场数据报告》，2025年互联网婚恋市场规模估计为103.5亿元，同比增长5.61%。其同口径序列由2021年的75亿元升至2025年的103.5亿元，年度增速下降。')
p('这一数字提供互联网婚恋的市场背景；本项目的线下高资产私宴、顾问和旅行服务仍需按具体产品评估。商业社群与线下高端婚恋没有取得可核的统一收入统计，报告不填一个推算总额。')
table(['要回答的问题','本轮查到什么'],[
 ['谁是行业头部','网经社2025百强榜把复爱合缘和珍爱网列为婚恋交友第一梯队，伊对、陌陌、探探、Soul列为第二梯队。这是多指标榜单分组。'],
 ['前五占多少','原始发布页没有给五家公司同口径收入和行业收入分母，不能据此计算CR5。榜单也不是份额排行。'],
 ['对线下创业的意义','大众入口有品牌与客源积累；小团队更应细分同城客群，把审核、同行者条件和一次线下交付讲清楚。我们是否能卖出高价，要回到客户与具体产品。']],[4.1,12.7]);cite(['E01','E02'])

page('4 客户画像怎样写到可用的程度')
p('把品牌定位、单场准入和公开个案分开后，客户差异很清楚。以下样例均由机构公开介绍，不作为其整体客户统计。')
table(['产品或样例','年龄 工作 收入条件','性格与兴趣的原始信息'],[
 ['Selected本次女性','95后为选项之一；海归或一本为另一项；家庭2000万或年入100万为第三项，选二','需面试；没有兴趣和性格分布'],
 ['HIMMR 10月21日回国匹配','在美国科技／AI公司工作的华人男性；借回国假期见面；收入仅称稳定','需先确认双方未来是否接受海外生活；未列统一年龄、薪资和兴趣'],
 ['HIMMR 8月24日公开男嘉宾','清华本科、UCLA硕士、UIUC博士；互联网AI算法岗；年龄及收入未列','机构介绍生活规律、踏实；这是单一样例'],
 ['金凤凰 8月22日公开女性','1995年生，深圳事业单位，211本科、硕士；收入未列','介绍温婉重家庭；爬山、短途游、羽毛球、瑜伽、舞蹈、游泳及象棋'],
 ['金凤凰 8月20日公开女性','1996年生；职业收入未列；香港住房为父母名下','介绍性格温和，喜欢音乐、旅游、美食、摄影']],[4.3,7.1,5.4]);cite(['BOL','M06','PX03','PX05'])
for row in D['old_profiles']:
 if any(k in row[0] for k in ['2026-10-21','2026-08-24','2026-08-22','2026-08-20']):
  q=d.add_paragraph();link(q,'公开样例原文 '+row[0].replace('\n',' '),row[5]);q.paragraph_format.space_after=Pt(2)
p('这些差异决定能否一起参加：长期在海外工作的人需讨论居住选择；再婚者需考虑家庭安排；年轻家族接班人与职业专业人士的财富口径也可能不同。活动招募应按实际关系目标与时间安排筛选，不能仅写“高端单身”。')

page('4 资产条件和参加目的怎样同时使用')
p('对本项目，A8及以上是既定目标，仍需明确个人还是家庭资产。对标的财富门槛不能直接当作我们的客户资产：Selected本场男士家庭5000万元或年入100万元，女性的资产条件只是三选二中的一项；钻石也同时服务千万级与亿元级客户。借鉴重点是审核和服务流程。');cite(['BOLSCREEN','C01'])
table(['拟服务的人','客群识别重点','适合的活动与购买理由'],[
 ['有明确恋爱目标的单身客户','本人真实单身、资产口径、年龄阶段、所在城市、居住计划与稳定关系意愿','审核私宴与单身旅行。第一次付款主要为同行条件、认识机会及被照顾的交流节奏。'],
 ['已婚、恋爱中或无婚恋意向的经营者','实际经营身份、正在面对的问题、可分享的经验及参与时间','主题小桌、分享会、运动或参访。付款理由是具体议题与有共同处境的同行。'],
 ['有职业专长的高资产人士','本人职业、时间安排及希望交流的主题，与家庭资产来源分别了解','可参加跨行业主题活动。由其专业与议题决定同桌，而不默认职业头衔意味着收入或资产。'],
 ['愿意共同出行的会员及伴侣','出行日期、独住或同住、活动强度、饮食与独处偏好、是否接受陌生同行','普通社交旅行或伴侣同行产品；单身互选另成专场，避免活动目的混淆。']],[4.2,6.1,6.5])
p('上表是根据已核案例提出的客户服务分组，不是本公司已有购买者统计。性格询问落实到交流偏好，爱好询问落实到过去实际投入；收入、个人资产和家庭资产分别记录，不能互相换算。')

page('5 社交旅行的市场与收费')
p('旅游补充检索已覆盖10组关键词、395条搜索记录、去重340篇；完整正文另核6条。比较结果包含普通同行、女性小团、单身主题和独行者旅行。真正要判断的是：谁负责筛人，谁负责地面交付，谁为两者收费。')
table(['产品','收费与主要含项','它与我们有什么关系'],[
 ['奇语济州岛4天3晚','“1k+”地面起价；四钻酒店、地面车及指定门票；机票另付','小团队与内容运营对标；该产品不限年龄性别'],
 ['OWR墨西哥城4天3晚','1400英镑；四星双住、部分餐、接送及活动；独住+225，机票另付','30—49岁独自报名者；学行前组织与同行体验'],
 ['心动计划5天4晚','正文6280、封面7280元起；8人、双住、4早2正；主持兼管家','单身主题和互动设计可学，报价差异及成团未核'],
 ['星七金海湖一日','198元；交通、领队、互动及保险，餐饮自理','低成本单身外出产品参照'],
 ['墨鱼App订阅','黄金23元／月、蓝钻38元／月','属于平台功能订阅，旅行团费另收'],
 ['32号及游侠客','既有具体线路、含项和退改资料继续保留','旅行地面交付参照，非单凭旅行社身份入选综合社群']],[4.2,7,5.6]);cite(['QI','OWR','HEART','STAR','MOYU','TD17','TD19'])
p('婚恋与旅游结合的证据：心动计划有单身筛选和完整五日设计；如爱、爱旅有低价单身出游招募。钻石山野局与我主良缘派聚是延伸线索，当前尚没有一份能够同时核清团费、天数、承办关系和经营实绩的高资产旅行档案。')

page('5 旅行公司怎样组织不同人群')
p('国内外供给可以按谁来组织同行者区分。下表继续复用已有如爱、心动计划、奇语及旅行社资料；新补官网产品和条款，不因新增来源改掉原主对标。')
table(['主导方式','代表与客群','收费与交付重点'],[[r[0],r[1]+'。'+r[2],r[3]+' '+r[4]]for r in D['travel_forms']],[3.2,6.5,7.1]);cite(['R09','R10','TD08','HEART','R12','R06','OWR'])
p('我们的判断：旅行承办方可以买到车、房、餐和目的地活动；社群还要承担同行筛选、行前沟通和现场相处。定价应把这两部分各自交付写清，才能讨论客户愿意多付的钱对应什么。',True)

page('5 一次同行可以容纳哪些关系状态')
picture('旅行者结构.png')
p('FTLO面向25—39岁，官网自述平均30岁，多数27—34岁，女性占70%—80%。其三项不同指标见图：独自报名80%，自述单身且愿交友60%，大城市居住75%。官网没有列对应样本量和年份，职业主要以年轻专业人士概括。');cite(['R06'])
p('这为我们的产品提供两个具体选择。普通社交旅行可服务独自报名、有伴侣但单独出行，或与伴侣同来的人，报名时说明共同兴趣和活动强度；以恋爱为目的的旅行则另行确认真实单身、关系意愿及相处安排。它们可以共用旅行供应，但应分别招募。')
p('性格与爱好怎样用于组织：OWR要求对认识同行者有兴趣，并接受共同年龄阶段中的职业、文化与观点差异；FTLO允许客人自由活动后再一起用餐。我们可以询问喜欢多人讨论还是少量深聊、是否愿意运动、需要多少独处时间，据此安排活动。上述为产品设计建议，不编写客户性格分布。');cite(['R04','R06'])

page('5 城市社群延伸旅行的完整例子')
p('The Offline Club与法国Simply Canvas Farm合作，2027年5月20—25日住五晚。社群带来客人并安排一位社区主持人，农场由业主提供场地、餐食和日常照料。这个分工比单看一条酒店线路更接近我们想研究的社群旅行。')
table(['五月场房型','每人全价','包含及付费方式'],[
 ['自带帐篷或房车','895欧元','使用场地及已列活动、餐食；自带住宿装备'],
 ['双人帐篷／指定双住','995欧元','同性双床按报名类型安排；每人单独买名额'],
 ['房屋内独立房间','1095欧元','一个人住，参加共同餐食及活动'],
 ['独立小屋','独住1125欧元；伴侣双住每人995','伴侣双住合计1990欧元；均先付每人300欧元定金']],[4.9,5,6.9])
p('售票页按钮显示的300欧元是定金，具体全价写在各房型说明中。上表采用房型票明示金额；活动正文还留有895—1195欧元总区间。往返交通、部分外出交通另计，特定饮食加价见主办方规则。');cite(['R12','R07'])
p('谁来参加：设计为16—20人，往期年龄23—68岁，男女及不同关系状态均可。官网展示33岁设计师、27岁创业者、47岁教师的体验样例。其共同点是愿意暂离手机、休息并自愿交流，未设置资产或职业级别门槛。')
p('怎样相处：入场收存手机，白天可散步、冥想、创作或独处，晚餐和篝火提供自然相聚；没有强制满排的破冰节目。它适合借鉴“有共同节奏、也有自由空间”的形式，不能据此照搬为高资产婚恋产品。')

page('5 旅行付款和退出怎样安排')
table(['产品','付款','旅客取消规则摘要'],[
 ['OWR','英镑定金200／美元250；出发前70天付清，不足70天报名全款','超过70天保留定金，可转后续旅行；69—30天扣50%；不足30天扣100%。'],
 ['Offline法国五月','每人先付300欧元，尾款由主办方另发通知，可分期','提前超过8周扣100欧元取消费；4—8周退50%；不足4周不退。'],
 ['32号海南六日','按选定日期、人数、房型确认；付费报名，一单一团2人起','前11日及以上0%；10—7日20%；6—4日40%；3—1日60%；当日70%。不足实际预订损失时按实际损失，人数变化重核费用。'],
 ['FTLO日本','线路定金；常规分期前95／65／35天累计付50%／75%／100%','日本适用95天档：95天以上定金转旅行额度；94—65天25%不可退；64—35天50%不可退；34天内不退。另有供应商损失条款。']],[3,6.1,7.7]);cite(['R05','R12','R09','R06'])
p('研究判断：客户收款节奏要与酒店、车辆及承办付款对应，先确认成团再锁不可退资源。自身合同应逐日写清取消边界；例如OWR网页未单列恰70天，部分退改还叠加供应商实际损失，不能只摘一个“可退比例”当成全部退款承诺。',True)

page('5 商业与婚恋扩搜新增的机构线索')
X=D.get('discovery',{});reviews=X.get('reviews',[])
def metric(v):return str(v)if v not in [None,'']else'未返回'
def social_label(r):return r['name']+'\n'+r['date']+'\n粉丝'+str(r['fans']if r['fans']is not None else r.get('fans_status','未采集'))+'\n赞'+metric(r['likes'])+' 评'+metric(r['comments'])+'\n分享'+metric(r['shares'])
p('本次追加四组泛词检索，分别按热度和最新排序，共返回'+str(X.get('cards',0))+'条卡片，去重'+str(X.get('unique_notes',0))+'篇；进一步读取'+str(X.get('complete_notes',0))+'篇正文。搜索范围包括商业社群、企业家饭局、高端单身活动与验资单身旅行；互动和粉丝为10月24日快照。')
commercial=[r for r in reviews if r['group']=='商业社交']
table(['账号 发布日期与互动','客户与服务','本次取舍'],[[social_label(r),r['people']+'。'+r['flow'],r['judgment']]for r in commercial],[3.6,8.2,5.0]);cite([r['source']for r in commercial]+[r['profile_source']for r in commercial if r.get('profile_source')])
p('Pitch值得补入活动形式库：围绕具体生意问题组织分享，再连接企业走访和同城聚会，学习主题比泛泛的“高端饭局”更具体。但这篇为2025年资料，当前收费与2026年履约仍未补齐，先保留为形式参照。')

page('5 婚恋新增线索与用户体验')
marriage=[r for r in reviews if r['group']=='婚恋交友']
table(['账号 发布日期与互动','客户和组织方式','本次取舍'],[[social_label(r),r['people']+'。'+r['flow'],r['judgment']]for r in marriage],[3.6,8.2,5.0]);cite([r['source']for r in marriage]+[r['profile_source']for r in marriage if r.get('profile_source')])
feedback=[r for r in reviews if r['group']=='用户体验']
p('体验反馈可用来检查主持与轮换：一位广州参与者一直与朋友玩游戏，未认识异性；另一篇深圳受邀体验提到晚餐开始安静、后段才自由交流。这两篇不用于统计机构的经营表现。');cite([r['source']for r in feedback])
p('新线索先进入名录与产品表。BoL及Selected继续作为综合主对标；钻石用于顾问服务，Founder Feast用于小桌组织。')

page('第二部分 对标',1);h('6 为什么选择这五家机构')
table(['判断维度','怎样使用'],[
 ['与业务的关系','是否持续组织同一类人线下相处；单身、商务和普通同行分别看'],
 ['启动可学性','能否先靠筛选、组织、主持和合作场地实行；不依赖现成大IP'],
 ['收费证据','至少明确一个产品的金额、单位及主要含项；模糊起价保留原文'],
 ['客户精度','能否把准入写到年龄、职业、资产／收入或关系目的'],
 ['交付证据','有公开活动流程、实际往期记录或具体团期，而非只有品牌介绍'],
 ['经营表现','会员、到场、续费及利润逐项判断；没有实绩不评为经营成功']],[4,12.8])
table(['方向','主对标与分工'],[['商业','BoL：本地会籍与活动分层；Founder Feast：小桌匹配、订阅与单次酒会'],['婚恋','BoL Selected：半天审核私宴；钻石：成熟高资产顾问服务及资源层级'],['旅行','奇语：企业小团与内容运营；OWR：独行年龄带、行前服务与收费节点']],[3,13.8])
p('备选处理：HIMMR保留高知交友市场参照，金凤凰保留具体低价活动；Inner Table Club与Founder Feast同类，前者暂缺同等具体活动与应用记录；长江、EO、马会和木兰仍放在成熟参照中。大超的IP获客、正和岛及优家荟不进入本版启动建议。')

for index,c in enumerate(D['cases'],7):
 page(str(index)+' '+c['name'])
 p(c['what']);p('研究用途：'+c['role'],True)
 h('谁会参加',3);table(['画像字段','能够核实的内容'],c['people'],[3.7,13.1])
 h('怎么收费和买到什么',3);table(['产品','费用','包含及另付'],c['fees'],[4,4.7,8.1])
 cite(c['sources'])
 if c['id']=='BOLB' and D.get('bol_membership_rights'):
  page('BoL两档会籍的具体权益')
  p('星光488元／年，星辰2888元／年。两档都列全年10+场酒会免费入场；差别集中在沙龙、运动及星辰的额外服务。以下按用户提供的会籍页逐项展开。')
  table(['权益','星光488元／年','星辰2888元／年','收费与使用说明'],D['bol_membership_rights'],[2.8,3.4,4.9,5.7])
  p('会籍页用每次金额乘12展示价值，同时在权益说明中写全年10+场；二者不能改写成已经交付12场。1288元位于私宴价值栏，会员实付票价未公开。')
  cite(['BOLSCREEN'])
 page(str(index)+' '+c['name']+'怎样交付和经营')
 h('活动与服务流程',3)
 for i,t in enumerate(c['flow'],1):p(str(i)+'  '+t)
 h('客户从哪里来',3);p(c['acquisition'])
 h('经营表现',3);p(c['performance'])
 h('我们的判断',3);p(c['judgment'])
 cite(c['sources'])

page('13 对标之后可以作出的判断')
p('结论是：这个项目值得以小规模、可交付的线下活动验证。现有供给已经证明商业餐桌、筛选私宴和陌生同行有付费形式；我们的机会在于持续服务同一批适配客户，并把同行条件与服务内容说明白。')
for row in D['conclusions'][:3]:
 h(row[0]+'：'+row[1],3);p('依据：'+row[2]);p('建议：'+row[3])
cite(['BOLSCREEN','R01','R03','R06','R12'])

page('13 首阶段向谁学，先做什么')
for row in D['conclusions'][3:]:
 h(row[0]+'：'+row[1],3);p('依据：'+row[2]);p('建议：'+row[3])
p('建议的启动顺序是先做好一类同城小活动，形成可再次邀请的人，再确定第一条旅行的同行构成。具体采用单身私宴还是经营者小桌，取决于两位合伙人首先能稳定触达哪类客户。商业与婚恋分别报名、共同兴趣可以跨活动延续。')
p('改变建议的条件：若持续筛选后仍无法形成彼此愿意再见的客群，应收窄主题或人群；若客户只愿买景点和酒店服务，可采用旅行合作产品；若有人明确购买长期顾问服务，再评估顾问团队、候选资源和合同。')
cite(['BOLSCREEN','C01','R01','R04','R09','R12'])

page('13 这些机构经营到了哪一步')
table(['机构与现有事实','能够作出的判断','对本项目的含义'],[
 ['BoL：两档年费、单独售卖Selected、完整准入与私宴流程；10万会员为品牌自述','产品组合已经清晰；没有真实付费会员和续费资料支持“会籍经营成功”的结论','适合优先学产品组合与筛选，初期不照搬全年高频免费承诺'],
 ['Founder Feast：当前三城价格及隔周晚餐，App预约与聊天，往期照片、预售酒会','有具体交付和产品迭代，尚缺付费人数、留存与利润','小桌组织容易拆解，规模与获客效果要另验证'],
 ['钻石：30／60／100万元起定制产品，当期私宴及山野招募','高价服务的专业分工较完整，公开标价不能证明成交数量或收益','借鉴筛选、测评及辅导；启动阶段不复制其资源库投入'],
 ['奇语：抖音8143粉丝、89件作品，10月21日阿勒泰招募，分性别主题小团','可确认持续发内容和在售招募，未核实单团实到、复购和贡献','学同伴主题与内容展示，酒店和高资产审核需自行设计'],
 ['OWR：具体日期、可付款商品、酒店、行前及退改条款','产品交易和交付规则较成熟；利润与复购未公开','最适合学旅行预期、付款及退改；A8服务标准另定']],[6.2,5.3,5.3]);cite(['BOLSCREEN','BOL','R01','R03','B14','C01','QI','OWR','R05'])
p('因此本报告不排“最好公司”。综合优先学BoL，轻量组织再看Founder Feast，婚恋专业能力看钻石，旅行交付看OWR；奇语与Offline Club各提供可拆解的内容和社群旅行形式。这是按学习用途分配优先级，依据是具体产品与可复制工作。')

page('13 客户第一次究竟付哪几笔钱')
table(['产品','首笔或总价','已包含与另付'],[[r[0],r[1],r[2]+'；另付：'+r[3]]for r in D['first_payment']],[3.5,5,8.3]);cite(['BOLSCREEN','R03','R11','OWR','R05','R12'])
p('比较时先看自己的使用方式：只想参加一次晚宴的人，未必适合购买全年多场权益；想独住的人，要看单房差后的总价；希望商谈具体问题的人，会在意组桌结果。年费、订阅和团费各有成立条件，现阶段不根据几家标价直接给我们的产品定价。')
p('我们的两个本地活动与一条旅行仍放在配套模型中，1680／680／19800元用于讨论成本和人数变化。建议把“符合准入的人是否愿意在具体日期为这些含项付款”作为下一步定价验证。')

page('13 对首轮收费测试的具体建议')
p('建议先测试一次同城私宴，保留1680元作为内部测试价。它与Selected该场1688元展示价接近，且能用既有餐食、组织和获客预算计算不同人数结果。收费依据应写成具体同行条件和交付，不因A8标签直接加价。',True)
Fdata=json.loads((B.parent/'execution-followthrough-20261024/执行依据与测算数据.json').read_text())
table(['初案及研究试价','基准情景','本轮建议'],[
 ['单身晚宴\n1680元／人','20人收入33600元；项目贡献9492元，团队合计70小时','作为首轮报价测试。以一席餐食、行前沟通、主持和自愿后续联系解释价格。'],
 ['经营者分享会\n680元／人','24人收入16320元；贡献1590.4元，16人时亏1606.4元；团队52小时','先确认足够具体的议题与分享者，优先合作场地；当前固定成本下盈利空间薄，不靠多办场掩盖低贡献。'],
 ['三亚四天三晚\n19800元／人','14人收入277200元；贡献35104元，贡献率12.7%；团队190小时，周转需求141296元','作为后续产品测算，先有合格同行者与承办报价再启动。机票另付，三晚独住和社交组织写进含项。']],[4,6.5,6.3])
p('这些是配套模型的情景结果，已计人员项目工时、获客预算和3%退款预留；项目贡献尚未扣公司年度管理费用及相关税费。酒店采用11月19—22日单间三晚7870元零售基准，其他费用仍有预算假设。')
p('三亚旅行并没有因为每人收19800元就形成很厚的利润：在14人基准下，售价下调10%会使贡献由35104元降至8215.6元。研究判断是先证明同行体验值得购买，再谈折扣和会员优惠。')
p('同线路比较的当前结论：官网与携程可核到同酒店三晚价，但旅行社的2748或5980元起价对应不同天数、酒店及同住方式。现有证据不足以计算相同规格旅行的社交溢价。供应商正式询价继续列外部验证，不阻塞本版公开研究交付。')
cite(['BOLSCREEN','R09','R10','MARBB','CTRBB'])

page('附录 前轮重点候选的取舍')
table(['前轮候选','继续保留的证据','本轮用途与取舍'],[
 ['Hampton','15000美元／年；按营收、融资或出售企业条件筛选；同城固定小组','保留商业机制参照。Sam Parr播客及已有网络是获客条件，不能和我们等量；小规模餐桌优先拆Founder Feast。'],
 ['Flash Pack','按30—49及50—65岁分组；问卷、提前4周同行群及领队组织','保留精品陌生同行参照。未取得动态指定团期价，收费比较以能读到完整金额的OWR和FTLO补足。'],
 ['The Offline Club','法国2027产品895—1125欧元，按房型；16—20人；含厨师全餐','重点保留城市社群延伸度假的形式。客群未设高资产或婚恋条件；已核5月20—25日、五晚及具体房型票；社群主持人与农场合作交付。'],
 ['FTLO Travel','日本九日6899美元起，含8晚8早3午4晚；独住另加','保留同龄精品旅行与报价参照。国际机票、机场接送及保险另付；起价需与具体日期配套。']],[3.5,6.3,7]);cite(['HAM','FLASH','OFF','FTLO'])
p('这四家没有因为新增候选而被删除。主对标用于优先拆解当前可实行的产品；成熟或相邻模式仍负责解释市场、客户预期和其他活动形式。')

page('附录 既有国外案例继续怎样使用')
p('前轮九家国外会员组织继续保留，用来研究成熟场景、会员权益和活动形式。它们不挤占小团队运营对标，也不因新发现Founder Feast而被删除。')
for j,c in enumerate(D['legacy_foreign']):
 if j==5:page('附录 既有国外案例继续怎样使用续')
 h(c['name'],3)
 text=c['what']+' '+c['fees']
 if c['name']=='The Supper Club':text='城市会员餐叙与活动组织。当前Quattro为2500美元／年另750入会，限同州4场，不开放And Beyond旅行；Explorer5000美元／年另750入会，可跨城市8场。适合学习次数、地域和旅行准入分层。'
 p(text)
 cite(['T'+x for x in c['sources']])

if any(q['id']=='CTR-BB'for q in D.get('cost_quotes',[])):
 page('附录 三亚同日住宿报价补核')
 p('比较条件为2026年11月19日入住、22日退房，三晚、1间房、搜索1位成人。官网特大床青樾湾景观房与携程青樾湾景观大床房均为73平方米；两端名称相近，但赠送项目与取消条件不同。')
 table(['报价','三晚总价及含项','取消与付款'],[
  ['官网普通含早套餐','7870元，含税；双早及皮划艇、珊瑚船、旅拍等套餐项目','10月25日23:59后取消预计扣6750元；显示预付款6750元，税费付款时点未单列'],
  ['官网预付含早','7240元，含税；早餐份数及全部赠送未展开','10月25日23:59前可免费取消，之后预付款不退；预订不可更改'],
  ['携程已登录会话含早','7240.86元，三晚页面应付；两份早餐，含皮划艇／桨板、潜水体验及旅拍等','10月24日23:59前免费取消，此后不可取消；页面在线支付7240.86元，未另列税费']],[3.8,6.1,6.9])
 cite(['MARBB','CTRBB'])
 p('携程列表每晚显示2414元，预订详情精确到2413.62元，三晚合计7240.86元。它比官网普通含早套餐低629.14元，与官网预付含早只差0.86元；不能据此判断携程在所有相同条件下更便宜。')
 p('费用模型仍采用7870元住宿基准，另一档报价留作可修改情景。先确认整团房量、套餐交付及退改，再替换成本。1间零售价没有证明14间能按同价取得；完整旅行还需用车、餐饮、承办及社交组织成本。')

page('附录 采集覆盖与证据使用')
table(['工作','本轮及既有成果','使用方式'],[
 ['机构地图',str(len(D['market']))+'条；复用原36家及既有国外案例，加入补查品牌线索','同品牌不因会籍或城市重复计数；名录不等于全量深访'],
 ['商业婚恋网页','本轮'+str(len(D['ledger']))+'个网页任务，'+str(sum(bool(x.get('success'))for x in D['ledger']))+'次有效正文','Crawl4AI保存原文、时间与哈希；失败任务保留，钻石定制页本轮已补采成功'],
 ['商业婚恋小红书','4词、80条搜索卡片；完整正文按返回情况留存','万象博联、钻石婚恋、HIMMR、木兰汇；本轮最热排序不是随机抽样'],
 ['追加泛词扩搜',str(D.get('discovery',{}).get('cards',0))+'卡片，'+str(D.get('discovery',{}).get('unique_notes',0))+'去重内容，'+str(D.get('discovery',{}).get('complete_notes',0))+'正文','热度与最新分别保留；新增账号的互动不当成成交'],
 ['旅游小红书','10词、395卡片、去重340篇，完整正文6条','围绕单身旅行、旅行社交和小团；已完成用户登录与验证'],
 ['前轮婚恋','73条帖子记录、5个账号快照及57条价格记录','大超留在历史资料；本版只使用与研究目的相关的有效资料'],
 ['BoL微信','10月24日会籍和Selected逐页读取','价格、准入、流程和商品计数均有截图；没有网页直链'],
 ['总Excel','两张阅读主表，三张追溯附件','市场调研、对标产品、机构名录、来源与缺口、内容互动']],[3.5,6.9,6.4])
p('真实性控制：金额保留币种、计费单位、适用活动与采集日期；行业统计优先官方，客群调查保留样本方法；品牌自述、页面销售展示、招募计划和实际活动分别记录。空白指标不填零，历史热帖不参与近期活跃判断。',True)
p('链接使用：Word与PDF内资料名均可点击。BoL为微信小程序，资料入口写操作路径并在证据包附截图；小红书和抖音可能要求登录。文件中超链接目标会逐项核验，网页采集成功不能保证平台未来不要求登录。',True)
p('尚需外部核实的事项：部分机构真实成交价、合同中的完整退改、实际付费客户分布、成团与复购、成本利润，以及我们自己的客户和资源验证。旧版1680／680／19800元仅是研究试价，本版不将其写成最终产品售价。')

# Source index contains only sources actually cited, keeping the inherited full ledger in Excel.
page('附录 资料索引')
for i,k in enumerate(used):
 if i and i%11==0:page('附录 资料索引续')
 s=S[k];q=d.add_paragraph()
 if s['url']:link(q,k+' '+s['name'],s['url'])
 else:q.add_run(k+' '+s['name']+'  微信搜索BoL万象博联')
 period=re.sub(r'(20\d\d-\d\d-\d\d)T(\d\d:\d\d):[0-9.]+[+-]\d\d:\d\d',r'\1 \2采集',s['period'])
 q2=p(period+'；'+s['kind']+'。'+s['note'])
 q2.paragraph_format.line_spacing=Pt(13.5)
 for r in q2.runs:r.font.size=Pt(9)

# Add internal bookmarks and linked reading directory.
heading_map={q.text:q for q in d.paragraphs if q.style.name.startswith('Heading')}
aliases={'3 商业社交供给与收费':'3 商业社交有哪些供给','4 婚恋市场与客群比较':'4 婚恋市场要分清哪些产品','5 旅行市场与产品比较':'5 社交旅行的市场与收费','6 选择依据与名单':'6 为什么选择这五家机构','7 BoL商务会籍':'7 BoL万象博联','13 综合判断':'13 对标之后可以作出的判断','附录 既有成果与资料入口':'附录 既有国外案例继续怎样使用'}
for i,(q,label)in enumerate(toc,1):
 target=heading_map.get(aliases.get(label,label))
 if target is None:continue
 anchor='section_'+str(i);a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(i));a.set(qn('w:name'),anchor);target._p.insert(0,a);z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(i));target._p.append(z)
 q.clear();link(q,label,None,anchor,size=11.5)
out=O/'高端线下社群市场调研与对标报告.docx';d.save(out)
(B/'报告引用来源.json').write_text(json.dumps([S[k]for k in used],ensure_ascii=False,indent=2))
print('Saved',out,'cited',len(used))
