from pathlib import Path
from copy import deepcopy
import json,re,shutil
from docx import Document
from docx.oxml.ns import qn
B=Path(__file__).resolve().parent;O=B/'交付文件';H=B/'历史版本/删除两家前';H.mkdir(parents=True,exist_ok=True)
for f in O.glob('*'):
 if f.suffix in ['.docx','.pdf','.xlsx']:shutil.copy2(f,H/f.name)
bad=['正和岛','优家荟'];has=lambda s:any(x in s for x in bad)
def txt(el):return ''.join(el.itertext()) if False else ''.join(n.text or '' for n in el.iter(qn('w:t')))
report=O/'中国高端会员社群市场研究及创业可行性评估报告.docx';d=Document(report)
# Restore the full-market taxonomy, independently from the selected deep cases.
old=Document(B.parent/'social-research-20261021/交付资料/历史版本/20261021'/report.name)
classification=next(t for t in old.tables if [c.text for c in t.rows[0].cells]==['机构类型','代表对象','研究重点'])
if not any([c.text for c in t.rows[0].cells]==['机构类型','代表对象','研究重点'] for t in d.tables):
 p=next(q for q in d.paragraphs if q.text.startswith('机构按主要组织方式归类'))
 p._p.addnext(deepcopy(classification._tbl))
# Remove the two entire case sections, retaining all other evidence and sources.
a=next(p._p for p in d.paragraphs if p.text=='4.4 正和岛');z=next(p._p for p in d.paragraphs if p.text=='4.6 我主良缘');active=False
for el in list(d._element.body):
 if el is z:break
 if el is a:active=True
 if active:d._element.body.remove(el)
replacements={
'八家对标分别承担具体研究任务':'六家对标分别承担具体研究任务：北京马会研究家庭会籍和日常使用，长江商学院研究共同学习与校友关系，木兰汇研究女性企业经营者活动；我主良缘、金凤凰和珍爱网用于研究专业匹配、活动收费与规模服务管理。EO保留为私董会专题。',
'木兰汇与优家荟分别':'木兰汇按企业家组织归类，女性是客户标签。旅行社及供应商另列合作资源；尚未核实主体和近期活动的账号保留为线索。完整机构地图和原始记录见Excel。',
'保留五家会员、学习与生活方式组织':'本报告重点研究北京香港马会会所、长江商学院、木兰汇，以及我主良缘、金凤凰深港婚恋和珍爱网。前三家分别提供会籍、共同学习和企业经营者活动的组织样本；后三家用于拆解专业婚恋的产品、收费与交付。',
'这八家属于不同用途':'这六家用于回答不同的经营问题。三家婚恋样本分别覆盖产品组合、区域活动和规模服务转型，具体适用客群与经营表现见逐家分析。',
'本项目小红书长江23篇':'本项目小红书长江23篇样本中，评论字段可见4篇合计5条，分享字段可见8篇合计33次。代表帖发表于2026年9月14日，记录2条评论、10次分享。采集日期为2026年10月21日。',
}
for p in list(d.paragraphs):
 for start,new in replacements.items():
  if p.text.startswith(start):p.text=new;break
 # Mixed source paragraphs retain the unrelated linked sources.
 if p.text.startswith(('依据  ','资料  ')):
  for e in list(p._p):
   if e.tag==qn('w:hyperlink') and has(txt(e)):p._p.remove(e)
 elif has(p.text):p._p.getparent().remove(p._p)
# Remove dedicated comparison rows, preserve a mixed row's remaining valid example.
for t in d.tables:
 for row in list(t.rows):
  if not has(' '.join(c.text for c in row.cells)):continue
  if row.cells[0].text=='成员提供内容与接待':row.cells[1].text='木兰汇';continue
  # Full-market taxonomy is a classification, not a recommendation.
  if [c.text for c in t.rows[0].cells]==['机构类型','代表对象','研究重点']:
   for c in row.cells:
    if has(c.text):c.text=c.text.replace('正和岛、','').replace('、正和岛','').replace('优家荟、','').replace('、优家荟','')
  else:row._tr.getparent().remove(row._tr)
renumber={'4.6 我主良缘':'4.4 我主良缘','4.7 金凤凰深港婚恋':'4.5 金凤凰深港婚恋','4.8 珍爱网':'4.6 珍爱网'}
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for n in p._p.iter(qn('w:t')):
     if n.text:
      for a,b in renumber.items():n.text=n.text.replace(a,b)
      n.text=n.text.replace('八家','六家')
for p in d.paragraphs:
 for n in p._p.iter(qn('w:t')):
  if n.text:
   for a,b in renumber.items():n.text=n.text.replace(a,b)
   n.text=n.text.replace('八家','六家')
d.core_properties.subject='六家对标及婚恋旅游专题更新 2026年10月22日';d.save(report)
structure=json.loads((B/'报告结构.json').read_text());structure['core_names']=[s for s in structure['core_names'] if not has(s)];structure['selection']=[r for r in structure['selection'] if not has(str(r))];structure['nav']=[(renumber.get(a,a).replace('八家','六家'),b) for a,b in structure['nav'] if not has(a)];(B/'报告结构.json').write_text(json.dumps(structure,ensure_ascii=False,indent=2))
# Extract the revised chapter using the same document relationships.
c=Document(report);active=False
for el in list(c._element.body):
 if txt(el)=='四 六家机构深度研究':active=True
 if txt(el)=='五 横向比较与公司机会':active=False
 if not active and el.tag!=qn('w:sectPr'):el.getparent().remove(el)
c.paragraphs[0].text='六家会员社群与婚恋机构对标案例档案';c.paragraphs[0].style='Title';c.paragraphs[0].paragraph_format.page_break_before=False;c.core_properties.title=c.paragraphs[0].text
c.save(O/'六家会员社群与婚恋机构对标案例档案.docx')
for x in O.glob('八家*'):x.unlink()
s=Document(O/'高端会员社群创业项目合伙人决策摘要.docx')
for p in s.paragraphs:
 if p.text.startswith('主研究样本调整为八家'):p.text='重点研究六家：北京香港马会会所、长江商学院、木兰汇、我主良缘、金凤凰深港婚恋和珍爱网。前三家研究会员组织与持续活动，后三家研究专业婚恋的产品、收费、匹配和履约。EO保留为固定小组专题。'
 else:
  for r in p.runs:r.text=r.text.replace('八家','六家')
for t in s.tables:
 for row in list(t.rows):
  if has(' '.join(c.text for c in row.cells)):row._tr.getparent().remove(row._tr)
s.save(O/'高端会员社群创业项目合伙人决策摘要.docx')
for file in O.glob('*.docx'):
 assert not has(txt(Document(file)._element.body)),file
# Prevent the previous build recipe from silently recreating the deleted selection.
p=B/'build_documents.py';z=p.read_text();z=z.replace("SOURCE=OLD/'中国高端会员社群市场研究及创业可行性评估报告.docx'","SOURCE=BASE.parent/'social-research-20261021/交付资料/历史版本/20261021/中国高端会员社群市场研究及创业可行性评估报告.docx'")
start=z.index('# Remove the duplicate general categorisation table;');end=z.index('# Remove EO from the main case chapter;',start);z=z[:start]+'# Preserve the full-market classification table alongside the case selection table.\n'+z[end:];p.write_text(z)
print('Restored taxonomy; six-case report, archive, summary saved.')
