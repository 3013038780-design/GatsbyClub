from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from xml.etree import ElementTree as E
import openpyxl,json
B=Path(__file__).resolve().parent;p=B/'交付文件/中国高端会员社群市场研究数据工作簿_31家机构.xlsx';w=openpyxl.load_workbook(p)
updates={}
s=w['文件与数据说明']
for row in s:
 label=row[0].value
 counts={'市场机构地图':w['市场机构地图'].max_row-4,'平台内容明细':w['平台内容明细'].max_row-4,'原文提取片段':w['原文提取片段'].max_row-4,'网页采集尝试':w['网页采集台账'].max_row-4}
 if label in counts:updates[row[1].coordinate]=counts[label]
 if label=='本次30家扩展':updates[row[0].coordinate]='当前机构覆盖';updates[row[1].coordinate]=31;updates[row[2].coordinate]='主表及价格、权益、画像、获客、流程和经营表现覆盖31家；其中六家为重点案例。'
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';rn='http://schemas.openxmlformats.org/officeDocument/2006/relationships';T=lambda s:'{'+ns+'}'+s;E.register_namespace('',ns);E.register_namespace('r',rn)
with ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
wr=E.fromstring(parts['xl/workbook.xml']);rel={r.get('Id'):r.get('Target') for r in E.fromstring(parts['xl/_rels/workbook.xml.rels'])};sh=next(s for s in wr.find(T('sheets')) if s.get('name')=='文件与数据说明');sp=rel[sh.get('{'+rn+'}id')].lstrip('/');sp=sp if sp.startswith('xl/') else 'xl/'+sp;rt=E.fromstring(parts[sp])
for c in rt.iter(T('c')):
 if c.get('r') in updates:
  v=updates[c.get('r')]
  for n in list(c):c.remove(n)
  if isinstance(v,str):c.set('t','inlineStr');E.SubElement(E.SubElement(c,T('is')),T('t')).text=v
  else:c.attrib.pop('t',None);E.SubElement(c,T('v')).text=str(v)
parts[sp]=E.tostring(rt,encoding='utf-8',xml_declaration=True)
with ZipFile(p,'w',ZIP_DEFLATED) as z:
 for n,v in parts.items():z.writestr(n,v)
print(updates)
