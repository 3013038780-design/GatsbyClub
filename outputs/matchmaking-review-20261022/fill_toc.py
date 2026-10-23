from pathlib import Path
import json,re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader
B=Path(__file__).resolve().parent;P=B/'交付文件/中国高端会员社群市场研究及创业可行性评估报告.docx';pdf=next((B/'qa/report').glob('*.pdf'));pages=PdfReader(pdf).pages
norm=lambda t:re.sub(r'\s+','',t)
texts=[norm(p.extract_text()) for p in pages];d=Document(P);nav=json.loads((B/'报告结构.json').read_text())['nav'];table=next(t for t in d.tables if [c.text for c in t.rows[0].cells]==['内容','页码']);result={}
for row,(title,key) in zip(table.rows[1:],nav):
 matches=[i+1 for i,t in enumerate(texts) if i!=1 and norm(title) in t]
 if not matches:raise ValueError(title)
 n=matches[0];row.cells[1].text=str(n);row.cells[1].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT;result[title]=n
# Keep first-level and case headings linked, page numbers are stable physical PDF pages.
d.save(P);(B/'qa/目录页码.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
