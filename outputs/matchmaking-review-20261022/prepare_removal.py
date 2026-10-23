from pathlib import Path
import openpyxl,json
B=Path(__file__).resolve().parent;P=B/'交付文件/中国高端会员社群市场研究数据工作簿_33家机构.xlsx';w=openpyxl.load_workbook(P)
changes={};names=['正和岛','优家荟'];has=lambda v:isinstance(v,str) and any(n in v for n in names)
def replace(v):
 if not isinstance(v,str):return v
 return v.replace('33家','31家').replace('八家','六家').replace('五家会员／学习／生活方式组织及三家婚恋机构','北京马会、长江商学院、木兰汇及三家婚恋机构')
for s in w:
 remove=[r for r in range(5,s.max_row+1) if any(has(c.value) for c in s[r])]
 substitutions=[(c.coordinate,replace(c.value)) for row in s for c in row if replace(c.value)!=c.value]
 if not remove and not substitutions:continue
 assert not any(c.data_type=='f' for row in s for c in row),(s.title,'unexpected formula in compacted data sheet')
 keep=[r for r in range(5,s.max_row+1) if r not in remove]
 rows=[[replace(s.cell(r,c).value) for c in range(1,s.max_column+1)] for r in keep]
 # Correct scalar sample counts and summaries in the project guide.
 for row in rows:
  if row[0]=='六家独立案例':row[1]=6
  if row[0]=='机构样本' and row[1]==33:row[1]=31
 changes[s.title]={'keep':keep,'rows':rows,'oldRows':s.max_row,'cols':s.max_column,'top':[(c,v) for c,v in substitutions if int(''.join(filter(str.isdigit,c)))<5]}
plan={'changes':changes,'rename':{n:n.replace('33家','31家') for n in w.sheetnames if '33家' in n}}
(B/'删除两家工作簿计划.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2,default=str))
print('sheets',len(changes),'deleted rows',sum(x['oldRows']-4-len(x['rows']) for x in changes.values()))
