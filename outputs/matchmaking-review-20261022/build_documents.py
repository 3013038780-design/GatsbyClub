from pathlib import Path
from copy import deepcopy
import re,json,sys
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from findings import SOURCES,PRICES,CASES,VIP,TRAVEL,BASE
OLD=BASE.parent/'social-research-20261021/交付资料/正式文件'
OUT=BASE/'交付文件';OUT.mkdir(exist_ok=True)
SOURCE=BASE.parent/'social-research-20261021/交付资料/历史版本/20261021/中国高端会员社群市场研究及创业可行性评估报告.docx'
d=Document(SOURCE);anchor=None
NAMES=['北京香港马会会所','长江商学院','木兰汇','正和岛','优家荟','我主良缘','金凤凰深港婚恋','珍爱网']
ROWS=[
[NAMES[0],'综合会籍、家庭同行与另付消费','现行个人与公司收费、家庭卡及文化活动','家庭和日常使用'],
[NAMES[1],'课程、共同学习、参访与校友','现行学费、2026参访与返校记录','学习和持续同伴'],
[NAMES[2],'女性企业家、年会与会员分享','2026分档票价、50余人参访记录','女性经营者'],
[NAMES[3],'商业内容、会员层次与地方组织','数字产品价、2026参访及大会记录','企业家内容获客'],
[NAMES[4],'女性生活方式与跨城共同体验','历史分层会籍、跨城项目与规模自述','生活体验和旅行'],
[NAMES[5],'匹配、情感、形象、活动与旅行品牌','线上商品价格、合同个案及派聚说明','专业婚恋产品组合'],
[NAMES[6],'国内与深港分层匹配、同城活动','VIP权益、2026活动价及报名、历史景区项目','区域婚恋及活动收费'],
[NAMES[7],'规模获客后的顾问服务与经营转型','五步服务、合同个案、外包调整和监管记录','转型及履约管理'],
]
def text(el):return ''.join(t.text or '' for t in el.iter(qn('w:t')))
def find(t):return next(e for e in d._element.body if text(e)==t)
def remove_range(start,end):
 a=find(start);b=find(end);active=False
 for el in list(d._element.body):
  if el is b:break
  if el is a:active=True
  if active:d._element.body.remove(el)
def place(el):
 if anchor is not None:anchor.addprevious(el)
 return el
def p(t='',style=None):
 q=d.add_paragraph(t,style);place(q._p);return q
def h(t,level=2,new=False):
 q=d.add_heading(t,level);place(q._p);q.paragraph_format.page_break_before=new;return q
def link(q,label,url):
 a=OxmlElement('w:hyperlink');a.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'295878');pr.append(c);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);q._p.append(a)
def refs(ids):
 q=p();q.paragraph_format.space_after=Pt(4);q.paragraph_format.line_spacing=Pt(13);q.add_run('资料  ').font.size=Pt(9)
 for i,k in enumerate(ids):
  if i:q.add_run('  ')
  link(q,k+' '+SOURCES[k][0],SOURCES[k][1])
def table(headers,rows,widths=None):
 widths=widths or [16.6/len(headers)]*len(headers);t=d.add_table(rows=1,cols=len(headers));t.autofit=False;place(t._tbl)
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v if v is not None else '未公开取得')
 for i,row in enumerate(t.rows):
  trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
  if i==0:trp.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1' if i==0 else ('F7F8FA' if i%2==0 else 'FFFFFF'));pr.append(shade)
   mar=OxmlElement('w:tcMar')
   for side,w in [('top',80),('bottom',80),('left',85),('right',85)]:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),str(w));x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   for q in c.paragraphs:
    q.paragraph_format.line_spacing=Pt(15);q.paragraph_format.space_after=Pt(0);q.paragraph_format.keep_with_next=(i==0)
    for run in q.runs:run.font.size=Pt(10);run.bold=i==0
 p('');return t
# Preserve the full-market classification table alongside the case selection table.
# Remove EO from the main case chapter; keep it as a short specialist appendix.
remove_range('4.2 EO中国华东分会','4.3 长江商学院')
for p0 in d.paragraphs:
 t=p0.text
 mapping={'4.3 长江商学院':'4.2 长江商学院','4.4 木兰汇':'4.3 木兰汇','4.5 正和岛':'4.4 正和岛','4.6 优家荟':'4.5 优家荟'}
 if t in mapping:p0.text=mapping[t]
# Replace selection chapter rather than appending a correction to a stale list.
a=find('3.2 六家为何值得学习');end=find('四 六家机构深度研究')
remove_range('3.2 六家为何值得学习','四 六家机构深度研究');anchor=end
h('3.2 八家对标对象及其用途')
p('保留五家会员、学习与生活方式组织，增加三家专业婚恋机构，共八家。选择依据是能否解释本项目的一项关键经营问题，以及是否有产品、收费、服务或实际活动证据。婚恋样本同时覆盖产品组合、区域活动和规模转型，避免只选择宣传成功的公司。')
table(['机构','入选理由','已取得的主要证据'],[r[:3] for r in ROWS],[3.3,5.9,7.4])
p('这八家属于不同用途的学习对象，其中三家婚恋机构均未公开证明A8客户规模。婚介高合同额、五星酒店场景及“高端”名称，均不作为资产准入或经营成功的证据。优家荟的现价仍缺，保留它是为了研究持续女性关系与跨城活动。')
h('3.3 专题参考和暂不升级的对象')
p('EO保留固定小组、保密和主持培训专题；其全球续会信息用于检验持续关系的组织方法，不占主案例名额。领教工坊和伟事达同属私董会专项。中欧、北大汇丰保留商学补充；观澜湖、黑骑士保留运动会员专项。')
p('MarryU存在活动和旅行宣传线索，但本轮未取得足以替换上述三家的现行收费、旅行交付与独立经营数据。纯旅行供应商继续作为合作资源，不因目的地或酒店昂贵进入高端社群核心样本。')
anchor=None
# Update lead text while preserving sources, illustrations and other five cases.
for p0 in d.paragraphs:
 t=p0.text
 if t=='四 六家机构深度研究':p0.text='四 八家机构深度研究'
 elif '候选发现清单包含55条记录' in t:p0.text='木兰汇与优家荟分别按企业家组织和生活方式组织理解，女性是客户标签。旅行社及供应商另列合作资源；尚未核实主体和近期活动的账号保留为线索。完整机构地图和原始记录见Excel。'
 elif t.startswith('这一章先把六家'):p0.text='本章先说明各家实际做什么、谁会购买，再看价格、权益、获客、活动与经营。前三类婚恋事实分别标明当前公开商品、历史合同个案和已公布活动，详细出处可以点击。'
 elif t.startswith('六家对标分别解决不同问题'):p0.text='八家对标分别承担具体研究任务：马会研究家庭会籍和日常使用，长江研究共同学习，木兰研究女性经营者与活动分层，正和岛研究商业内容获客，优家荟研究生活体验及旅行；我主良缘、金凤凰和珍爱网补齐专业匹配、活动收费与规模服务管理。EO改为私董会专题。'
 elif t.startswith('资料复核截至') or t.startswith('资料截至2026年10月21日'):p0.text='婚恋机构及旅游结合专题复核截至2026年10月22日 其他市场资料保留原统计期'
 elif t.startswith('六家没有一项'):p0.text='八家收费分别对应会籍资格、持续服务、课程、单场体验和婚介合同。婚恋价格进一步分为线上功能、顾问匹配、增值课程和活动票；同一客户累计购买金额与单份合同分别比较。'
 elif t.startswith('可见会员经营数据中'):p0.text='经营比较先看已经交付了什么。马会有现行会籍规则，商学院和会员社群有近期活动，金凤凰有单场报价与报名；婚恋机构的合同金额、服务变更和退款记录补充了收入质量这一面。EO全球续会仍作为专题证据。'
 elif t.startswith('建议把“高端生活社交”'):p0.text='建议把高端生活社交作为品牌范围，首期在生活社交、单身主题社交与独立顾问匹配中选择一个主产品。若团队已有经营者关系与主持能力，私董会可以先行；婚恋有明确需求时，应与旅行同等进入试点，而非始终作为最后追加的服务。'
 elif t.startswith('六项业务分阶段推出'):p0.text='六项业务按各自的付费和交付证据推进。旅行按次售卖，单身主题活动与一般会员旅行分别招募；持续婚介另签服务期、介绍次数及反馈合同。伴侣与已婚会员购买生活、学习或家庭服务，不进入单身匹配名单。'
# Replace cross-case tables by their identifying header or row labels.
for t in list(d.tables):
 head=[c.text for c in t.rows[0].cells]
 if head==['收费结构','对应案例','新公司需要明确的内容']:
  anchor=t._tbl
  table(head,[['入会费加月费加消费','北京马会','使用资格、提名人数、家属和实际消费分别定价'],['课程费','长江','完整周期、固定上课及参访范围'],['分档活动票','木兰汇','席位、餐饮、赠票及展示权益'],['数字学习到会员','正和岛','内容购买者与核心会员分别维护'],['会籍加跨城项目','优家荟','常态权益和旅行额外支出'],['匹配加课程加活动','我主良缘、珍爱','各模块价格、次数、工时及退款价值分配'],['VIP分层加单场活动','金凤凰','寻访地域、团队、会期与活动票分开']],[4.3,3.8,8.5]);t._tbl.getparent().remove(t._tbl);anchor=None
 elif head and head[0]=='组织机制':
  anchor=t._tbl;table(['组织机制','对标证据','公司可执行的做法'],[['固定成员重复见面','长江班级、EO专题小组','固定日期与缺席回访'],['共同做事和小范围交流','马会DIY、景区活动分区互动','安排共同任务和自由交流时间'],['成员提供内容与接待','木兰、正和岛、优家荟','会员分享及企业互访'],['活动后独立匹配','金凤凰活动与VIP分开','只向自愿单身客户提供后续婚介'],['客户进入恋爱阶段','珍爱暂停条款与履约争议','匹配暂停，生活会员可继续；合同分别处理']],[4.2,5.6,6.8]);t._tbl.getparent().remove(t._tbl);anchor=None
 elif head==['机构','粉丝或互动取得情况','已确认经营证据及范围']:
  oldrows=[[c.text for c in r.cells] for r in t.rows[1:] if 'EO' not in r.cells[0].text]
  oldrows += [['我主良缘','本轮未新增可核验的平台指标','全国连锁规模自述；历史合同和课程交付个案；有效付费及财务未公开'],['金凤凰','本轮未取得已核验机构主号指标','官网活动目录、99元59报名；1月免费35报名；实到、VIP及利润未公开'],['珍爱网','本轮未新增可核验的平台指标','1.4亿为官网累计用户自述；2025模式调整及服务合同个案']]
  anchor=t._tbl;table(head,oldrows,[3.0,5.4,8.2]);t._tbl.getparent().remove(t._tbl);anchor=None
 elif head and head[0]=='方向':
  anchor=t._tbl;table(head,[['生活与学习会员社群','连续主题活动、参访及旅行','已有同城客户、内容与接待能力','只购买目的地、不愿继续同群参加'],['单身主题社交与旅行','短途共同体验、分组交流、回访','真实单身客户、男女结构与匹配意愿','依赖赠送名额，或筛选和组织溢价无人付费'],['顾问婚恋服务','期限内推荐、双方同意约见和跟进','合适人选供给、顾问产能、清楚合同','合格约见不足、退款及交付工时持续侵蚀毛利']],[3.3,4.3,4.6,4.4]);t._tbl.getparent().remove(t._tbl);anchor=None
# Add a current-knowledge note for the already explored membership system.
anchor=find('4.2 长江商学院')
p('家庭与消费补充核查：香港个人会籍、北京内地会籍、公司提名及配偶独立会籍分别适用规则。餐饮、住宿、运动课程和马术按项目另付，会费不视为消费额度；免费附属卡与独立全费配偶会籍不能互换。完整三图另附，主比较仍以北京会所人民币产品为准。')
anchor=find('五 横向比较与公司机会')
for i,c in enumerate(CASES,6):
 h(f'4.{i} {c["name"]}',2,True)
 for title,paras,ids in c['sections']:
  h(title,3)
  for tx in paras:
   if tx.startswith(('我们可以学习在共同活动中安排','你们的绩效应同时看首次付款')):continue
   p(tx.replace('恋练同款','恋综同款'))
  if title=='具体收费和买到的东西' or title=='具体收费和会员权益':
   rr=[x for x in PRICES if x[0]==c['name'] and x[1] not in ['现行旅游相亲产品','福州旅游相亲活动']]
   table(['收费项目','金额及单位','包含内容和时间'],[[r[1],('未公开取得' if r[2] is None else f'{r[2]:,}元')+'／'+r[4],r[6]+'。'+r[5]+'；'+r[9]] for r in rr],[4.1,4.0,8.5])
  if title=='具体收费和会员等级':
   table(['会员层次','服务区别','收费与使用说明'],VIP,[3.2,6.5,6.9])
  if c['name']=='金凤凰深港婚恋' and title=='旅游结合和对我们的启发':ids=['M18']
  if c['name']=='珍爱网' and title=='我们具体学什么':ids=['M19','M21']
  refs(ids)
anchor=None
# Insert a separate, explicit travel comparison before the strategy chapter.
anchor=find('六 进入方案与经营验证')
h('5.5 婚恋与旅游结合的实际证据')
p('三家都能找到旅游结合的表述或历史项目，但本轮没有一家同时公开近期多日高端行程、完整售价、实到规模与重复购买结果。当前最具体的收费案例是景区单日相亲，不应把它扩写为成熟的高净值婚恋旅行市场。')
table(['机构及项目','已确认的产品与价格','研究判断'],[[r[0],r[2],r[4]] for r in TRAVEL],[4.0,7.0,5.6]);refs(['M02','M08','M14','M15','M17','M18','M24'])
h('5.6 对本公司旅游产品的具体建议')
p('单身旅行应先确定报名条件、年龄范围、男女名额及交友目的，再设计行程。普通会员旅行可接受伴侣和已婚成员，主题以共同兴趣为主。两种产品分别招募和沟通，避免客户以为自己参加的是单身团，到场后才发现目的不同。')
table(['阶段','建议产品与交付','需要记录的经营数据'],[['首轮测试','本城半日体验或周边一日活动；共同任务、轮换小组、自由交流和自愿回访','申请、符合条件、付款、取消、实到、每位服务工时及单场贡献'],['重复购买测试','同一批客户第二次收费活动；愿意者进入两日一夜主题旅行','同批第二次付款、退款、出游后继续联系与再次参加'],['多日旅行','住宿房型与单房差、交通、餐食、保险、主持与顾问分别列明；成团与取消有明确节点','成团率、供应商成本、会员服务成本、退款及归来后再购'],['婚介升级','旅行后自愿申请顾问匹配，单独签合同和收费','双向同意约见、实际约见、有效反馈、暂停及退款']],[3.0,8.0,5.6])
p('这是拟测试的产品方案，未作为竞品实际流程。价格应由真实供应商报价、顾问工时和客户付款测试决定。先用短途确认人群与组织价值，再决定是否值得增加跨城、多日和高住宿成本。')
anchor=None
# Update evidence gaps and remove old mixed ranking language.
for t in list(d.tables):
 if [c.text for c in t.rows[0].cells]==['对象','本轮已查到哪里','仍缺什么及用途']:
  rr=[[c.text for c in r.cells] for r in t.rows[1:] if 'EO' not in r.cells[0].text]
  rr += [['我主良缘','四品牌、App商品、历史合同与派聚说明','当期完整合同、付费人数、约见完成与旅行成团'],['金凤凰','VIP层次、2026活动价及报名、历史景区项目','现行VIP费用与会期、实到、活动到VIP转化、A8客户规模'],['珍爱网','官网五步流程、合同及2025模式调整','当期VIP报价、活跃付费、服务网络及外包后单位经济性']]
  anchor=t._tbl;table(['对象','已取得的具体资料','仍需取得的经营证据'],rr,[3.1,6.2,7.3]);t._tbl.getparent().remove(t._tbl);anchor=None
# Add source list and specialist appendix at the end.
h('C EO固定小组的专项参考',2)
p('EO不进入八家主案例，保留为私董会专题。官方Forum由6至10位企业经营者组成固定小组，每月见面，由受训主持人组织，并有保密及经验分享规则；另有配偶和生活伴侣Forum。它对我们最有用的是稳定组员、主持培训与出席维护。')
p('原研究取得的华东109位会员与全球FY2024/25续会91.6%分别保留原地区和观察期。全球续会可以支持固定关系产品的研究，不能说明它能替代婚恋匹配，也不能直接用于预测本公司留存。')
q=p();link(q,'EO官方Forum机制','https://eonetwork.org/membership/forum/')
h('D 婚恋及旅游专题新增资料',2)
p('价格证据分为当前公开商品、机构公布活动价、历史合同报道和消费者单方线索。只有前三类已核对来源的事实进入核心比较；单方线索仅保留验证用途。本轮对三家财务、有效付费及A8客群规模均未获得完整披露。')
p('本轮检索机构网站、规则、活动页、原版报纸和监管发布，Crawl4AI另渲染我主良缘活动页与珍爱首页。MediaCrawler访问小红书首页超时，本轮没有新增可核验粉丝、评论或分享数据；旧平台快照仍保留2026年10月21日采集日期。')
for k,(title,url,scope,kind) in SOURCES.items():
 q=p();q.paragraph_format.space_after=Pt(5);q.paragraph_format.line_spacing=Pt(14)
 link(q,k+' '+title,url);q.add_run('\n'+scope+'；'+kind+'；核查2026-10-22').font.size=Pt(9)
# Clean stale inline occurrences, maintain existing hyperlinked source runs.
for p0 in d.paragraphs:
 for r in p0.runs:
  updated=r.text.replace('六家中的证据','对标证据').replace('六家深度对标','八家深度对标').replace('六家为深度','八家为深度')
  if updated!=r.text:r.text=updated
# Build clean internal TOC from final chapter headings and case headings.
toc=find('目录');nextchapter=find('二 市场与客户需求');active=False
for el in list(d._element.body):
 if el is nextchapter:break
 if active:d._element.body.remove(el)
 if el is toc:active=True
# Remove prior bookmarks to avoid stale destinations.
for tag in ['bookmarkStart','bookmarkEnd']:
 for el in list(d._element.body.iter(qn('w:'+tag))):el.getparent().remove(el)
nav=[]
for idx,p0 in enumerate(d.paragraphs,1000):
 if p0.style.name in ['Heading 1','Heading 2'] and p0.text!='目录':
  key=f'review_{idx}';b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(idx));b.set(qn('w:name'),key);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(idx));p0._p.insert(0,b);p0._p.append(e)
  if p0.style.name=='Heading 1' or re.match(r'4\.[1-8] ',p0.text):nav.append((p0.text,key))
anchor=nextchapter
t=table(['内容','页码'],[[a,''] for a,b in nav],[15.1,1.5])
for row,(_,key) in zip(t.rows[1:],nav):
 q=row.cells[0].paragraphs[0];s=q.text;q.clear();a=OxmlElement('w:hyperlink');a.set(qn('w:anchor'),key);r=OxmlElement('w:r');tx=OxmlElement('w:t');tx.text=s;r.append(tx);a.append(r);q._p.append(a)
anchor=None
# Existing first-page insight table: remove old EO-forward narrative.
for t in d.tables:
 if [c.text for c in t.rows[0].cells]==['发现','对公司的含义']:
  rr=[['婚恋需要专业交付','三家新增案例补齐匹配、辅导、服务期、退款和转型管理。'],['旅行结合已有实践','先研究单日真实产品；近期多日高端经营结果尚缺。'],['家庭与单身服务分开','一般会员活动接受伴侣家庭；婚恋匹配只面向自愿单身。'],['高收费不等于高净值客群','每家分别核对客群条件、有效人数与实际履约。'],['首期投入依赖付款与成本','连续活动验证人群及再购，合同明确各项服务如何完成。']]
  for row,vals in zip(t.rows[1:],rr):
   for c,v in zip(row.cells,vals):c.text=v
# Remove spacer-only paragraphs and keep source notes with the evidence they support.
from docx.text.paragraph import Paragraph
for el in list(d._element.body):
 if el.tag==qn('w:p') and not text(el) and not any(list(el.iter(qn('w:'+tag))) for tag in ['drawing','pict','br','fldChar']):d._element.body.remove(el)
for q in d.paragraphs:
 if q.text.startswith(('资料  ','依据  ')):
  prev=q._p.getprevious()
  if prev is not None and prev.tag==qn('w:p'):Paragraph(prev,d._body).paragraph_format.keep_with_next=True
normal=d.styles['Normal'].paragraph_format;normal.line_spacing=Pt(18);normal.space_after=Pt(6)
d.styles['Heading 3'].paragraph_format.space_before=Pt(8);d.styles['Heading 3'].paragraph_format.space_after=Pt(5)
# Save final report and maintain common styles.
d.core_properties.title='中国高端会员社群市场研究及创业可行性评估报告';d.core_properties.subject='八家对标及婚恋旅游专题更新 2026年10月22日'
for p0 in d.paragraphs:
 if p0.style.name.startswith('Heading'):p0.paragraph_format.keep_with_next=True
report=OUT/'中国高端会员社群市场研究及创业可行性评估报告.docx';d.save(report)
# Create case archive using same package, relationships and styles.
c=Document(report);body=c._element.body;start=False
for el in list(body):
 tx=text(el)
 if tx=='四 八家机构深度研究':start=True
 if tx=='五 横向比较与公司机会':start=False
 if not start and el.tag!=qn('w:sectPr'):body.remove(el)
c.paragraphs[0].text='八家会员社群与婚恋机构对标案例档案';c.paragraphs[0].style='Title';c.paragraphs[0].paragraph_format.page_break_before=False
c.core_properties.title=c.paragraphs[0].text;c.save(OUT/'八家会员社群与婚恋机构对标案例档案.docx')
# Partner summary independently updated, not stale six-case prose.
d=Document(report)
for el in list(d._element.body):
 if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
anchor=None
p('高端会员社群创业项目\n合伙人决策摘要','Title');p('婚恋及旅游专题更新 2026年10月22日','Subtitle')
p('主研究样本调整为八家：保留五家会员与学习组织，新增我主良缘、金凤凰深港婚恋和珍爱网。EO改为固定小组专题。三家婚恋机构的入选理由是能拆解产品与交付，不是已经证明其服务大量A8客户。')
h('一 本轮新增发现',1)
table(['机构','最值得研究的内容','经营判断'],[['我主良缘','匹配、形象、情感与派聚社交组合；iOS68元月卡、198元年卡；历史3万元4对象合同','组合与我们接近；顾问和课程交付重，当前定制报价及有效客户未公开'],['金凤凰','VIP按对象地域与团队分层；9月99元活动59人报名；历史景区包价','适合研究深港区域与活动到婚介；未证实A8客群规模'],['珍爱网','五步红娘服务；2024年一年3.44万元合同个案；2025服务外包调整','规模服务与转型案例；需要反向研究获客承诺、交付和退款']],[2.8,7.1,6.7])
refs(['M06','M07','M11','M14','M19','M21','M22'])
h('二 旅游结合值得做什么验证',1)
p('我主良缘保留派聚旅行社交定位；金凤凰参与过景区相亲打包，2019年男200元、女100元，含门票、大巴和午餐；珍爱网有2011年漓江历史活动。近期同时具备完整报价、行程、实到与复购的多日高端婚恋产品，本轮未取得。')
p('建议先测试单身周边一日活动，观察同批客户是否愿意再次付款，再推出两日一夜。普通会员旅行与单身专场分别招募；匹配、辅导与旅行按具体交付定价。金凤凰“山海皆可赴”实际为罗湖室内联谊，不计作旅行。')
refs(['M02','M18','M15','M24'])
h('三 八家研究分工',1,True)
table(['对象','用来解决的经营问题'],[[r[0],r[1]] for r in ROWS],[3.5,13.1])
h('四 合伙人这次需要确定的事情',1)
p('先确定首个城市和首个产品：生活学习会员、单身主题活动或顾问婚介。再确认能邀请到的客户、筛选负责人、主持及顾问、旅游承接方和预算。婚恋业务有明确需求时，应直接进入付费试点。')
table(['指标','记录方式'],[['客户质量','申请、符合条件、真实单身、双方意愿；资产口径另行定义'],['销售与体验','付款、实到、取消、30天内同批第二次付款'],['婚介交付','双向同意、实际约见、反馈、暂停与退款'],['旅行经营','成团、实到、供应商成本、服务工时、项目贡献及再购']],[3.5,13.1])
p('三家均缺完整当期定制价、有效付费、财务及高净值客户规模。下一轮最有价值的验证是取得正式合同与报价，并访谈最近真正付费和退出的客户；本轮未外联、未填咨询表单。')
d.core_properties.title='高端会员社群创业项目合伙人决策摘要';d.save(OUT/'高端会员社群创业项目合伙人决策摘要.docx')
(BASE/'报告结构.json').write_text(json.dumps({'core_names':NAMES,'selection':ROWS,'nav':nav,'sources':SOURCES},ensure_ascii=False,indent=2))
print('created',report,'and cases + summary')
