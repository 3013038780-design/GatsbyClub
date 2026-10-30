from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).parent;data=json.loads((B/'启动方案内容.json').read_text());d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
font='Source Han Serif CN'
for name,size,bold in [('Normal',11,False),('Title',24,True),('Heading 1',18,True),('Heading 2',13,True),('Subtitle',12,False)]:
 st=d.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor(0,0,0)
 for key in ['ascii','hAnsi','eastAsia','cs']:st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+key),font)
 st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.25
for st in d.styles:
 for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
f=sec.footer.paragraphs[0];f.alignment=2;r=f.add_run('俱乐部启动工作方案  ');r.font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for idx,p in enumerate(data['pages']):
 if idx:d.add_page_break()
 d.add_paragraph(p['title'],'Title'if idx==0 else'Heading 1')
 if idx==0:d.add_paragraph('2026年10月30日  合伙人讨论与修改')
 for para in p['paras']:d.add_paragraph(para)
 if p['headers']:
  t=d.add_table(rows=1,cols=len(p['headers']));t.autofit=False
  widths=([3.2,8.4,5.4]if len(p['headers'])==3 else[3.5,13.5]);
  if len(p['headers'])==3 and idx==13:widths=[7,5,5]
  for j,width in enumerate(widths):t.columns[j].width=Cm(width)
  for j,h in enumerate(p['headers']):t.rows[0].cells[j].text=h
  for row in p['rows']:
   c=t.add_row().cells
   for j,v in enumerate(row):c[j].text=v.replace(' ','\n',1)if idx==3 and j==0 else v
  for i,row in enumerate(t.rows):
   prop=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');prop.append(cant)
   if i==0:prop.append(OxmlElement('w:tblHeader'))
   for j,c in enumerate(row.cells):
    c.width=Cm(widths[j]);pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7ECE8'if i==0 else('F5F5F0'if i%2 else'FFFFFF'));pr.append(sh)
    for q in c.paragraphs:
     q.paragraph_format.space_after=Pt(5);q.paragraph_format.space_before=Pt(4);q.paragraph_format.line_spacing=1.15
     for run in q.runs:run.font.size=Pt(10);run.bold=(i==0)
  d.add_paragraph('')
 for para in p['notes']:
  q=d.add_paragraph(para);q.paragraph_format.space_after=Pt(7)
  for r in q.runs:r.font.size=Pt(10)
d.core_properties.title='俱乐部启动工作方案';d.core_properties.author='GatsbyClub';d.save(B/'交付文件/俱乐部启动工作方案.docx');print('Manual DOCX built')
