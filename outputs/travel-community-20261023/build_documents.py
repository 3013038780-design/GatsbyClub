from pathlib import Path
import json,math,shutil
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
B=Path(__file__).resolve().parent;O=B/'deliverables';D=json.loads((B/'研究与产品数据.json').read_text());M=json.loads((B/'模型核查.json').read_text())
S={s['id']:s for s in D['sources']};F=FontProperties(fname='/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
plt.rcParams.update({'axes.spines.top':False,'axes.spines.right':False,'axes.spines.left':False,'axes.grid':False,'font.size':11})
fig,ax=plt.subplots(figsize=(10,4.3));cats=['晚宴','商务分享会','三亚旅行'];x=range(3)
for i,p in enumerate(D['products']):
 vals=[p['price']*n*.97-n*(p['var']+p['cac'])-p['fixed'] for n in [p['low'],p['base'],p['full']]]
 for j,v in enumerate(vals):ax.bar(i+(j-1)*.23,v/10000,width=.21,color=['#ADB9C3','#26475A','#9B7138'][j],label=['低于预期','基准','满团'][j] if i==0 else None)
ax.set_xticks(list(x),cats,fontproperties=F);ax.axhline(0,color='#909AA2',lw=.7);ax.set_ylabel('单次贡献 万元',fontproperties=F);ax.legend(prop=F,frameon=False,ncol=3);ax.set_title('三种人数下的贡献  全部为规划假设',fontproperties=F,pad=20);fig.tight_layout();fig.savefig(B/'qa/人数与贡献.png',dpi=180);plt.close(fig)
fig,ax=plt.subplots(figsize=(10,4));n=14;p=D['products'][2];items=['直接交付','获客','固定及工时','退款情景','单团贡献'];vals=[n*p['var'],n*p['cac'],p['fixed'],p['price']*n*.03,M['results']['travel']['contribution']]
for i,(label,v,c) in enumerate(zip(items,vals,['#8FABB8','#C1CBD0','#566B79','#D0B992','#26475A'])):
 ax.barh([i],v/10000,color=c,height=.55);ax.text(v/10000+.14,i,f'{v/10000:.2f}',ha='left',va='center',fontproperties=F,fontsize=11)
ax.set_xlim(0,max(vals)/10000*1.17);ax.set_yticks(range(5),items,fontproperties=F);ax.invert_yaxis();ax.set_xlabel('万元',fontproperties=F);ax.set_title('三亚基准14人  27.72万元收款如何分配',fontproperties=F,pad=18);ax.spines['bottom'].set_visible(False);fig.tight_layout();fig.savefig(B/'qa/旅行收入分配.png',dpi=180);plt.close(fig)

def new(title,subtitle):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2.0);s.bottom_margin=Cm(1.9);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
 for name,size,bold,cn in [('Normal',11,False,'Songti SC'),('Title',24,True,'Heiti SC'),('Subtitle',12,False,'Songti SC'),('Heading 1',17,True,'Heiti SC'),('Heading 2',14,True,'Heiti SC'),('Heading 3',11.5,True,'Heiti SC')]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),cn);st.paragraph_format.line_spacing=Pt(18.5);st.paragraph_format.space_after=Pt(6)
  for z in st.element.xpath('.//w:pBdr'):z.getparent().remove(z)
  if name.startswith('Heading'):st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(14)
 d.add_paragraph(title,'Title');d.add_paragraph(subtitle,'Subtitle');d.add_paragraph('2026年10月23日  私有工作资料','Subtitle')
 h=s.header.paragraphs[0];h.text='高端单身旅行与线下社群研究';h.style=d.styles['Normal'];h.runs[0].font.size=Pt(9)
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 return d
def p(d,t,boldlead=False):
 q=d.add_paragraph()
 if boldlead and '：' in t:
  a,b=t.split('：',1);q.add_run(a+'：').bold=True;q.add_run(b)
 else:q.add_run(t)
 return q
def h(d,t,level=1):return d.add_paragraph(t,'Heading '+str(level))
def link(q,text,url):
 if not url:q.add_run(text);return
 z=OxmlElement('w:hyperlink');z.set(qn('r:id'),q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'255878');pr.append(col);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);z.append(r);q._p.append(z)
def cite(d,ids):
 q=d.add_paragraph();q.paragraph_format.space_after=Pt(9);q.paragraph_format.line_spacing=Pt(14)
 for i,k in enumerate(ids):
  if i:q.add_run('  ')
  s=S[k];link(q,k+' '+s['name'],s['url'])
 for r in q.runs:r.font.size=Pt(9)
 return q
def table(d,headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for r in rows:
  for c,v in zip(t.add_row().cells,r):c.text=str(v)
 pr=t._tbl.tblPr;bor=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bor.append(e)
 pr.append(bor)
 for i,r in enumerate(t.rows):
  tr=r._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  if i==0:tr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(r.cells):
   if widths:c.width=Cm(widths[j])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tc=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tc.append(mar)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF1');tc.append(sh)
   for q in c.paragraphs:
    q.paragraph_format.line_spacing=Pt(15);q.paragraph_format.space_after=Pt(0)
    if len(t.rows)<=8 and i<len(t.rows)-1:q.paragraph_format.keep_with_next=True
    if i==0:q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for z in q.runs:z.font.size=Pt(10);z.bold=(i==0)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def picture(d,file):
 q=d.add_paragraph();q.paragraph_format.line_spacing=1;q.add_run().add_picture(str(B/'qa'/file),width=Cm(16.2))
def cost(p,n):return p['price']*n*.97-n*(p['var']+p['cac'])-p['fixed']

d=new('高端单身旅行与线下社群研究报告','公开研究与首发产品初案')
p(d,'建议先验证北京小型单身晚宴，再决定是否启动三亚旅行。商务学习线独立验证重复付费。现有证据支持“筛选同行者、组织持续相处”的产品方向；首轮试价仍需客户反馈和同口径供应商报价。')
p(d,'本版复用36家机构地图、五家婚恋深挖与国外九家案例；新增旅行社型社群和具体旅游商品核查。Olivia已确认深圳案例为BoL／BoL Selected，目前暂无已确认交付资源。')
h(d,'阅读顺序',2)
for v in ['一 合伙人决策摘要','二 市场与首发客群','三 对标选择与国内案例','四 国外九家案例的可用做法','五 两场活动和一条旅行','六 费用比较与经营测算','七 七个工作阶段与分工','附录 咨询访谈问题及来源']:p(d,v)
d.add_page_break()
h(d,'一 合伙人决策摘要')
p(d,'首发建议：用一场1680元试价、20人基准的北京单身主题晚宴验证筛选、主持与再次见面的能力。旅行先保留四天三晚三亚方案，研究试价19800元／人、机票自理、每人独住。商务分享会以680元试价单独核算，优先压低场地和讲师固定成本。',True)
table(d,['需要决定','建议及成立条件'],[
['首发客群','北京常住、30至42岁、真实单身且愿意建立稳定关系；个人可投资资产1000万元是待讨论口径，家庭资产单列。'],
['首条线路','三亚四天三晚为唯一旅行初案；北京周边两天一晚用于报价比较与备选，暂不同时开发第二款。'],
['活动与试价','晚宴1680元；商务分享680元；旅行19800元。都是研究试价，无现有成交支持。'],
['预算','基准周转金晚宴10728元、分享8682元、旅行75216元。若全部同时准备，简单相加约9.46万元；建议分阶段投入。'],
['执行分工','研究与材料由我整合；Olivia负责供应商询价及潜客访谈；格格核实资源；两位合伙人确定预算和交付责任。']],[3.0,13.4])
p(d,'最值得先学的六个对象：32号和游侠客负责旅行交付参照，钻石婚恋和我主良缘负责婚恋筛选与旅行延伸，BoL／BoL Selected和大超负责双线活动与主理人获客。其余如爱、金凤凰、南充囍上媒捎、北京爱旅、MeetKnow保留为具体产品样本；稻草人继续候选。六个对象是研究用途组合，尚未形成六家同等完整的旅行经营案例。')
p(d,'当前最关键的三个缺口：相同日期和房型的供应商报价、具体客户的愿付反馈、可核的单身旅行实际成团与复购。它们决定19800元能否成立，不能由浏览量或豪华酒店照片替代。')
h(d,'二 市场与首发客群')
h(d,'2.1 最新资料改变了什么',2)
p(d,'2025年结婚登记676.5万对，同比增长10.8%，采用年度公报口径。婚恋需求仍有规模，但它没有给出高资产单身旅行的付费人数。首期经营规模需要从北京可触达名单、日期与实际付款逐步建立。');cite(d,['M05'])
p(d,'胡润2026年调查覆盖470位千万资产人士，78%来自一线城市，平均家庭总资产6100万元、平均可投资资产2300万元。受访者计划服务体验支出增加12%，每年假期21天。对产品的含义是：用有限假期完成有质量的共同体验，减少赶路和强制互动；这份样本没有单独证明我们的单身客群愿付19800元。');cite(d,['M04'])
p(d,'现有选择包括朋友组局、商学院或商会、兴趣旅行、红娘服务和私人旅行定制。我们的购买理由应具体到“同行者条件可理解、先见面再出行、独住和隐私舒适、现场有人照顾交流”。六项服务不必一次推出，先检验单身晚宴与旅行之间是否存在真实购买承接。')
h(d,'2.2 客户分组及首轮访谈假设',2)
p(d,'以下是我们拟验证的客群，不是竞品实际购买者统计。年龄和预算是讨论范围；收入只记录客户自述或可核资料，不凭职业推定。资产统一记录个人总资产、个人可投资资产、家庭总资产三个字段。')
table(d,['客群','工作生活与参与动机','偏好与需验证内容'],[
['经营者单身客\n约32至42岁','企业经营者、业务负责人、专业服务合伙人；工作时间不稳定，希望缩短认识合适人的筛选成本。','偏好小桌与私密介绍；可能喜欢高尔夫、酒食或轻运动，逐人询问。验证能否提前两周锁定周四至周日。'],
['专业职业单身客\n约30至40岁','金融、科技、法律、咨询、医疗等岗位或独立专业人士；收入与资产分别问，希望摆脱单一工作圈。','偏好有共同话题和明确边界，内向外向均可；验证独住、性别结构、年龄差和拍摄授权。'],
['离异后重新建立关系\n约35至48岁','有过婚姻、可能有子女或照料责任；重视安全、关系意愿与出行时间。首轮不与全部年龄段混招。','确认带娃安排与周末可用性，避免公开婚史。是否独立开场由访谈与人数决定。'],
['商务学习客\n约30至50岁','经营者、管理者与专业服务合伙人；单身、恋爱中和已婚均可，围绕具体经营问题参加。','验证主题、讲师可信度与二次参加；兴趣不强行设统一标签，不把商业联系人承诺为婚恋资源。']],[3.2,6.6,6.6])
p(d,'性格记录应描述行为：喜欢一对一还是大组、愿不愿公开自我介绍、能接受多久集体活动、遇到陌生人是否希望主持人引导。兴趣记录最近半年真的花过时间或钱的项目、频次与金额。这样才能用于排座、活动节奏和旅行同伴安排。')
h(d,'三 对标选择与国内案例')
h(d,'3.1 选择依据与资料覆盖',2)
p(d,'先看是否能提供一项可借鉴的实际服务，再看收费、组织和经营证据是否足够。旅行社型学线路与领队，婚恋型学筛选和关系服务，社群型学获客与持续参与。纯高价旅行社只有在交付比较时进入，不作为高端社群成功案例。')
table(d,['研究组','重点对象','现阶段能学什么','还缺什么'],[
['旅行社主导','32号、游侠客','路线、含项、领队、报名与退改','同条件团期报价、独立实收人数、成本与复购'],
['婚恋机构主导','钻石、我主良缘；补充南充门店、如爱、金凤凰','财富与婚恋筛选、低价活动入口、长期服务另购','至少两家完整旅行经营档案尚未取得'],
['社群或主理人','BoL／BoL Selected、大超；补充爱旅线索','商务与单身分线、内容转线下、主持和回访','两家当前旅行商品和实际成团未取得'],
['成熟体系与候选','马会、长江等沿用；稻草人候选','权益设计、学习关系、青年同行','只作相应用途，不推导我们的会费或销量']],[2.8,4.2,4.8,4.6])
for i,c in enumerate(D['cases'],1):
 h(d,f'3.{i+1} {c["name"]} {c["product"]}',2)
 for label,key in [('做什么和谁参加','customer'),('怎样收费','price'),('付费买到什么','rights'),('怎样组织','flow'),('怎样获客','acquisition'),('经营表现','performance'),('我们的判断','judgment')]:p(d,label+'：'+c[key],True)
 cite(d,c['sources'])
h(d,'四 国外九家案例的可用做法')
p(d,'国外案例主要提供活动、会籍和关系延续的做法。以下现价以2026年10月23日可读官方页面为准；未公开金额保留为空。活动模块可以轻资产试做，空间投入和会费另行评估。')
for i,c in enumerate(D['foreign'],1):
 h(d,f'4.{i} {c["name"]}',2);p(d,c['what']);p(d,'收费与权益：'+c['fees'],True);p(d,'组织与获客入口：'+c['activity'],True);p(d,'经营逻辑与我们可以做的部分：'+c['judgment'],True);cite(d,c['sources'])
h(d,'五 两场活动和一条旅行')
p(d,'以下为团队产品初案。日期仅用于统一研究报价，人数、门槛、试价和流程尚未对外发布。晚宴与旅行共享单身申请资料及自愿联系机制，商务学习线使用独立报名目的。')
for i,c in enumerate(D['products'],1):
 h(d,f'5.{i} {c["name"]}',2)
 p(d,'谁参加：'+c['customer'],True);p(d,'为什么付款：'+c['reason'],True)
 table(d,['收费及名额','包含及另付'],[[f'试价{c["price"]:,}元／人\n基准{c["base"]}人，上限{c["capacity"]}人',c['rights']]],[4,12.4])
 p(d,'从哪里来：首轮通过合伙人直接认识或可信推荐的潜客邀请，辅以展示真实流程的品牌内容。未确认客户名单，模型将获客支出列为预算。',True)
 for step in c['flow']:p(d,step)
 p(d,'人员与工作量：'+c['team']+f'；规划总工时{c["hours"]}小时，需在试点计时校正。',True)
 p(d,'结束后为什么再来：'+c['retention'],True)
 if c['id']=='travel':
  p(d,'成团条件建议：至少12名、目标14名、上限16名；关系意愿、年龄范围和可相互交流的结构均须满足。财务保本6人不等于适宜成团人数。男女比例目标接近1比1，成团前公开实际结构，客户不同意时按约定退出。',True)
  p(d,'退改初案：供应商尚未报价，暂不对外承诺免费退期。申请筛选不收费；成团与合同确认后再收款，支付进度、不可退成本、天气取消和替换人员条件由承办合同确定。模型30%首款只用于现金演示。',True)
h(d,'六 费用比较与经营测算')
h(d,'6.1 同一线路怎样比较',2)
table(d,['条件','三亚研究组','北京周边研究组'],[['日期','2026年11月19至22日','2026年11月14至15日'],['天数及人数','4天3晚，16名，三亚集合','2天1晚，16名，北京集合'],['住宿','同一豪华酒店基础房，每人独住','同一度假酒店基础房，每人独住'],['共同成本','同车餐次、体验、保险与接送；机票另列','同车餐次、体验、保险'],['比较对象','单身社交团、精品私家团、DIY','单身短途活动、精品私家团、DIY']],[3.2,6.6,6.6])
p(d,'目前取得的5980元海南六日私家团是两人一间、海口环岛，含项和行程与三亚四日独住不一致。它可以解释为什么起价低，尚不足以计算我们的溢价。2026年8月酒店促销同样只保留历史背景。本轮未取得上表相同条件下三类可成交报价，同线路价差结论暂不成立。');cite(d,['D17','M02'])
p(d,'消费者账：旅行初案19800元，往返机票另按2500元情景演示，总支出22300元；升级房、超出饮品和个人消费另计。2500元不是实时机票报价。私家团与DIY总额留待同口径询价，不能填成0元。')
h(d,'6.2 单人和固定成本',2)
table(d,['项目','晚宴','商务分享','三亚旅行'],[['试价 元／人',1680,680,19800],['单人交付 元',550,160,8700],['单人获客预算 元',180,100,1200],['固定及工时成本 元',8500,8000,48000],['退款金额／收款情景','3%','3%','3%'],['总工时 小时',70,52,190]],[5,3.8,3.8,3.8])
p(d,'三亚8700元／人拆分为：独住房三晚4500、全程餐饮1600、体验1100、交通700、保险物料300、服务协调500。固定48000元为项目工时190小时×150元、团队机酒10500、场地5000、器材4000。以上均为规划输入，没有计入尚未确认的免费资源。')
p(d,'晚宴固定8500元为工时70小时×100元及场地1500元；分享会固定8000元为工时52小时×100元、讲师2000元和场地800元。固定成本一并纳入测算，避免把合伙人的时间算成免费。')
h(d,'6.3 不同人数下怎样赚钱',2)
rr=[]
for c in D['products']:
 z=M['results'][c['id']];rr.append([c['name'],f'{c["low"]}人／{cost(c,c["low"]):,.0f}元',f'{c["base"]}人／{cost(c,c["base"]):,.0f}元',f'{c["full"]}人／{cost(c,c["full"]):,.0f}元',str(z['breakeven'])+'人'])
table(d,['产品','低于预期贡献','基准贡献','满团贡献','财务保本'],rr,[4.2,3.3,3.3,3.3,2.3]);picture(d,'人数与贡献.png')
p(d,'单次贡献等于扣除退款情景后的收入，再减直接交付、获客、固定及工时成本；尚未扣公司日常开支和税费。商务分享会16人时亏1606元、24人只剩1590元，说明它目前更适合控制成本做需求验证。旅行14人基准贡献82284元依赖19800元试价，价格与人数必须一起验证。')
picture(d,'旅行收入分配.png')
p(d,'价格反证：旅行14人时，若票价比试价低15%即16830元，贡献降至41951元；若退款金额比例从3%增至10%，贡献为62880元。若同等体验的竞争者报价显著更低，而客户不愿为同行者质量另付费，应缩短产品或降低成本，再重新测算。')
h(d,'6.4 收钱和花钱的先后',2)
table(d,['时间','收支假设','管理含义'],[['收款前','先付40%固定成本及全部获客费','旅行基准先需36000元；尚未成团的投入也要预算'],['首款到账','客户付30%，供应商付70%直接成本，并结清固定成本','旅行累计现金为负66900元；加8316退款预留需75216元周转'],['尾款及履约','客户付其余70%，结清30%直接成本','全程结清并扣3%退款后剩82284元'],['取消变化','按已发生成本、供应商退款及客户退款分别记录','3%预留不是退款上限，不能动用预收款掩盖亏损']],[2.8,7.1,6.5])
p(d,'应把旅游合同、收款及供应商履约由具备相应许可的旅行承办方具体落实，我们的筛选与社交服务怎样签约收费另行明确。旅游法要求旅行社许可并向合格供应商订购；采用合作不等于借用许可证。资产核验可能涉及金融账户等敏感信息，应先确定必要性、单独同意、存储权限及删除机制，优先不留存完整证明。');cite(d,['L01','L02'])
p(d,'现金测算保守按项目工时报酬全部支付计算，包含合伙人工时预算；若延期支付，需按实际付款时间重算周转金。')
h(d,'七 七个工作阶段与分工')
p(d,'七个工作日按依赖顺序安排。公开研究与初案已推进到可讨论版本，外部访谈、资源和报价尚未开展，因此不把本次交付标为七天验证全部完成。')
table(d,['阶段','本次结果','仍需完成及责任'],[['1 复用与对标','36家历史名录、12对象13产品、九家国外文字；BoL身份确认','继续查具体交易记录，我负责'],['2 扩搜与主体','32号主体和商品、游侠客、婚恋旅行招募已核','独立社群旅行承办主体仍缺，我负责'],['3 报价与交付','已有三亚及周边统一询价条件、具体旅行含项','Olivia取得两家同口径报价，补完整经营案例'],['4 客户与资源','资源确认暂无；访谈及资源模板已备','Olivia访谈，格格核实资源；我归纳'],['5 首发产品','两场活动和一条旅行初案已完成','两位合伙人校准客群及交付'],['6 经济性','可修改模型、原生敏感性、现金及工时','报价与反馈到位后替换假设'],['7 材料与决策','报告、简化五表、网站、讲解稿保持同一数据','会议确定测试价格、预算、分工及是否试点']],[2.5,6.6,7.3])
h(d,'附录 咨询访谈问题及来源')
h(d,'供应商统一询价单',2)
for q in ['请按三亚11月19至22日、16位客人、每人独住房基础房、酒店三晚，分别报含税总价、单人价、最低成团人数和有效期；北京周边按11月14至15日同样填写。','请逐项写酒店和替代酒店、床型早餐、车辆座位和时长、餐标、体验门票、领队、保险以及司机领队住宿；每项是否含税和服务费。','10人、12人、14人、16人的总成本分别是多少？哪些固定、哪些随人增加？独住与双住分别多少钱？','谁签合同、谁收钱、谁开票、谁承办？请提供主体与旅行许可；社群筛选和主持由谁负责、如何计费？','定金尾款何时付？减员、未成团、航班取消、恶劣天气及酒店替换如何退费，已发生不可退成本怎样证明？','全年哪些月份可重复执行？酒店旺日和节假日如何加价？组织者优惠、佣金与免费名额有什么条件？']:p(d,q)
h(d,'客户访谈与资源记录',2)
for q in ['最近12个月真实参加过哪三场活动或旅行？组织者、价格、同行者、日期和是否再买。','年龄、职业岗位、典型每周工作安排、照料责任；个人收入与个人／家庭资产分别记录，是否愿意提供何种最少证明。','最近半年投入最多的两项兴趣是什么？每月几次、和谁参加、花费多少。偏好一对一还是小组，需要怎样的主持引导。','能否参加给定周四至周日旅行？独住和双住怎样选？看到匿名同行构成后，晚宴1680与旅行19800会选择哪项、理由是什么；记录反对与替代选择。','最担心谁参加、被拍摄、隐私、过度推销、住宿或退改中的哪一点？上次不再参加一项服务的原因是什么。','资源记录统一填名称、联系人、可提供事项、价格、条件、有效期、书面凭据，状态为已确认可用／需要协商／目前只有线索。当前没有已确认可用项。']:p(d,q)
h(d,'下一次会议五项决定',2)
for q in ['首发客群及资产核验口径，是否把离异群体另开场。','首发旅行是否选三亚；接受的天数、独住与总支出范围。','先晚宴还是同时安排分享会，各自的最低报名和服务标准。','试价、首轮可承受亏损及准备金；若客户不接受，怎样调整。','Olivia和格格的询价、访谈、资源及执行分工，反馈截止日和下一次会议日期。']:p(d,q)
h(d,'资料入口',2)
for s in D['sources']:
 q=d.add_paragraph();link(q,s['id']+' '+s['name'],s['url']);q.add_run('。'+s['kind']+'；'+s['period']+'。'+s['note'])
 for r in q.runs:r.font.size=Pt(9.5)
for i,q in enumerate([q for q in d.paragraphs if q.style.name=='Heading 1'],1):
 anchor='chapter_'+str(i);z=OxmlElement('w:bookmarkStart');z.set(qn('w:id'),str(i));z.set(qn('w:name'),anchor);q._p.insert(0,z);z2=OxmlElement('w:bookmarkEnd');z2.set(qn('w:id'),str(i));q._p.append(z2)
 for toc in d.paragraphs:
  if toc.style.name=='Normal' and toc.text==q.text:
   toc.clear();hl=OxmlElement('w:hyperlink');hl.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=q.text;r.append(t);hl.append(r);toc._p.append(hl)
d.save(O/'高端单身旅行与线下社群研究报告.docx')

e=new('高端单身旅行与线下社群决策摘要','下一次合伙人会议讨论材料')
p(e,'本次建议先验证北京单身晚宴，三亚只保留一条可报价的旅行初案。商务学习活动单独核算。BoL／BoL Selected已由Olivia确认，现无已确认供应商或客户访谈结论。')
table(e,['产品','研究试价','基准人数','单次贡献','周转金'],[[c['name'],f'{c["price"]:,}元',c['base'],f'{M["results"][c["id"]]["contribution"]:,.0f}元',f'{M["results"][c["id"]]["cash"]:,.0f}元'] for c in D['products']],[5.3,2.5,2.3,3.2,3.1])
p(e,'所有价格与成本均为规划输入，已含工时机会成本，未计公司固定开支和税费。旅行含三晚独住，往返机票自理。分享会约21人才保本，旅行财务保本6人但建议至少12人才满足社交交付。')
h(e,'重点学谁',1)
for q in ['32号、游侠客：线路、领队、报价与退改。32号海南六日5980元起、双住且多数正餐自理，不是我们的同条件价格。','钻石、我主良缘：高资产与婚恋筛选、活动和深度服务分开收费；旅行经营表现尚缺。','BoL／BoL Selected、大超：双业务线组织、内容获客与线下承接；尚无可核的旅行利润。']:p(e,q)
h(e,'本次会议应决定',1)
for q in ['客群：30至42岁北京常住真实单身的范围与资产口径。','线路：三亚四天三晚，是否接受独住及机票另购。','活动：先做晚宴，分享会是否待讲师和低成本场地确认。','价格预算：是否用1680／680／19800做客户访谈试价，最多投入多少验证成本。','分工：Olivia询价和访谈，格格资源；确定反馈与碰头日期。']:p(e,q)
h(e,'决定前还需两类证据',1)
p(e,'取得至少两家相同条件的供应商报价；完成约12位潜客访谈，记录过去真实消费及对具体日期、同行条件和试价的选择。当前同线路价差、实际付费转化及资源优势尚未成立。')
e.save(O/'高端单身旅行与线下社群决策摘要.docx')

g=new('高端单身旅行与线下社群会议讲解稿','建议讲解时长十五分钟')
for title,txt in [
('一 先说建议 三分钟','我们这一轮把方向收到了可测试的产品：北京先做一场小型单身晚宴，确认我们能否把合适的人请来并服务好。三亚先设计一条四天三晚的旅行，不同时开多条线。商务分享会单独组织，允许已婚客户参加，收费和结果也分开。'),
('二 为什么选这些对标 三分钟','32号和游侠客告诉我们怎样交付旅行；钻石和我主良缘帮助拆婚恋筛选；BoL和大超帮助拆品牌入口、内容与线下。这个组合不是“六家最赚钱”的排名。32号海南六日5980元起可以查到具体住宿和不含项，但它是双住私家团，不能拿来直接证明我们的价格便宜或贵。我们仍缺两家同口径报价。'),
('三 客户到底买什么 三分钟','我们需要验证的客户约30至42岁、北京常住、真实单身，有明确稳定关系意愿。不要只问他有多少钱，还要问他最近参加过什么、和谁旅行、什么日期有空、愿不愿独住、喜欢一对一还是小组。产品的钱主要要买到可信的同行者、舒服的组织和足够长的相处。商务活动则买有用的主题和经营同行，不能暗中把已婚商务客当相亲供给。'),
('四 账怎么算 三分钟','晚宴研究试价1680，20人基准贡献9492元；分享680，24人贡献约1590元，说明分享会必须控制固定成本。旅行19800，14人贡献82284元，但这依赖尚未验证的试价。模型含直接交付、获客和项目工时，尚未扣公司日常成本和税费。旅行基准需约75216元周转金，预收款不是已经赚到的钱。'),
('五 今天拍板什么 三分钟','今天确定客群及核验口径、是否以三亚为首条、先做哪两场活动、访谈用什么试价和最高验证预算、双方分工。Olivia随后拿统一问题询价和访谈，格格落实具体资源。拿到反馈后再改输入、确定对外价格及合同，最后安排招募。年费、黑卡和多团并行放到复购之后讨论。')]:h(g,title);p(g,txt)
g.save(O/'高端单身旅行与线下社群会议讲解稿.docx')
print('documents created')
