from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import json

B=Path(__file__).resolve().parents[1];O=B/'交付资料';P=O/'正式文件';P.mkdir(exist_ok=True)
src=B/'formal-rebuilt/中国高端社群市场研究与创业进入评估_具体表达版.docx'
title='中国高端会员社群市场研究及创业可行性评估报告'
d=Document(src);d.paragraphs[0].runs[0].text='中国高端会员社群市场研究\n及创业可行性评估报告';d.core_properties.title=title
d.save(P/(title+'.docx'))

d=Document(src);body=d._element.body;keep=False
for el in list(body):
 text=''.join(t.text or '' for t in el.iter(qn('w:t')))
 if text=='四 六家机构深度研究':keep=True
 if text=='五 横向比较与公司机会':keep=False
 if not keep and el.tag!=qn('w:sectPr'):body.remove(el)
q=d.paragraphs[0];q.text='六家高端会员社群对标案例档案';q.style='Title';q.paragraph_format.page_break_before=False
d.core_properties.title=q.text;d.save(P/'六家高端会员社群对标案例档案.docx')

original=Document(src);d=Document(src)
for el in list(d._element.body):
 if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
def p(t,style=None):return d.add_paragraph(t,style)
def h(t):return d.add_heading(t,1)
def tb(idx):
 el=deepcopy(original.tables[idx]._tbl);d._element.body.insert(len(d._element.body)-1,el);p('')
def link(q,label,url):
 a=OxmlElement('w:hyperlink');a.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 for k,v in [('color','295878'),('sz','19')]:e=OxmlElement('w:'+k);e.set(qn('w:val'),v);pr.append(e)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);q._p.append(a)
p('高端会员社群创业项目\n合伙人决策摘要','Title');p('资料截至2026年10月21日','Subtitle')
p('建议先在一个城市，为一批有明确共同兴趣和参加目的的客户，组织连续见面的收费活动。确认他们愿意再次付费、团队能够稳定交付，再推出年度会员。旅行、游学与企业参访按共同需求单独报价，婚恋服务由单身成员另行自愿申请。')
h('一 六家分别值得学什么')
tb(6)
p('六家提供的服务有所侧重。我们应组合可以自己做好的部分：EO的固定同伴和主持，长江的共同学习，马会的日常使用和家属规则，木兰汇的成员分享与分层票价，正和岛的商业内容，优家荟的会员企业互访和生活体验。')
h('二 目前能够支持决策的事实')
p('收费已经能拆到具体产品。北京马会2026个人入会费60万元、月费2200元；EO全球年费2630美元、首次入会费3500美元，华东地方完整收费仍需确认；长江2026级EMBA学费99.8万元、申请费2000元，服务两年；木兰2026年会普通、VIP、SVIP标准票分别为8800元、17800元和37800元。')
p('正和岛App数字学习卡1598元／年、金卡海邻3998元／年，核心线下会籍的2026完整报价未公开取得。优家荟已取得2017历史白金卡2800元／年、黑金卡18000元／年，以及2025新加坡两日项目250新币／人；这些分别用于理解历史分层和项目收费，现行会籍另取报价。')
p('EO全球FY2024/25续会率91.6%，是本轮较直接的持续付费指标。华东分会2026年9月底有109位会员。长江自述累计培养13000余位EMBA学员；木兰汇自述近1000位成员；优家荟深圳分会自述十年吸纳3000余名女性、举办1000余场活动。各项保留原地区、时期和累计口径。')
p('平台数据用来判断内容与传播。长江EMBA认证小红书账号快照为3304粉丝；正和岛同名账号为66299粉丝，主体仍未充分核实。优家荟官方微博索引为15086粉丝，一条历史艺术旅程视频有283条评论、352次转发。核心会员成交、复购与获客成本需要后续取得。')
h('三 先选择哪一种起步方式')
tb(22)
p('城市选择先看团队实际能触达的人：谁最近买过同类服务、每月能参加几次、愿意与什么人持续见面。选择经营者固定组，需要有合适主持和保密讨论；选择生活体验，需要可靠策划与接待；选择女性方向，需要自己的女性客户网络和持续内容。')
h('四 合伙人本次需要作出的决定')
p('第一，首发服务哪一类人，主要解决他们工作以外的生活社交，还是经营中的同行交流。第二，谁负责找到首批客户，谁负责主持与内容，谁负责现场与供应商。第三，批准客户访谈和三场收费体验的预算、时间及停止条件。年度价格、固定场地和跨城投入留到试点后决定。')
p('价格先从权益和成本倒算：服务多少个月、固定组几次、开放活动几席、能否带伴侣、哪些旅行另收费。分别记录收款、已完成服务的收入、未交付权益和退款。试点阶段重点看实付、实到、第二次付费、每场剩余贡献和团队工时。')
h('五 九十天验证安排')
tb(25)
p('试点开始前确定通过条件。正式报告提出实到80%、30天第二次付费30%、单场贡献为正作为讨论参数；合伙人应结合人数、拒绝原因和实际工时确定最终目标。年度续会须等会员到期后观察。')
h('六 资料入口')
S=json.loads((B/'formal-rebuilt/核查索引.json').read_text())['sources']
for label,ids in [('北京马会',['R01','K01']),('EO',['K05','K06','K07']),('长江EMBA',['K03','R14']),('木兰汇',['R02','R03']),('正和岛',['K09','R11']),('优家荟',['U14','U15','R09'])]:
 q=p(label+'  ')
 for i,k in enumerate(ids):
  if i:q.add_run('  ')
  link(q,S[k]['title'],S[k]['url'])
d.core_properties.title='高端会员社群创业项目合伙人决策摘要';d.save(P/'高端会员社群创业项目合伙人决策摘要.docx')
print('\n'.join(str(p) for p in P.glob('*.docx')))
