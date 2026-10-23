from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import re,json
r=Path(__file__).resolve().parents[1]/'formal-rebuilt';path=r/'中国高端社群市场研究与创业进入评估.docx';d=Document(path)
d.styles['Subtitle'].font.italic=False
changes={
'官网员工总量和平均值出现异常，本报告不使用该字段。':'',
'主要价格差异落在初始入会门槛和人数，并非已经证明甲类享有三倍活动或更高服务等级。':'公司档位按登记人数区分，入会费和月费按整份会籍报价。',
'110除以380约为29%，只是以报道约数计算的潜在会员参加占比。它说明活动具有招募功能；没有加入结果，不能填写成交率。这是北亚区域案例，其中有华东成员参与，不能当作华东本地活动或本地活跃率。':'这场北亚区域活动将学习、参访和会员招募放在同一体验中，其中有华东成员参与。对本公司，值得学习的是让潜在会员直接体验组织的内容与交流方式，并明确记录之后的申请和付款结果。',
'2026-10-21；人数、年龄、全球收费':'2026-10-21采集；人数及年龄统计截至2026-09-30；全球收费',
}
for q in d.paragraphs:
 for run in q.runs:
  for a,b in changes.items():
   if a in run.text:run.text=run.text.replace(a,b)
for table in d.tables:
 for i,row in enumerate(table.rows):
  if not i:continue
  for c in row.cells:
   if re.fullmatch(r'\d+(?:\.\d+)?(?:人)?',c.text.strip()):
    for q in c.paragraphs:q.alignment=2
# Add direct note locators to make content metrics independently traceable.
links=[('U11a','长江EMBA代表笔记','https://www.xiaohongshu.com/explore/6aa7c749000000002700b528','2026-09-14发布；本项目2026-10-21快照，2评论、10分享'),('U11b','正和岛同名账号代表笔记','https://www.xiaohongshu.com/explore/6ad60b8b0000000027014c39','2026-10-20发布；本项目2026-10-21快照，2评论、1分享；主体未充分核验')]
for id,title,url,scope in links:
 q=d.add_paragraph();q.paragraph_format.line_spacing=Pt(14);q.paragraph_format.space_after=Pt(5);q.add_run(id+' ').bold=True
 el=OxmlElement('w:hyperlink');el.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True));rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz);col=OxmlElement('w:color');col.set(qn('w:val'),'295878');pr.append(col);rr.append(pr);tx=OxmlElement('w:t');tx.text=title;rr.append(tx);el.append(rr);q._p.append(el);q.add_run('\n'+scope)
 for run in q.runs:run.font.size=Pt(9.5)
d.save(path)
a=json.loads((r/'核查索引.json').read_text());a['sources']['K05']['scope']='2026-10-21采集；人数年龄截至2026-09-30；全球收费'
for id,title,url,scope in links:a['sources'][id]={'title':title,'url':url,'scope':scope}
(r/'核查索引.json').write_text(json.dumps(a,ensure_ascii=False,indent=2))
print('Finalized',path)
