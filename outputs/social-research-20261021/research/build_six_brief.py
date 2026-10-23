import json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]/'community-supplement'
d=json.loads((ROOT/'curated-data.json').read_text());ts={t['name']:t for t in d['tables']}
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.3)
def font(style,size,bold=False,head=False):
 style.font.name='Arial' if head else 'Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 rf=style._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'Heiti SC' if head else 'Songti SC')
 for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme']:rf.attrib.pop(qn('w:'+k),None)
for name,size in [('Normal',12),('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
 st=doc.styles[name];font(st,size,name!='Normal',name!='Normal');f=st.paragraph_format
 f.line_spacing=Pt(20 if name=='Normal' else size+7);f.space_after=Pt(5 if name=='Normal' else 7);f.space_before=Pt(12 if name.startswith('Heading') else 0)
 f.keep_with_next=name!='Normal';f.widow_control=True
 for b in list(st._element.iter(qn('w:pBdr'))):b.getparent().remove(b)
def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,1)
def para(label,text):
 q=p('');q.add_run(label+' ').bold=True;q.add_run(text);return q
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 trpr=t.rows[0]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:tblHeader'))
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v if v is not None else '待补')
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side,val in [('top',85),('bottom',85),('left',95),('right',95)]:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');m.append(x)
   tcpr.append(m)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1' if i==0 else 'F7F9FA' if i%2==0 else 'FFFFFF');tcpr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=Pt(16);q.paragraph_format.keep_with_next=False
    for r in q.runs:
     r.font.size=Pt(10.5);r.font.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
   if i:
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 p('')
 return t
def link(q,text,url):
 el=OxmlElement('w:hyperlink');el.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'224F71');pr.append(color);r.append(pr);v=OxmlElement('w:t');v.text=text;r.append(v);el.append(r);q._p.append(el)

ROOT=Path(__file__).resolve().parents[1]/'six-community-review'
d=json.loads((ROOT/'六家筛选依据及对标数据.json').read_text());sources={r[0]:r for r in d['sources']}
def refs(ids):
 q=p('资料依据 ')
 q.paragraph_format.line_spacing=Pt(16);q.paragraph_format.space_after=Pt(0)
 for i,key in enumerate(ids):
  if i:q.add_run('；')
  link(q,key+' '+sources[key][1],sources[key][2])
 for r in q.runs:r.font.size=Pt(10)
def newpage(title):
 doc.add_page_break();h(title)
p('六家高端社群筛选与对标分析','Title')
p('合伙人讨论简报　2026年10月21日')
p('建议将北京香港马会会所、长江商学院、EO中国华东分会、木兰汇、正和岛、阿拉善SEE列为六家重点学习对象。六家分别提供生活方式与家庭、客户筛选、固定小组、女性参与、分层商业连接和长期会员投入的参考。')
para('选择依据','机构及品牌身份、持续会员关系、目标客群关联、独特经营能力和可核对的规模或经营证据。此次不设置主观评分或流量总榜。')
table(['入选机构','最值得学习的能力','入选判断'],[[c['name'],c['learn'],c['reason']] for c in d['cases']],[3.4,4.2,8.6])
para('名单调整','原六家中的京城俱乐部、观澜湖转为专项参考；正和岛与阿拉善SEE进入六家。长江继续保留。')

newpage('一 平台声量与经营证据')
p('小红书数据采集于2026年10月21日。评论与分享列使用一条明确标识的代表内容；账号粉丝、内容互动和会员经营各自列示。')
rows=[
 ['北京香港马会会所','待核','待核','待核','个人入会费60万元；2025会刊记载三类实际文化活动'],
 ['长江商学院','3304','2','10','EMBA累计13000余人；99%以上为企业核心决策层'],
 ['EO中国华东分会','待核','待核','待核','华东109会员；全球2024／2025续会91.6%'],
 ['木兰汇','待核','待核','待核','近1000成员、29省市、112位核心成员企业市／估值20亿元以上'],
 ['正和岛同名账号','66299*','2','1','开发者自述深度服务亿元营收以上客户企业9000余家'],
 ['阿拉善SEE','待核','待核','待核','2025有效会员475、续费70%、会费收入2280万元']]
table(['机构','小红书粉丝','评论','分享','规模及经营证据'],rows,[3.8,2,1.4,1.4,7.6])
p('代表内容：长江为2026年9月14日非洲出海校友案例，17赞、2评论、10分享；正和岛为10月20日组织用人内容，9赞、2评论、1分享。正和岛同名号的运营主体尚待交叉核验，标记为*。其他四家已完成关键词结果及官网核对，尚无经核验主号指标。')
para('平台判断','现有小红书证据支持拆解内容主题与咨询入口，尚未形成六家可比的90天完整互动样本。正和岛粉丝较多，但本轮可见近期互动较少，因此此次入选的理由是组织与产品体系。')
para('经营判断','EO全球续会率和SEE会员经营数据，对持续付费与服务设计的参考更直接。木兰汇、长江及正和岛的规模资料，主要用于判断客群和组织基础。')
refs(['K04','K05','K07','K08','K09','K11'])

for idx,c in enumerate(d['cases'],1):
 newpage(str(idx+1)+' '+c['name'])
 para('学习重点',c['learn']+'。'+c['reason'])
 for label,key in [('经营与规模','performance'),('价格与收费','price'),('会员权益','rights'),('客户画像','customer'),('获客路径','acquisition'),('活动与交付','flow'),('研究判断','judgment'),('建议采用的做法','transfer')]:para(label,c[key])
 refs(c['sources'])

newpage('八 对拟设公司的建议')
para('先建立稳定的社群关系','优先借鉴EO的固定小组、主持训练和月度节奏，再用马会的文化与生活活动拓展成员之间的连接。首发产品需要说清楚成员和谁长期相处、多久见面、组织者提供什么服务。')
para('按参与动机组织客户','企业家同行、独立女性成长、伴侣家庭生活和单身关系探索，应有各自清晰的活动入口。成员可以跨组参加共同活动，单身匹配采用单独的报名和交付规则。')
para('让六项服务形成先后关系','私董会和固定社交小组承担持续关系；文化及兴趣活动增加见面频次；旅行、游学和企业参访围绕具体主题按项目交付；婚恋服务在单身客户中单独验证需求与支付。')
para('收费按交付拆分','分别设计年度服务、单次体验和高成本项目。长江教育项目、马会入会资格、EO年度全球会费与正和岛数字订阅，收费单位和交付均不同；报价模型应基于拟设公司的实际服务量、成本和愿付证据。')
para('把会员贡献做成服务的一部分','参考SEE的项目参与、企业互访、组织治理和嘉许，让会员能发起、能分享、能共同策划。组织者负责质量、规则与持续跟进。')
para('经营看板','至少跟踪申请人数、面谈通过、首次付费、实到、90天重复参与、到期续费及推荐成交。每项记录统计期和分母；内容粉丝及互动作为获客前端指标。')
h('专项参考名单')
for r in d['rejected']:para(r[0],r[1]+'。'+r[2]+'。')

newpage('附录 统计口径与来源')
para('小红书样本','长江EMBA认证账号采集24条，其中23条在2026年7月23日至10月21日。可见评论值4条、合计5；可见分享值8条、合计33。正和岛同名号近90天9条，可见评论值1条、合计2；可见分享值3条、合计5。缺项保持未知。代表帖不作为账号累计或平均值。')
para('会员与时间','EO华东人数为2026年10月21日，91.6%续会率为2024／2025全球数据。SEE 2025经营数据与2024上半年活动资料分开列示。机构累计服务和成员规模保留原披露口径。')
for row in d['sources']:
 q=p('');q.paragraph_format.line_spacing=Pt(16);q.paragraph_format.space_after=Pt(5)
 q.add_run(row[0]+' ').bold=True;link(q,row[1],row[2]);q.add_run('　'+row[3])
 for r in q.runs:r.font.size=Pt(10.5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.author='';doc.core_properties.title='六家高端社群筛选与对标分析'
out=ROOT/'六家高端社群筛选与对标分析.docx';doc.save(out);print(out)
