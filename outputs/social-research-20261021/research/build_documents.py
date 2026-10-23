from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from report_content import sections,summary

ROOT=Path(__file__).resolve().parent.parent
D=json.loads((ROOT/'research/data.json').read_text())
SOURCES={s['id']:s for s in D['sources']}
FONT='Hiragino Sans GB'
NAVY='233D55'

def font(run,size=11.5,bold=False,color='111111'):
    run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)

def base(short):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.70);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
    sec.header_distance=Inches(.30);sec.footer_distance=Inches(.30)
    for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        s=doc.styles[st];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    for s in doc.styles:
        for border in list(s._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11.5)
    pf=doc.styles['Normal'].paragraph_format;pf.line_spacing=1.25;pf.space_after=Pt(7);pf.widow_control=True
    for n,size in [('Title',24),('Heading 1',17),('Heading 2',13)]:
        st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True
        st.paragraph_format.space_before=Pt(15 if n!='Title' else 0);st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(h.add_run(short+'　研究截止 2026年10月21日'),8.5,color='596573')
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(f.add_run('创业决策案头研究　|　'),8.5,color='596573')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    doc.core_properties.title=short;doc.core_properties.subject='中国高端社交创业决策研究';doc.core_properties.author='研究整理'
    return doc

def p(doc,text,size=None,gray=False):
    para=doc.add_paragraph();font(para.add_run(str(text)),size or 11.5,color='596573' if gray else '111111');return para

def heading(doc,text,level=1):
    # Word headings deliberately use black text and no decorative rules.
    return doc.add_heading(re.sub(r'[：:—\/、]',' ',text),level)

def link(para,label,url):
    rid=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid)
    run=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'254E76');rp.append(c)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');rp.append(sz)
    rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Arial');rf.set(qn('w:eastAsia'),FONT);rp.append(rf)
    run.append(rp);t=OxmlElement('w:t');t.text=label;run.append(t);el.append(run);para._p.append(el)

def refs(doc,text):
    pp=p(doc,'证据索引　'+text,9,True);pp.paragraph_format.space_after=Pt(12)

def table(doc,headers,rows,widths=None):
    if not rows:return
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or ([1.35,2.65,2.94] if len(headers)==3 else [6.94/len(headers)]*len(headers))
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for j,h in enumerate(headers):t.rows[0].cells[j].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):cells[j].text='未公开核实' if v is None else str(v)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b)
            margins=OxmlElement('w:tcMar')
            for side,val in [('top','85'),('bottom','85'),('left','95'),('right','95')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),NAVY if i==0 else ('F3F6F8' if i%2==0 else 'FFFFFF'));pr.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=1.15
                if i==0:pp.paragraph_format.keep_with_next=True
                if len(rows)<=4:pp.paragraph_format.keep_with_next=i<len(t.rows)-1
                if j==0 and widths[j]<1.1:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in pp.runs:font(r,10.5,bold=i==0,color='FFFFFF' if i==0 else '111111')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def title(doc,text,intro):
    doc.add_paragraph(text,'Title');p(doc,'中国大陆　北京 上海 深圳重点观察　|　2026年10月21日',9.5,True)
    p(doc,intro)

def appendix(doc,ids=None,compact=False):
    heading(doc,'来源索引')
    p(doc,'源ID与附件工作簿一致。点击标题可访问原始页面。统计期与采集日分开；搜索索引、转载和机构自述的限制保留，来源数量不代表独立样本数。完整发现及访问说明见工作簿来源台账。',10.5)
    for s in D['sources']:
        if ids is not None and s['id'] not in ids:continue
        pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(3 if compact else 5);pp.paragraph_format.line_spacing=1.12
        font(pp.add_run(s['id']+'　'),9.5,True);link(pp,s['title'],s['url'])
        font(pp.add_run('\n'+s['publisher']+'　|　'+s['period']+'　|　'+s['kind']),9,color='596573')
        # Relevant limits stay attached to the source; detailed audit remains in the workbook.
        font(pp.add_run('\n'+s['limit']),9,color='596573')

report=base('中国社交婚恋与高端会员服务市场研究')
title(report,'中国社交婚恋与高端会员服务市场研究',
      '本报告供创始人与合伙人讨论公司进入方向、首发产品和验证投入。研究从全国社交婚恋背景，收敛到高端客群、相邻服务与三种经营模式。当前建议优先验证城市兴趣社交体验；年费会员、三城扩张和复杂产品组合仍需真实付费与交付证据。')
p(report,'交付范围：41个机构或品牌体系、15家结构化比较、三家逐字段案例；价格、产品、客户、获客、流程、用户及活跃均有独立表。内容和活动采集未全部达到计划目标，相关缺口在正文及附件明确披露。',10.5)
heading(report,'阅读导航')
p(report,'第1至2节说明结论和方法；第3至9节为宏观与市场地图；第10至16节为必查模块；第17至19节为三家案例；第20至25节比较经济性、进入条件、公司方向与验收结果。建议合伙人先读决策摘要，再用工作簿核查报价与关键假设。',10.5)
for i,s in enumerate(sections,1):
    heading(report,f'{i:02d}　{s["title"]}')
    for text in s['paras']:p(report,text)
    if s.get('table'):
        n=len(s['headers']);w={2:[1.9,5.04],3:[1.5,2.45,2.99],4:[1.25,1.85,1.75,2.09]}.get(n)
        table(report,s['headers'],s['table'],w)
    if s['refs']:refs(report,s['refs'])
heading(report,'附件使用与验收')
table(report,['交付件','内容','当前边界'],[
 ['正式报告','行业与模式分析 方向条件 证据索引','全国高端综合社交无统一统计口径'],
 ['合伙人决策摘要','建议选择 分歧 待决定问题','只支持下一阶段验证'],
 ['三家案例档案','每家A—F六模块102字段与两条活动观察','未知保留；活动并非均为实际完整记录'],
 ['数据工作簿','17个工作表 含价格来源缺口与可调模型','20篇逐篇内容只覆盖耀悦'],
 ['后续验证清单','工作簿证据缺口表与正文验证章节','不外联不购买；未执行下一阶段活动']
])
appendix(report)
report.save(ROOT/'中国高端社交市场研究正式报告.docx')

memo=base('合伙人决策摘要')
title(memo,'高端社交公司合伙人决策摘要','本次需要决定的是下一阶段验证方向与投入条件。现有资料支持相邻付费需求存在，尚不足以证明综合年费社交模式能够持续盈利。建议优先验证城市兴趣社交体验，并用真实重复支付检验会籍价值。')
for h,text in summary:
    heading(memo,h,2);p(memo,text)
p(memo,'会谈记录：首发城市＿＿＿　首发产品＿＿＿　验证预算＿＿＿　客户负责人＿＿＿　交付负责人＿＿＿　扩大条件＿＿＿　停止条件＿＿＿',10.5)
refs(memo,'完整链接与证据限制见正式报告来源索引及工作簿来源台账；所有模型输入均为研究假设')
memo.save(ROOT/'高端社交公司合伙人决策摘要.docx')

cases=base('三家深度案例与逐字段核验')
title(cases,'三家高端服务模式案例档案','本档案对京城俱乐部、伟事达中国、耀悦逐字段记录已知证据和未知项，用于比较经营机制及复制条件。每家包含价格、产品、客户、获客、活动、用户活跃六模块，共306条字段。没有取得公开资料的项保留未知，不视作不存在或经营为零。')
p(cases,'每个字段后的源ID指向本档案末尾及工作簿来源台账。查找范围包括官网、公开产品、近期活动和平台线索；缺失的具体证据由各模块问题、活动表及证据缺口表说明。公开标价与定位不能替代真实成交和客户抽样。',10.5)
labels={'价格':'A 价格与权益','产品':'B 产品与收费模式','客户':'C 客户定位与实际画像','获客':'D 获客与成交路径','活动':'E 活动及交付流程','用户活跃':'F 用户量与活跃'}
case_src=set()
for k,name in enumerate(['京城俱乐部','伟事达中国','耀悦']):
    hh=heading(cases,f'案例 {k+1}　{name}')
    if k:hh.paragraph_format.page_break_before=True
    sec=next(s for s in sections if s['title']==f'案例{["一","二","三"][k]} {name}')
    for text in sec['paras']:p(cases,text)
    rows=[r for r in D['fields'] if r[0]==name]
    known=sum(r[5]!='未知' for r in rows)
    p(cases,f'核验统计：102个字段中{known}项取得部分证据，{102-known}项仍未知。部分证据常只覆盖产品标准、品牌定位或个别活动，不表示完整核实。',10.5,True)
    for mod,label in labels.items():
        heading(cases,label,2)
        rr=[r for r in rows if r[1]==mod]
        if mod=='客户':p(cases,'本表主要反映机构希望吸引的人；实际客户占比、资产、预算、购买触发和留存动机尚无可靠抽样证据。',10.5)
        if mod=='用户活跃':p(cases,'粉丝不是会员；累计人数不是当前人数；没有观察期、有效会员分母及行为记录，不计算活跃或续费率。',10.5)
        table(cases,['字段','观察与证据边界','源ID'],[[r[2],r[3],r[4] if r[5]!='未知' else r[4]+' 已查未得'] for r in rr],[1.45,4.3,1.19])
        for r in rr:case_src.update(re.findall(r'(?:S\d{2}|X\d{2}(?:-\d{2})?)',r[4]))
    heading(cases,'两条代表性活动观察',2)
    ee=[e for e in D['events'] if e[1]==name]
    for e in ee:
        p(cases,e[0]+'　'+e[2]+'　'+e[3],11,False)
        p(cases,'记录类型：'+e[5]+'。'+e[12]+'。目标人群：'+e[6]+'。',10.5)
        table(cases,['流程阶段','已知内容或未知'],[[f[3],f[4]] for f in D['flows'] if f[0]==e[0]],[1.45,5.49]);refs(cases,e[11]);case_src.add(e[11])
    heading(cases,'下一阶段最需要补的证据',2)
    needed={
      '京城俱乐部':'现行会籍书面报价、包含及另收费清单、家庭及宾客规则、两场近期实际会员活动议程、有效付费与续费定义。',
      '伟事达中国':'CE SE KE实际合同价及退改；本地可加入组与教练配置；两场组会或游学真实议程；中国区域有效会员、续费和服务产能。',
      '耀悦':'具体团期费用清单与退改；四人私享是否自组同行；异常报价书面核实；真实出团人数、复购与社交筛选是否单独收费。'}
    p(cases,needed[name]);p(cases,'本轮未提交咨询或外联。上述为验证清单，不是已取得结果。',10,True)
appendix(cases,case_src|{'X03'}|{f'X03-{j:02}' for j in range(1,21)},compact=True)
cases.save(ROOT/'三家高端服务模式深度案例档案.docx')
print(json.dumps({'sections':len(sections),'sources':len(SOURCES),'case_fields':len(D['fields']),'files':[p.name for p in ROOT.glob('*.docx')]},ensure_ascii=False))
