import json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]/'community-supplement'
d=json.loads((ROOT/'curated-data.json').read_text());ts={t['name']:t for t in d['tables']}
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.3)
def font(style,size,bold=False,head=False):
 style.font.name='Arial' if head else 'Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 rf=style._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'Heiti SC' if head else 'Songti SC')
 for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme']:rf.attrib.pop(qn('w:'+k),None)
for name,size in [('Normal',12),('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
 st=doc.styles[name];font(st,size,name!='Normal',name!='Normal');f=st.paragraph_format
 f.line_spacing=Pt(20 if name=='Normal' else size+7);f.space_after=Pt(5 if name=='Normal' else 7);f.space_before=Pt(12 if name.startswith('Heading') else 0)
 f.keep_with_next=name!='Normal';f.widow_control=True
 for b in list(st._element.iter(qn('w:pBdr'))):b.getparent().remove(b)
def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):return doc.add_heading(text,1)
def para(label,text):
 q=p('');q.add_run(label+' ').bold=True;q.add_run(text);return q
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 trpr=t.rows[0]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:tblHeader'))
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v if v is not None else '待补')
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side,val in [('top',85),('bottom',85),('left',95),('right',95)]:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');m.append(x)
   tcpr.append(m)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1' if i==0 else 'F7F9FA' if i%2==0 else 'FFFFFF');tcpr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcpr.append(borders)
   for q in c.paragraphs:
    q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=Pt(16);q.paragraph_format.keep_with_next=False
    for r in q.runs:
     r.font.size=Pt(10.5);r.font.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
   if i:
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 p('')
 return t
def link(q,text,url):
 el=OxmlElement('w:hyperlink');el.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'224F71');pr.append(color);r.append(pr);v=OxmlElement('w:t');v.text=text;r.append(v);el.append(r);q._p.append(el)

p('中国高端社群补充调研','Title')
p('数据更新 '+d['generated_at'][:16].replace('T',' '))
p('本轮比较综合会员俱乐部、学习与校友组织、企业家小组、女性社群和高尔夫会员组织，重点补充收费、权益、人群、获客、活动交付和活跃度。机构完整对照、逐条内容数据及来源见配套工作簿。')
h('一 主要发现')
for title,body in d['findings']:para(title,body)
para('与客群设计直接相关的做法','EO全球Forum除了企业家小组，也设配偶及生活伴侣小组。京城和香港马会提供家庭活动，深商按女性及年轻企业家分设组织。对拟设公司而言，可据此比较共同活动与分组服务如何搭配。来源 C10 C29 C30。')
h('二 机构比较')
table(['机构与层级','公开资料与研究价值'],[[r[0]+'\n'+r[1],r[3]+'。'+r[4]+'。来源 '+r[5]] for r in ts['机构与研究重点']['rows']],[4,12.2])
h('三 价格与会员权益')
pr=[r for r in ts['价格与收费']['rows'] if r[3] is not None]
table(['机构与产品','金额及单位','适用时间及来源'],[[r[0]+'\n'+r[2],f'{r[3]:,} '+('元' if r[4]=='CNY' else '美元')+'\n'+r[5],r[6]+'\n'+r[10]] for r in pr],[6.3,4.2,5.7])
para('长江及中欧','长江EMBA公开费用覆盖课程、教材讲义、课间茶点及学院资源。中欧GEMBA还列明各模块午餐茶歇、三次领导力住读食宿；HEMBA列海外部分地面费用，海外机票和住宿另付。各项目独立列示报名费。来源 C01 C02 C04 C05。')
para('香港马会','会刊列个人及三类公司会籍的入会费。子女附属卡为8—21周岁设置，年费120元；12—21岁可申请指定场所签账权。主卡会员承担相关消费。来源 C07 C08。')
para('EO与侨商','EO公开全球年度会费和一次入会费，地方分会收费另列。深圳侨商按会员和组织职务公布不同年度会费。来源 C13 C22。')
para('会费之外的持续参与','观澜湖将10场平日赛事组成套餐，部分已有会籍的会员报名价为4000元，嘉宾为13300元，包含5晚酒店的嘉宾套餐为16300元，并配合消费积分及生日礼遇。黑骑士提供四类会籍。来源 C11 C20 C21 C34。')
h('四 客户与获客方式')
for brand in ['京城俱乐部','北京香港马会会所','长江商学院','EO中国华东分会','木兰汇']:
 profile=next(r for r in ts['客户画像']['rows'] if r[0]==brand)
 acq=next(r for r in ts['获客与转化']['rows'] if r[0]==brand)
 doc.add_heading(brand,2)
 para('客户',profile[2]+'。'+profile[3]+'。来源 '+profile[6]+'。')
 para('获客与成交',acq[1]+'。'+acq[2]+'。'+acq[3]+'。来源 '+acq[6]+'。')
doc.add_heading('观澜湖及中欧的补充做法',2)
p('观澜湖通过球赛、匹克球、雪茄试驾、音乐与家庭活动拓展会员见面场景。中欧MBA按医疗、科技及海归客群分别发布约20席的小聚，每场2.5小时。用户留言或私信报名，由课程部通过邮件确认；内容由校友和在读学生分享。来源 C20 C21 C31 C32 C33。')
h('五 活动交付方式')
for r in ts['活动流程']['rows']:
 if r[1] not in ['香茗雅韵·春季下午茶沙龙','征戈会六周年暨第三届班际驿传','2026第十八届商界木兰年会','年度私人董事会小组','Forum入门培训及月度小组','2026高尔夫马拉松']:continue
 doc.add_heading(r[0].replace('EO全球Forum','EO全球')+' '+r[1],2)
 intro=p(r[2]+'，'+r[3]+'。'+r[4]+'。');intro.paragraph_format.keep_with_next=True
 flow=para('环节',r[5]+'。');flow.paragraph_format.keep_with_next=True
 para('分工',r[6]+'。来源 '+r[8]+'。')
h('六 规模与平台内容')
ms=[r for r in ts['规模与活跃度']['rows'] if r[2] is not None and r[1]!='成立年份']
def metric(r):return f'{r[2]*100:.1f}%'+('以上' if r[3]=='以上' else '') if r[3] in ['比例','以上'] else f'{r[2]:,} {r[3]}'
table(['机构及指标','数值','统计范围与时间'],[[r[0]+'\n'+r[1],metric(r),r[4]+'。'+r[6]] for r in ms],[6.2,3,7])
accounts=[r for r in ts['平台机构账号']['rows'] if '主页蓝色认证标识可见' in r[4]]
table(['小红书机构账号','粉丝','累计获赞与收藏'],[[r[1],r[2],r[3]] for r in accounts],[8.2,4,4])
p('两家账号认证标识均已核对，各采集24条主页内容，其中各23条在近90天内。剔除中欧1条本科招生内容后，共45条近期相关内容。明细保留发布时间、采集时间、点赞、收藏、评论、分享及公开链接，按同一账号样本内点赞排序。')
h('七 后续需要取得的经营数据')
for r in ts['待补数据']['rows']:
 para(r[0],r[2]+'。'+r[4]+'。')
h('附录 来源')
for r in ts['来源台账']['rows']:
 q=p('');q.paragraph_format.line_spacing=Pt(17);q.paragraph_format.space_after=Pt(8)
 run=q.add_run(r[0]+' '+r[1]+' ');run.bold=True;run.font.size=Pt(10.5)
 link(q,r[2],r[3]);q.add_run('。'+r[4]+'。'+r[7]+'。').font.size=Pt(10.5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.author='';doc.core_properties.title='中国高端社群补充调研'
out=ROOT/'中国高端社群补充调研.docx';doc.save(out);print(out)
