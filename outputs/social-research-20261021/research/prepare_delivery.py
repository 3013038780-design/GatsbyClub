from pathlib import Path
import json,csv,re,hashlib,shutil
from collections import OrderedDict
from docx import Document
from docx.oxml.ns import qn

B=Path(__file__).resolve().parents[1];O=B/'交付资料';O.mkdir(exist_ok=True)
T=O/'数据明细';T.mkdir(exist_ok=True)
load=lambda p:json.loads(p.read_text())
base=load(B/'research/data.json');sup=load(B/'community-supplement/curated-data.json');latest=load(B/'formal-rebuilt/核查索引.json');integrated=load(B/'formal-integrated/研究与模型核查.json')
tables={t['name']:t for t in sup['tables']};report=Document(B/'formal-rebuilt/中国高端社群市场研究与创业进入评估_具体表达版.docx')
sources=[]
for s in base['sources']:sources.append([s['id'],s['title'],s['url'],s.get('publisher',''),s.get('period',''),s.get('kind',''),s.get('finding',''),s.get('limit',''),s.get('collected','')])
for r in tables['来源台账']['rows']:sources.append([r[0],r[2],r[3],r[1],r[4],r[6],r[7],'',r[8]])
for k,s in latest['sources'].items():sources.append([k,s['title'],s['url'],'',s['scope'],'正式报告采用',s['scope'],'','2026-10-21'])
sm={r[0]:r for r in sources}
url=lambda keys:'\n'.join(dict.fromkeys(sm[k][2] for k in re.findall(r'[A-Z]+\d+[a-z]?',str(keys)) if k in sm))
sheets=[]
def add(n,h,r,note='',widths=None):
    sheets.append(dict(name=n,headers=h,rows=r,note=note,widths=widths))
def st(n,h,r,note=''):add(n,h,r,note)
names=['北京香港马会会所','EO中国华东分会','长江商学院','木兰汇','正和岛','优家荟']
cases=[];cur=None;sub='';
for q in report.paragraphs:
 if re.match(r'4\.[1-6] ',q.text):
  cur={'name':q.text.split(' ',1)[1],'sections':OrderedDict(),'links':[]};cases.append(cur);sub=''
 elif q.text.startswith('五 '):cur=None
 elif cur:
  if q.style.name=='Heading 3':sub=q.text;cur['sections'][sub]=[]
  elif q.text.startswith('打开资料'):
   for a in q._p.findall(qn('w:hyperlink')):
    cur['links'].append({'label':''.join(a.itertext()).split('打开资料')[-1], 'url':q.part.rels[a.get(qn('r:id'))].target_ref})
  elif q.text and not q.text.startswith('依据'):cur['sections'].setdefault(sub,[]).append(q.text)
# Read report tables rather than outdated comparison records.
comparison=[[c.text for c in r.cells] for r in report.tables[6].rows]
add('六家对照',comparison[0],comparison[1:],'资料截至2026年10月21日。按机构服务与经营事实比较，未设置综合排名。',[24,64,70])
price=[]
def priceadd(n,product,amount,currency,unit,period,rights,extras,keys,status='正式报告采用'):
 price.append([n,product,amount,currency,unit,period,rights,extras,keys,url(keys),status])
for typ,num,fee,monthly in [('个人',1,600000,2200),('公司丙类C',1,700000,2200),('公司乙类B',2,1300000,4400),('公司甲类A',3,1800000,6600)]:
 priceadd(names[0],typ+'入会费',fee,'CNY','整份会籍／一次','2026-01-01起',f'{num}位登记使用人；设施依规则使用',f'月费{monthly}元；餐饮住宿及另付活动另外消费','R01')
 priceadd(names[0],typ+'月费',monthly,'CNY','整份会籍／月','2026-01-01起',f'{num}位登记使用人','另收入会费及消费','R01')
priceadd(names[0],'子女附属卡',120,'CNY','人／年','现行公开表单','8至21岁，依年龄适用使用与签账规则','主会籍另计','K02')
for pr,amt,unit in [('全球年费',2630,'人／年'),('全球首次入会费',3500,'人／一次')]:priceadd(names[1],pr,amt,'USD',unit,'2026-10-21取得公开页','全球会员组织；Forum等依规则参加','地方会费及另付项目另计','K05 K15')
priceadd(names[1],'华东地方会费',None,'CNY','人／年','2026-10-21','需取得地方当期完整报价','Fact Sheet与资格页未共同确认完整现价','K05 K15','未知')
priceadd('EO北京分会','地方年费对照',9200,'CNY','人／年','2026-07-01至2027-06-30','仅为北京地方收费','全球费另计；不替代华东报价','R05')
priceadd(names[2],'2026级EMBA学费',998000,'CNY','人／两年完整项目','2026级','课程、教材讲义、课程茶点、图书馆及网络设施','申请费及具体项目的旅行支出另列','K03')
priceadd(names[2],'申请费',2000,'CNY','人／次','2026级','申请及笔面试流程','学费另计','K03')
for prod,amt,rights in [('普通',8800,'会议简餐，无姓名桌卡'),('普通早鸟',5800,'3月31日前；普通票权益'),('VIP',17800,'3次自助餐、前排课桌姓名卡、赠普通票1张、感言视频及合影'),('VIP早鸟',14800,'3月31日前；VIP权益'),('SVIP',37800,'嘉宾桌餐1次、自助餐2次、前排沙发、赠普通票2张、正装照、感言视频及合影')]:priceadd(names[3],'2026安吉年会'+prod,amt,'CNY','人／单场年会','2026年4月','年会票；'+rights,'年度会籍另列','R02')
for prod,amt,unit in [('数字学习卡',1598,'人／年'),('金卡海邻',3998,'人／年'),('直播',198,'人／场')]:priceadd(names[4],prod,amt,'CNY',unit,'2026-10-21 App内购页','数字内容与学习产品','核心线下会籍另列','K09')
priceadd(names[4],'核心绿色岛邻',30000,'CNY','人／年','2015历史制度','当年会员活动及组织参与权','不作2026报价','K10','历史价格')
for prod,amt,unit,period,rights,key,curcy in [('白金卡',2800,'人／年','2017','历史会员权益','R06','CNY'),('黑金卡',18000,'人／年','2017','历史分层权益及品牌专场等','R07','CNY'),('新加坡两日周年活动',250,'人／项目','2025-11-14至15','白金黑金；计划50席；参访论坛晚宴','U15','SGD'),('新加坡会员访京沪',500,'人／六日项目','2023','京沪会员交流；历史招募','R08','SGD')]:priceadd(names[5],prod,amt,curcy,unit,period,rights,'机酒、交通或餐费等按各项目明细另计',key,'历史价格或项目')
add('六家价格', ['机构','收费项目','金额','币种','收费单位','适用时间','包含权益','另付与限制','来源编号','原文网址','资料状态'],price,'金额按原币保存。未知金额留空；年费、入会费、课程费、活动票分别记录。',[24,27,14,10,24,26,66,60,18,52,24])
for name,selector in [('会员权益',lambda k:'权益' in k or '产品' in k),('客户画像',lambda k:'这家做什么' in k or '画像' in k),('获客方式',lambda k:'获客' in k),('活动流程',lambda k:'活动' in k or '项目执行' in k),('经营表现',lambda k:'经营表现' in k)]:
 rows=[]
 for c in cases:
  for h,ps in c['sections'].items():
   if selector(h):
    for j,p in enumerate(ps):rows.append([c['name'],h,j+1,p,'\n'.join(dict.fromkeys(z['url'] for z in c['links']))])
 add(name,['机构','内容模块','段落序号','具体记录与分析','资料入口'],rows,'与最新正式报告一致。事实、历史制度、方案与分析判断在原文中分别说明。',[24,32,12,110,65])
steps=[]
for idx,n,activity in [(8,names[0],'NEWBY茶沙龙'),(10,names[1],'北亚北京企业家节'),(12,names[2],'迈邦生物参访'),(14,names[3],'杭州PingPong参访'),(16,names[4],'安徽元琛科技参访'),(18,names[5],'新加坡周年活动公布方案')]:
 t=report.tables[idx]
 for j,r in enumerate(t.rows[1:]):steps.append([n,activity,j+1]+[c.text for c in r.cells])
add('活动环节明细',['机构','活动','顺序','环节','实际内容或公布安排','人员分工或组织作用'],steps,'新加坡案例为已公布方案，其余按活动回顾；完整日期、条件及出处见活动流程表。',[24,32,10,24,75,65])
macro=[r+[url(r[5])] for r in base['macro']]
add('宏观市场',['指标','年份','数值','单位','地区','来源编号','资料性质','统计说明','原文网址'],macro,'不同市场与人群口径独立使用。',[30,12,15,12,16,14,24,65,60])
mp=integrated['market_map'];headers=['id','name','category','city','model','target','source','role','status','tier']
add('市场机构地图',['机构编号','机构','类别','城市','经营方式','目标人群','来源编号','研究关系','资料状态','研究层级','原文网址'],[[r.get(k,'') for k in headers]+[url(r.get('source',''))] for r in mp],'55家机构与品牌线索，含替代服务、旅行供应方和待核新候选；六家深度对标见首表。')
for old,new in [('价格与收费','扩展历史价格'),('平台内容明细','平台内容明细'),('平台机构账号','平台账号'),('扩展内容线索','扩展内容线索'),('扩展账号初筛','扩展账号初筛'),('扩展采集覆盖','平台采样覆盖'),('抖音已观察内容','抖音观察')]:
 t=tables[old];rows=t['rows'];hs=t['headers']
 if '源ID' in hs:
  ind=hs.index('源ID');rows=[r+[url(r[ind])] for r in rows];hs=hs+['原文网址']
 add(new,hs,rows,'扩展记录保留采集时的内容和核验状态。最新六家结论以本工作簿前部及正式报告为准。')
page=[];ev=[];links=[];manifest=[]
R=O/'公开网页存档';R.mkdir(exist_ok=True)
for path in sorted(B.glob('collection-*/collection.json')):
 x=load(path);batch=path.parent.name
 for r in x.get('pages',[]):
  rel=r.get('原文文件','');stored=''
  src=path.parent/rel
  if rel and src.is_file() and src.suffix.lower() in ['.md','.txt','.pdf']:
   dest=R/batch/src.name;dest.parent.mkdir(exist_ok=True);shutil.copy2(src,dest);stored=str(dest.relative_to(O))
   manifest.append([stored,src.stat().st_size,hashlib.sha256(src.read_bytes()).hexdigest()])
  page.append([batch,r.get('机构',''),r.get('网址',''),r.get('状态',''),r.get('HTTP状态',''),r.get('正文字数',''),r.get('采集时间',x.get('collected_at','')),stored,r.get('错误','')])
 for r in x.get('evidence',[]):ev.append([batch,r.get('机构',''),r.get('字段',''),r.get('原文片段',''),r.get('来源',''),r.get('PDF页码',''),r.get('采集时间',''),r.get('证据类型','')])
 for r in x.get('activity_links',[]):links.append([batch,r.get('机构',''),r.get('链接标题',''),r.get('详情网址',''),r.get('发现页',''),r.get('采集时间',''),r.get('状态','')])
add('来源台账',['来源编号','资料名称','原文网址','发布机构或归属','适用时间与范围','资料类别','主要内容或定位','使用限制','采集时间'],sources,'同一网址可能支持多个来源编号。编号对应原研究阶段，未将重复引用算成独立证据。')
add('网页采集台账',['采集批次','机构','网址','状态','HTTP状态','正文字数','采集时间','包内原文文件','错误记录'],page,'保留全部采集尝试。已采集原文可在公开网页存档中查找。')
add('原文提取片段',['采集批次','机构','字段','原文片段','来源网址','PDF页码','采集时间','核验状态'],ev,'自动提取片段用于定位资料，可能包含导航和重复文本。正式采用的数据见前部专题表。')
add('活动与会籍线索',['采集批次','机构','链接标题','详情网址','发现网址','采集时间','核验状态'],links,'从网页发现的线索，保留尚未判定或尚未打开的记录。')
g=[[c.text for c in r.cells] for r in report.tables[27].rows]
add('证据缺口',g[0],g[1:],'公开资料仍缺的具体字段及用途。')
add('后续验证',['阶段','工作','判断依据'],[[c.text for c in r.cells] for r in report.tables[25].rows[1:]],'按正式报告的90天验证安排执行，尚未开展外联。')
add('文件与数据说明',['资料','记录数或范围','使用说明'],[
 ['正式报告','27页','最新六家判断、市场分析及进入评估'],['六家独立案例','6家','与正式报告一致，可单独阅读'],['合伙人决策摘要','按本轮报告更新','建议、分歧和下一步验证'],['市场机构地图',len(mp),'包含候选、替代机构与相邻服务，研究层级已列'],['平台内容明细',len(tables['平台内容明细']['rows']),'核验状态逐条保留；缺失互动不填零'],['扩展内容线索',len(tables['扩展内容线索']['rows']),'搜索卡片记录，含重复内容和个人发布者'],['网页采集尝试',len(page),'包括成功、失败及重复复核'],['可读原文存档',len(manifest),'仅公开网页正文，文件相对路径见台账'],['原文提取片段',len(ev),'机器提取，需回到来源复核'],['时间','2026-10-21','本套为既有采集结果整理，非新增一轮实时爬取'],['价格更新','以六家价格表为准','扩展历史价格保留旧记录，勿当成现价'],['登录状态','不在交付包中','本机登录状态继续保留于原项目环境']
 ],'文件名和内部标题统一按正式工作材料命名。',[25,32,95])
# Place guide first without duplicating detailed records.
sheets=[sheets[-1]]+sheets[:-1]
for s in sheets:
 with (T/(s['name']+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows([s['headers']]+s['rows'])
with (T/'公开原文文件清单.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows([['包内相对路径','字节数','SHA256']]+manifest)
(B/'research/交付工作簿数据.json').write_text(json.dumps({'sheets':sheets,'cases':cases,'counts':{'institutions':len(mp),'sources':len(sources),'pages':len(page),'raw':len(manifest),'excerpts':len(ev),'platform_contents':len(tables['平台内容明细']['rows']),'discovery_records':len(tables['扩展内容线索']['rows'])}},ensure_ascii=False,indent=2))
print(json.dumps({'sheets':len(sheets),'rows':sum(len(s['rows']) for s in sheets),'raw_files':len(manifest)},ensure_ascii=False))
