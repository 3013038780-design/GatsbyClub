from pathlib import Path
import json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).parent;R=B.parents[1];D=json.loads((R/'outputs/integrated-research-20261024/统一研究数据.json').read_text());F=json.loads((B/'执行依据与测算数据.json').read_text());O=B/'交付文件';O.mkdir(exist_ok=True)
d=Document();sec=d.sections[0];sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2);sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for name,size,bold in [('Normal',11.5,False),('Title',22,True),('Heading 1',17,True),('Heading 2',13,True)]:
 st=d.styles[name];st.font.name='Source Han Serif CN';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.line_spacing=Pt(19);st.paragraph_format.space_after=Pt(8)
 fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
 for old in list(fonts.attrib):del fonts.attrib[old]
 for edge in st.element.xpath('.//w:pBdr'):edge.getparent().remove(edge)
 for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Source Han Serif CN')
 if name=='Title':st.paragraph_format.line_spacing=Pt(30)
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(10);st.paragraph_format.keep_with_next=True
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('合伙人会议讲解稿  ').font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def p(t):d.add_paragraph(t)
def h(t):d.add_paragraph(t,'Heading 2')
def link(label,url):
 q=d.add_paragraph();q.paragraph_format.line_spacing=Pt(13);z=OxmlElement('w:hyperlink');z.set(qn('r:id'),q.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));run=OxmlElement('w:r');prop=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'255878');prop.append(col);size=OxmlElement('w:sz');size.set(qn('w:val'),'18');prop.append(size);run.append(prop);text=OxmlElement('w:t');text.text=label;run.append(text);z.append(run);q._p.append(z)
J=json.loads((B/'口语讲稿内容.json').read_text())
qa=json.loads((R/'outputs/integrated-research-20261024/qa/交付核验.json').read_text())
assert qa.get('pdf_internal_targets_verified')==16 and qa.get('report_pages')==50
for i,section in enumerate(J['sections']):
 q=d.add_paragraph(section['title'],'Heading 1')
 if i:q.paragraph_format.page_break_before=True
 else:
  q.insert_paragraph_before(J['title'],'Title')
  q.insert_paragraph_before('2026年10月24日  '+J['duration'])
 for text in section['paragraphs']:p(text)
out=O/'高端线下社群合伙人会议讲解稿.docx';d.save(out);print(out)
