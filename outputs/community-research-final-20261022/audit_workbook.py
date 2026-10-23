from pathlib import Path
import openpyxl,json,csv
B=Path(__file__).resolve().parent;old=B.parent/'matchmaking-review-20261022/交付文件/中国高端会员社群市场研究数据工作簿_31家机构.xlsx';w=openpyxl.load_workbook(old)
sheets={s.title:{'rows':[list(r) for r in s.values],'headers':[c.value for c in s[4]],'dimensions':[s.max_row,s.max_column]} for s in w}
(B/'原工作簿数据.json').write_text(json.dumps(sheets,ensure_ascii=False,indent=2,default=str))
archive=B/'旧版明细存档';archive.mkdir(exist_ok=True)
for s in w:
 with (archive/(s.title+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(list(s.values))
print('Original',len(w.sheetnames),'sheets archived. Fragmented report paragraph rows:',sum(w[n].max_row-4 for n in ['会员权益','客户画像','获客方式','活动流程','经营表现']))
